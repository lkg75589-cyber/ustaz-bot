# -*- coding: utf-8 -*-
"""
🎓 المدرّس الذكي — النسخة 3
────────────────────────────
• يدرس كتب المنهج والوزاريات (مع سنة كل وزاري) ويجيب منها مع ذكر المصدر.
• يبحث عن الملازم في قناة المالك ويرسل الروابط المناسبة فقط + رأي صريح من السوشيال ميديا (دون تحليل ملف الملزمة).
• نصائح عن مدرّسي اليوتيوب. قراءة الصور (OCR) تذهب مباشرة للنموذج دون عرض النص.
• تطبيق مصغّر: رئيسية، مواد (فصول/ملازم/مدرّسين/يوتيوب)، تحديات، حسابي — مع عرض المعادلات بشكل مرتّب.
• حماية من الضغط: طوابير محدودة، حدّ لكل طالب، قواطع للمزوّدين، خادم ويب بخيوط محدودة.
"""

import sys
import concurrent.futures
import os, re, io, json, math, time, base64, random, hashlib, hmac, logging, sqlite3, html, functools, itertools
import threading, subprocess, platform, tempfile, unicodedata, mimetypes, selectors, socket, queue, secrets, struct, zlib, socket
from contextlib import contextmanager
from pathlib import Path
from datetime import datetime, timedelta
from collections import defaultdict, deque
from concurrent.futures import ThreadPoolExecutor
from http.server import HTTPServer, BaseHTTPRequestHandler
from urllib.parse import parse_qsl, urlparse, unquote

def _load_env_files():
    """يقرأ .env أو env.txt (الهاتف لا يعرض الملفات التي تبدأ بنقطة، فيمكن رفعه باسم env.txt)."""
    here = os.path.dirname(os.path.abspath(__file__))
    for name in (".env", "env.txt", "env", "config.env"):
        fp = os.path.join(here, name)
        if not os.path.isfile(fp): continue
        try:
            from dotenv import load_dotenv
            load_dotenv(fp, override=True)
        except ImportError:
            for ln in open(fp, encoding="utf-8", errors="ignore"):
                ln = ln.strip()
                if ln and not ln.startswith("#") and "=" in ln:
                    k, v = ln.split("=", 1); os.environ[k.strip()] = v.strip().strip("'\"")
        break
_load_env_files()

import requests
import telebot
from telebot import types, apihelper

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None
try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None


# ═══════════════════════════════════════════════════════════════════
#                              الإعدادات
# ═══════════════════════════════════════════════════════════════════

def _env(name, default=""):
    return os.environ.get(name, default).strip().strip("'\"").strip()

def _int(name, default, lo=1, hi=10**7):
    try: v = int(_env(name) or default)
    except ValueError: v = default
    return max(lo, min(hi, v))

BOT_TOKEN = _env("BOT_TOKEN")
OWNER_IDS = {int(x.strip()) for x in _env("OWNER_ID").split(",") if x.strip().lstrip("-").isdigit()}

def _keys(prefix, *extra):
    """كل مفاتيح مزوّد: PREFIX ثم PREFIX_1 … PREFIX_20 (بلا تكرار)."""
    out = []
    for k in list(extra) + [_env(f"{prefix}_{i}") for i in range(1, 21)]:
        if k and k not in out: out.append(k)
    return out

GROQ_KEYS = _keys("GROQ_API_KEY", _env("GROQ_API_KEY"))
GROQ_KEY = GROQ_KEYS[0] if GROQ_KEYS else ""
GROQ_BASE = "https://api.groq.com/openai/v1"
GROQ_MODELS = [m.strip() for m in _env(
    "GROQ_MODELS", "openai/gpt-oss-120b,openai/gpt-oss-20b,llama-3.3-70b-versatile,allam-2-7b").split(",") if m.strip()]
WHISPER_MODELS = ["whisper-large-v3-turbo", "whisper-large-v3"]

GEMINI_KEYS = _keys("GEMINI_API_KEY", _env("GEMINI_API_KEY"))
GEMINI_MODEL = _env("GEMINI_MODEL", "gemini-2.5-flash")

OPENROUTER_KEYS = _keys("OPENROUTER_API_KEY", _env("OPENROUTER_API_KEY"))
OPENROUTER_BASE = "https://openrouter.ai/api/v1"
OPENROUTER_MODELS = [m.strip() for m in _env("OPENROUTER_MODELS", "openrouter/free,meta-llama/llama-3.3-70b-instruct:free").split(",") if m.strip()]

MISTRAL_KEYS = _keys("MISTRAL_API_KEY", _env("MISTRAL_API_KEY"))
MISTRAL_BASE = "https://api.mistral.ai/v1"
MISTRAL_MODELS = [m.strip() for m in _env("MISTRAL_MODELS", "mistral-small-latest,open-mistral-nemo").split(",") if m.strip()]
MISTRAL_VISION_MODEL = _env("MISTRAL_VISION_MODEL", "mistral-small-latest")

OCR_KEY = _env("OCR_API_KEY")
OCR_URL = "https://api.ocr.space/parse/image"

APP_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(APP_DIR, "static")
DB_PATH = _env("DATABASE_PATH", os.path.join(APP_DIR, "teacher_bot.sqlite3"))

WEB = {"url": _env("WEBAPP_URL").rstrip("/")}
YT_KEY = _env("YOUTUBE_API_KEY")
GH_REPO = _env("GITHUB_REPO").strip().strip("/")          # owner/repo للتحديث التلقائي من GitHub
GH_BRANCH = _env("GITHUB_BRANCH", "main")
GH_TOKEN = _env("GITHUB_TOKEN")                           # فقط إذا المستودع خاص (Private)
GH_EVERY = max(2, int(_env("AUTO_UPDATE_MIN", "10") or 10))
NGROK_TOKEN = _env("NGROK_AUTHTOKEN")
NGROK_DOMAIN = _env("NGROK_DOMAIN").replace("https://", "").strip("/")
CF_TUNNEL_TOKEN = _env("CLOUDFLARE_TUNNEL_TOKEN")      # نفق ثابت (اختياري) لرابط لا يتغيّر
WEB_PORT = _int("PORT", 8080)

IG_HANDLE = _env("INSTAGRAM_HANDLE").lstrip("@")
IG_URL = _env("INSTAGRAM_URL") or (f"https://instagram.com/{IG_HANDLE}" if IG_HANDLE else "")

CHANNEL_ENV = [x.strip() for x in _env("CHANNEL_ID").split(",") if x.strip()]

MAX_LEN = 3900
MAX_FILE = 20 * 1024 * 1024          # حد تنزيل بوتات تيليجرام
MAX_STUDENT_FILE = 8 * 1024 * 1024   # ملفات الطلاب (تلخيص) أصغر حتى لا تُثقل الخادم
MAX_AUDIO = 20 * 1024 * 1024
MAX_HISTORY = 8
RATE_LIMIT = 3
DAILY_LIMIT = _int("DAILY_LIMIT", 80)          # الحد الأساسي لكل طالب
DAILY_MAX = _int("DAILY_MAX", 200)             # أقصى حد يصله الطالب بالنقاط
PUBLISH_CHANNEL = _env("PUBLISH_CHANNEL")      # قناة نشر مساهمات الطلاب (اختياري؛ وإلا أول قناة معتمدة)
BOT_USERNAME = ""
LLM_BUDGET = 70                      # أقصى زمن كلي (ثانية) لمحاولة الإجابة عبر كل المزودين
LLM_CONCURRENCY = _int("LLM_CONCURRENCY", 6)
JOB_WORKERS = _int("JOB_WORKERS", 8)
JOB_QUEUE_MAX = _int("JOB_QUEUE_MAX", 60)
HTTP_WORKERS = _int("HTTP_WORKERS", 48)
OCR_MAX_PAGES = _int("OCR_MAX_PAGES", 45)
MAX_UPLOAD = _int("MAX_UPLOAD_MB", 40, 1, 200) * 1024 * 1024     # رفع الكتب من التطبيق
UPLOAD_DEADLINE = 300
INITDATA_MAX_AGE = 24 * 3600
STALE_SECS = 300                     # رسائل أقدم من هذا تُتجاهل (بعد توقف/ضغط)

logging.basicConfig(level=logging.INFO, format="%(asctime)s │ %(levelname)-7s │ %(message)s", datefmt="%H:%M:%S")
log = logging.getLogger("mudarris")
logging.getLogger("urllib3").setLevel(logging.WARNING)

try:
    apihelper.CONNECT_TIMEOUT = 10
    apihelper.READ_TIMEOUT = 40
    apihelper.RETRY_ON_ERROR = False
except Exception:
    pass


# ═══════════════════════════════════════════════════════════════════
#                              المواد
# ═══════════════════════════════════════════════════════════════════

SUBJECTS = {
    "math":      {"name": "الرياضيات", "emoji": "📐"},
    "physics":   {"name": "الفيزياء",  "emoji": "⚛️"},
    "chemistry": {"name": "الكيمياء",  "emoji": "🧪"},
    "biology":   {"name": "الأحياء",   "emoji": "🧬"},
    "english":   {"name": "الإنكليزي", "emoji": "🔤"},
    "arabic":    {"name": "العربي",    "emoji": "📖"},
    "islamic":   {"name": "الإسلامية", "emoji": "🕌"},
    "history":   {"name": "التاريخ",   "emoji": "🏛️"},
    "geography": {"name": "الجغرافية", "emoji": "🌍"},
    "general":   {"name": "عام",       "emoji": "📚"},
}

SUBJECT_PROMPTS = {
    "biology":   "المادة الآن: الأحياء. ركّز على الخلية، الوراثة، الأنسجة، الأجهزة.",
    "physics":   "المادة الآن: الفيزياء. اكتب القوانين مع الوحدات واشرح الرموز.",
    "chemistry": "المادة الآن: الكيمياء. وازن المعادلات واشرح الخطوات.",
    "math":      "المادة الآن: الرياضيات. اعرض الحل خطوة خطوة بدقة.",
    "arabic":    "المادة الآن: اللغة العربية (نحو، صرف، بلاغة، أدب).",
    "english":   "المادة الآن: الإنكليزية. اشرح بالعربي وأعطِ أمثلة إنكليزية.",
    "islamic":   "المادة الآن: التربية الإسلامية. التزم بالدقة ولا تُفتِ بما لا تعلم.",
    "history":   "المادة الآن: التاريخ. رتّب الأحداث زمنياً.",
    "geography": "المادة الآن: الجغرافية.",
}

def subj_name(key):
    return SUBJECTS.get(key or "", {}).get("name", "عام")


# ═══════════════════════════════════════════════════════════════════
#                       معالجة النص العربي + البحث
# ═══════════════════════════════════════════════════════════════════

_TASHKEEL = re.compile("[\u0617-\u061A\u064B-\u0652\u0640\u0670]")
_AR_DIGITS = {**{0x660 + i: str(i) for i in range(10)}, **{0x6F0 + i: str(i) for i in range(10)}}
_STOP = set("""من في على الى إلى عن هو هي هم ما ماذا هذا هذه ذلك تلك الذي التي الذين كان كانت يكون ان إن أن او أو ثم قد لا لم لن
بين عند كل اي أي مع او اذا إذا كيف لماذا ليش شنو شلون شكو وين اشرح اشرحلي وضح وضّح عرف عرّف اريد أريد ممكن لو بس هل
لي لنا انا اني احنا عليه عليها منه منها له لها به بها فيه فيها حتى او ايضا أيضا كما لكن غير بعد قبل ذلك هنا هناك""".split())

def ar_norm(s):
    s = unicodedata.normalize("NFKC", s or "")
    s = _TASHKEEL.sub("", s).translate(_AR_DIGITS).lower()
    s = re.sub("[\u0625\u0623\u0622\u0671]", "ا", s)
    return s.replace("ى", "ي").replace("ة", "ه")

_STOP_N = {ar_norm(w) for w in _STOP}

def _stem(w):
    if w.startswith("وال") and len(w) > 5: return w[3:]
    if w[:3] in ("بال", "كال", "فال") and len(w) > 5: return w[3:]
    if w.startswith("ال") and len(w) > 4: return w[2:]
    if w.startswith("لل") and len(w) > 4: return w[2:]
    return w

def tokens(text):
    out = []
    for w in re.findall(r"\w+", ar_norm(text)):
        if w in _STOP_N: continue
        w = _stem(w)
        if len(w) >= 2 or w.isdigit(): out.append(w)
    return out

_SUBJ_KW = [
    ("biology", ["احياء", "علم الاحياء", "biology"]),
    ("physics", ["فيزياء", "physics"]),
    ("chemistry", ["كيمياء", "chemistry"]),
    ("math", ["رياضيات", "تفاضل", "تكامل", "math"]),
    ("arabic", ["عربي", "نحو", "قواعد", "ادب", "بلاغه"]),
    ("english", ["انكليزي", "انجليزي", "english"]),
    ("islamic", ["اسلاميه", "اسلامي", "فقه", "قران", "حديث"]),
    ("history", ["تاريخ", "history"]),
    ("geography", ["جغرافيه", "جغرافية", "geography"]),
]
_SUBJ_KW = [(k, [ar_norm(x) for x in v]) for k, v in _SUBJ_KW]
_SUBJ_WORDS = {_stem(w) for _, kws in _SUBJ_KW for kw in kws for w in kw.split()}
_SUBJ_WORDS |= {t for v in SUBJECTS.values() for t in tokens(v["name"])}

def guess_subject(text):
    t = ar_norm(text)
    for key, kws in _SUBJ_KW:
        if any(kw in t for kw in kws): return key
    return None

def subject_key(s):
    s = (s or "").strip()
    if s in SUBJECTS: return s
    for k, v in SUBJECTS.items():
        if v["name"] == s: return k
    return guess_subject(s) or "general"

CALL_RE = re.compile("^\\s*[\u0623\u0627]ستاذ(?![\u0621-\u064A])[\\s,،:;\\-–—!.؟?]*")

def strip_call(text):
    """يرجع نص السؤال إذا بدأت الرسالة بـ «استاذ/أستاذ»، وإلا None."""
    t = _TASHKEEL.sub("", text or "")
    m = CALL_RE.match(t)
    return t[m.end():].strip() if m else None

def detect_year(*texts):
    """يستخرج سنة الوزاري (2022-2023 → 2023)."""
    for t in texts:
        ys = [int(y) for y in re.findall(r"(?<!\d)(19[89]\d|20[0-3]\d)(?!\d)", ar_norm(t or ""))]
        if ys: return max(ys)
    return None

def detect_round(*texts):
    m_ = {"اول": "الدور الأول", "1": "الدور الأول", "ثاني": "الدور الثاني", "2": "الدور الثاني", "ثالث": "الدور الثالث", "3": "الدور الثالث"}
    for t in texts:
        m = re.search(r"دور\s*(?:ال)?(اول|ثاني|ثالث|[123])(?!\d)", ar_norm(t or ""))
        if m: return m_[m.group(1)]
    return None


# ═══════════════════════════════════════════════════════════════════
#                  المعادلات: ترتيبها للتطبيق وتبسيطها لتيليجرام
# ═══════════════════════════════════════════════════════════════════

_LATEX_CMD = re.compile(r"\\(frac|dfrac|sqrt|sum|int|lim|times|cdot|le|ge|leq|geq|neq|pm|rightarrow|to|infty|alpha|beta|theta|pi|Delta|text|mathrm|left|right|vec|bar|overline|rightleftharpoons)\b")
_HAS_AR = re.compile("[\u0600-\u06FF]")

def normalize_math(t):
    """يوحّد حدود المعادلات إلى $...$ و $$...$$ ويجمع الأسطر الرياضية العارية المتتالية في معادلة واحدة."""
    t = re.sub(r"\\\[(.+?)\\\]", lambda m: "\n$$" + " ".join(m.group(1).split()) + "$$\n", t, flags=re.S)
    t = re.sub(r"\\\((.+?)\\\)", lambda m: "$" + " ".join(m.group(1).split()) + "$", t, flags=re.S)
    t = re.sub(r"\$\$(.+?)\$\$", lambda m: "\n$$" + " ".join(m.group(1).split()) + "$$\n", t, flags=re.S)   # معادلة $$ مكسورة على أسطر → سطر واحد
    out, buf = [], []
    def flush():
        if buf: out.append("$$" + " ".join(buf) + "$$"); buf.clear()
    for ln in t.split("\n"):
        st = ln.strip().strip("`")
        bare = bool(st) and "$" not in ln and not _HAS_AR.search(ln) and (_LATEX_CMD.search(ln) or (buf and re.fullmatch(r"[\w\\{}()^_=+\-*/.,;:|\[\]<>!' ]+", st)))
        if bare: buf.append(st)
        else:
            flush(); out.append(ln)
    flush()
    return "\n".join(out)

_SUP = str.maketrans("0123456789+-=()ni", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿⁱ")
_SUB = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")
_GREEK = {"alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "Delta": "Δ", "theta": "θ", "lambda": "λ", "mu": "μ", "pi": "π",
          "sigma": "σ", "Sigma": "Σ", "omega": "ω", "Omega": "Ω", "phi": "φ", "epsilon": "ε", "rho": "ρ", "tau": "τ", "eta": "η",
          "psi": "ψ", "chi": "χ", "nu": "ν", "kappa": "κ", "zeta": "ζ", "xi": "ξ"}
_SYM = {"times": "×", "cdot": "·", "div": "÷", "pm": "±", "mp": "∓", "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥", "neq": "≠", "ne": "≠",
        "approx": "≈", "infty": "∞", "rightarrow": "→", "to": "→", "Rightarrow": "⇒", "leftarrow": "←", "leftrightarrow": "↔",
        "rightleftharpoons": "⇌", "degree": "°", "circ": "°", "sum": "Σ", "int": "∫", "partial": "∂", "nabla": "∇", "in": "∈",
        "cup": "∪", "cap": "∩", "subset": "⊂", "therefore": "∴", "because": "∵", "ldots": "…", "dots": "…", "propto": "∝",
        "angle": "∠", "perp": "⊥", "parallel": "∥", "prime": "′", "quad": " ", "qquad": "  "}

def _script(m, table, mark):
    body = m.group(1) if m.group(1) is not None else m.group(2)
    conv = body.translate(table)
    return conv if conv.translate({ord(c): None for c in "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿⁱ₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎"}) == "" else f"{mark}({body})"

def latex_to_unicode(s):
    """تحويل LaTeX إلى نص مقروء في تيليجرام (الذي لا يعرض LaTeX)."""
    s = s.replace("$$", "\n").replace("$", "")
    for _ in range(3):
        s = re.sub(r"\\[dt]?frac\{([^{}]*)\}\{([^{}]*)\}",
                   lambda m: (f"{m.group(1)}/{m.group(2)}" if re.fullmatch(r"\w+", m.group(1)) and re.fullmatch(r"\w+", m.group(2))
                              else f"({m.group(1)})/({m.group(2)})"), s)
    s = re.sub(r"\\sqrt\[(\d)\]\{([^{}]*)\}", lambda m: {"3": "∛", "4": "∜"}.get(m.group(1), m.group(1) + "√") + f"({m.group(2)})", s)
    s = re.sub(r"\\sqrt\{([^{}]*)\}", lambda m: "√" + (m.group(1) if re.fullmatch(r"\w", m.group(1)) else f"({m.group(1)})"), s)
    s = re.sub(r"\\(?:text|mathrm|mathbf|mathit|textbf|operatorname|boldsymbol)\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\(?:left|right|big|Big|displaystyle)\b\s*", "", s)
    s = re.sub(r"\^\{([^{}]*)\}|\^(\w)", lambda m: _script(m, _SUP, "^"), s)
    s = re.sub(r"_\{([^{}]*)\}|_(\w)", lambda m: _script(m, _SUB, "_"), s)
    s = re.sub(r"\\([A-Za-z]+)", lambda m: _GREEK.get(m.group(1)) or _SYM.get(m.group(1)) or m.group(1), s)
    s = re.sub(r"\\[,;:! ]", " ", s).replace("\\\\", "\n").replace("{", "").replace("}", "")
    return re.sub(r"[ \t]{2,}", " ", s)


# ═══════════════════════════════════════════════════════════════════
#                           قاعدة البيانات
# ═══════════════════════════════════════════════════════════════════

_DB_LOCK = threading.RLock()
HAS_FTS = True

@contextmanager
def db():
    with _DB_LOCK:
        con = sqlite3.connect(DB_PATH, timeout=30)
        con.row_factory = sqlite3.Row
        try:
            con.execute("PRAGMA busy_timeout=30000")
            yield con
            con.commit()
        except Exception:
            con.rollback()
            raise
        finally:
            con.close()

def now_iso():
    return datetime.utcnow().isoformat(timespec="seconds")

def today_key():
    return datetime.now().strftime("%Y-%m-%d")

def _cols(con, table):
    return {r["name"] for r in con.execute(f"PRAGMA table_info({table})")}

def _add_col(con, table, col, ddl):
    if col not in _cols(con, table):
        con.execute(f"ALTER TABLE {table} ADD COLUMN {col} {ddl}")

def mat_norm(*parts):
    return " ".join(tokens(" ".join(p for p in parts if p)))

def db_init():
    d = os.path.dirname(DB_PATH)
    if d: os.makedirs(d, exist_ok=True)
    with db() as con:
        con.execute("PRAGMA journal_mode=WAL")
        con.execute("PRAGMA synchronous=NORMAL")
        con.executescript("""
        CREATE TABLE IF NOT EXISTS users(user_id INTEGER PRIMARY KEY, username TEXT, first_name TEXT, last_name TEXT,
            joined_at TEXT NOT NULL, last_seen TEXT NOT NULL, banned INTEGER NOT NULL DEFAULT 0,
            role TEXT NOT NULL DEFAULT 'student', grade TEXT, preferred_subject TEXT, private_started INTEGER NOT NULL DEFAULT 0);
        CREATE TABLE IF NOT EXISTS messages(id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL,
            chat_id INTEGER NOT NULL DEFAULT 0, role TEXT NOT NULL, content TEXT NOT NULL,
            source TEXT NOT NULL DEFAULT 'telegram', created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS materials(id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, subject TEXT,
            kind TEXT NOT NULL DEFAULT 'link', telegram_url TEXT NOT NULL, description TEXT DEFAULT '',
            added_by INTEGER NOT NULL DEFAULT 0, active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL,
            channel_id INTEGER, message_id INTEGER, source_id INTEGER);
        CREATE TABLE IF NOT EXISTS sources(id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, subject TEXT,
            kind TEXT NOT NULL DEFAULT 'book', pages INTEGER DEFAULT 0, chars INTEGER DEFAULT 0, chunks INTEGER DEFAULT 0,
            added_by INTEGER DEFAULT 0, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS channels(chat_id INTEGER PRIMARY KEY, title TEXT, username TEXT, approved INTEGER NOT NULL DEFAULT 0, added_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, event TEXT NOT NULL, details TEXT DEFAULT '', created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS daily_usage(user_id INTEGER NOT NULL, day TEXT NOT NULL, requests INTEGER NOT NULL DEFAULT 0, PRIMARY KEY(user_id,day));
        CREATE TABLE IF NOT EXISTS quiz_results(id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, subject TEXT, score INTEGER NOT NULL, total INTEGER NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS chapters(id INTEGER PRIMARY KEY AUTOINCREMENT, subject TEXT NOT NULL, title TEXT NOT NULL, ord INTEGER NOT NULL DEFAULT 0, UNIQUE(subject,title));
        CREATE TABLE IF NOT EXISTS teachers(id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, subject TEXT, youtube_url TEXT DEFAULT '',
            notes TEXT DEFAULT '', active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS mistakes(id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, subject TEXT, chapter_id INTEGER, q TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS activity(user_id INTEGER NOT NULL, day TEXT NOT NULL, points INTEGER NOT NULL DEFAULT 0, focus_min INTEGER NOT NULL DEFAULT 0, PRIMARY KEY(user_id,day));
        CREATE TABLE IF NOT EXISTS login_tokens(token TEXT PRIMARY KEY, code TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'pending', user_id INTEGER, ua TEXT, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS sessions(token_hash TEXT PRIMARY KEY, user_id INTEGER NOT NULL, created_at TEXT NOT NULL, last_seen TEXT NOT NULL, ua TEXT);
        CREATE TABLE IF NOT EXISTS focus_sessions(id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, subject TEXT, minutes INTEGER NOT NULL, note TEXT DEFAULT '', created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS admins(user_id INTEGER PRIMARY KEY, added_by INTEGER, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS opinions(key TEXT PRIMARY KEY, text TEXT NOT NULL, created_at TEXT NOT NULL, ttl_h INTEGER NOT NULL DEFAULT 168);
        CREATE TABLE IF NOT EXISTS daily_challenge(day TEXT PRIMARY KEY, subject TEXT, data TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS challenge_done(user_id INTEGER NOT NULL, day TEXT NOT NULL, score INTEGER NOT NULL, total INTEGER NOT NULL, PRIMARY KEY(user_id,day));
        """)
        global HAS_FTS
        try:
            con.execute("CREATE VIRTUAL TABLE IF NOT EXISTS chunks USING fts5(norm, raw UNINDEXED, source_id UNINDEXED, subject UNINDEXED, page UNINDEXED)")
            HAS_FTS = True
        except sqlite3.OperationalError:
            log.warning("FTS5 غير متاح في SQLite هنا — سأستخدم بحثاً بديلاً أبطأ قليلاً")
            con.execute("CREATE TABLE IF NOT EXISTS chunks(norm TEXT, raw TEXT, source_id INTEGER, subject TEXT, page INTEGER)")
            con.execute("CREATE INDEX IF NOT EXISTS idx_chunks_src ON chunks(source_id)")
            HAS_FTS = False
        # ترقية قواعد البيانات القديمة
        _add_col(con, "messages", "chat_id", "INTEGER NOT NULL DEFAULT 0")
        for col, ddl in (("channel_id", "INTEGER"), ("message_id", "INTEGER"), ("source_id", "INTEGER"), ("norm", "TEXT DEFAULT ''")):
            _add_col(con, "materials", col, ddl)
        for col, ddl in (("private_started", "INTEGER NOT NULL DEFAULT 0"), ("hide_lb", "INTEGER NOT NULL DEFAULT 0"), ("last_ctx", "TEXT")):
            _add_col(con, "users", col, ddl)
        _add_col(con, "sources", "year", "INTEGER")
        _add_col(con, "users", "focus_goal", "INTEGER DEFAULT 60")
        for col in ("mkind", "teacher", "part", "batch", "mtype", "file_ids", "msg_ids", "src"):
            _add_col(con, "materials", col, "TEXT DEFAULT ''")
        _add_col(con, "materials", "status", "TEXT DEFAULT 'approved'")
        _add_col(con, "sources", "round", "TEXT")
        _add_col(con, "quiz_results", "chapter_id", "INTEGER")
        _add_col(con, "quiz_results", "kind", "TEXT")
        con.execute("CREATE INDEX IF NOT EXISTS idx_messages_chat ON messages(user_id,chat_id,id DESC)")
        con.execute("CREATE INDEX IF NOT EXISTS idx_materials_msg ON materials(channel_id,message_id)")
        con.execute("CREATE INDEX IF NOT EXISTS idx_quiz_user ON quiz_results(user_id,subject,id DESC)")
        for r in con.execute("SELECT id,title,description FROM materials WHERE norm IS NULL OR norm=''").fetchall():
            con.execute("UPDATE materials SET norm=? WHERE id=?", (mat_norm(r["title"], r["description"]), r["id"]))

def db_event(uid, event, details=""):
    try:
        with db() as con:
            con.execute("INSERT INTO events(user_id,event,details,created_at) VALUES(?,?,?,?)", (uid, event, str(details)[:500], now_iso()))
    except Exception as e:
        log.warning("db_event failed: %s", e)

_UPS = {}
def db_upsert_user(user, source="telegram"):
    """user: كائن تيليجرام أو dict. الكتابة مخفّفة (مرة كل دقيقتين لكل مستخدم) لتقليل الحمل."""
    if not user: return
    g = (lambda k, d=None: user.get(k, d)) if isinstance(user, dict) else (lambda k, d=None: getattr(user, k, d))
    uid = int(g("id")); t = time.time()
    if source != "telegram_start" and t - _UPS.get(uid, 0) < 120: return
    _UPS[uid] = t
    now = now_iso(); started = 1 if source == "telegram_start" else 0
    with db() as con:
        con.execute(
            "INSERT INTO users(user_id,username,first_name,last_name,joined_at,last_seen,private_started) VALUES(?,?,?,?,?,?,?) "
            "ON CONFLICT(user_id) DO UPDATE SET username=excluded.username, first_name=excluded.first_name, "
            "last_name=excluded.last_name, last_seen=excluded.last_seen, private_started=MAX(users.private_started,excluded.private_started)",
            (uid, g("username"), g("first_name", "") or "", g("last_name", "") or "", now, now, started))

def db_get_user(uid):
    with db() as con:
        r = con.execute("SELECT * FROM users WHERE user_id=?", (int(uid),)).fetchone()
    return dict(r) if r else {}

_BAN = {}
def db_is_banned(uid):
    c = _BAN.get(uid)
    if c and time.time() - c[1] < 30: return c[0]
    with db() as con:
        r = con.execute("SELECT banned FROM users WHERE user_id=?", (int(uid),)).fetchone()
    v = bool(r and r["banned"]); _BAN[uid] = (v, time.time())
    return v

def db_save_message(uid, role, content, source="telegram", chat_id=0):
    if not content: return
    with db() as con:
        con.execute("INSERT OR IGNORE INTO users(user_id,joined_at,last_seen) VALUES(?,?,?)", (int(uid), now_iso(), now_iso()))
        con.execute("INSERT INTO messages(user_id,chat_id,role,content,source,created_at) VALUES(?,?,?,?,?,?)",
                    (int(uid), int(chat_id), role, str(content)[:12000], source, now_iso()))
        if random.random() < 0.1:
            con.execute("DELETE FROM messages WHERE user_id=? AND id NOT IN (SELECT id FROM messages WHERE user_id=? ORDER BY id DESC LIMIT 400)", (int(uid), int(uid)))

def db_history(uid, chat_id=0, limit=12):
    with db() as con:
        rows = con.execute("SELECT role,content,source,created_at FROM messages WHERE user_id=? AND chat_id=? ORDER BY id DESC LIMIT ?",
                           (int(uid), int(chat_id), int(limit))).fetchall()
    return list(reversed([dict(r) for r in rows]))

def db_clear_history(uid, chat_id=None):
    with db() as con:
        if chat_id is None: con.execute("DELETE FROM messages WHERE user_id=?", (int(uid),))
        else: con.execute("DELETE FROM messages WHERE user_id=? AND chat_id=?", (int(uid), int(chat_id)))

def db_usage_today(uid):
    with db() as con:
        r = con.execute("SELECT requests FROM daily_usage WHERE user_id=? AND day=?", (int(uid), today_key())).fetchone()
    return int(r["requests"]) if r else 0

def db_student_summary(uid):
    """ملخّص أداء الطالب الأخير (يُدمج في تعليمات المدرّس لهذا الطالب فقط)."""
    with db() as con:
        rows = con.execute("SELECT subject,score,total FROM quiz_results WHERE user_id=? AND COALESCE(kind,'')!='challenge' ORDER BY id DESC LIMIT 12", (int(uid),)).fetchall()
    agg = defaultdict(lambda: [0, 0])
    for r in rows:
        agg[r["subject"] or "general"][0] += r["score"]; agg[r["subject"] or "general"][1] += r["total"]
    return "، ".join(f"{subj_name(k)} {round(100 * a / b)}٪" for k, (a, b) in agg.items() if b)

# ── النشاط والسلسلة والتحديات ──
def touch_activity(uid, pts=1, focus=0):
    try:
        with db() as con:
            con.execute("INSERT INTO activity(user_id,day,points,focus_min) VALUES(?,?,?,?) ON CONFLICT(user_id,day) "
                        "DO UPDATE SET points=points+excluded.points, focus_min=focus_min+excluded.focus_min",
                        (int(uid), today_key(), int(pts), int(focus)))
    except Exception as e:
        log.warning("activity: %s", e)

def streak_of(uid):
    with db() as con:
        days = {r["day"] for r in con.execute("SELECT day FROM activity WHERE user_id=? AND (points>0 OR focus_min>0) ORDER BY day DESC LIMIT 400", (int(uid),))}
    d = datetime.now().date()
    if d.isoformat() not in days: d -= timedelta(days=1)
    n = 0
    while d.isoformat() in days: n += 1; d -= timedelta(days=1)
    return n

def activity_totals(uid):
    wk = (datetime.now() - timedelta(days=6)).strftime("%Y-%m-%d")
    with db() as con:
        w = con.execute("SELECT COALESCE(SUM(points),0) p, COALESCE(SUM(focus_min),0) f FROM activity WHERE user_id=? AND day>=?", (int(uid), wk)).fetchone()
        t = con.execute("SELECT COALESCE(SUM(points),0) p, COALESCE(SUM(focus_min),0) f FROM activity WHERE user_id=?", (int(uid),)).fetchone()
        td = con.execute("SELECT COALESCE(focus_min,0) f FROM activity WHERE user_id=? AND day=?", (int(uid), today_key())).fetchone()
    return {"week_points": w["p"], "total_points": t["p"], "focus_total": t["f"], "focus_today": td["f"] if td else 0}

RANKS = [(0, "مبتدئ"), (60, "مجتهد"), (200, "متميّز"), (500, "نابغة"), (1000, "عالِم"), (2000, "أسطورة")]

def rank_of(points):
    cur, nxt = RANKS[0], None
    for i, r in enumerate(RANKS):
        if points >= r[0]:
            cur = r; nxt = RANKS[i + 1] if i + 1 < len(RANKS) else None
    return {"title": cur[1], "from": cur[0], "next": nxt[1] if nxt else None, "next_at": nxt[0] if nxt else None}

def contrib_title(n):
    return "ناشر أسطوري" if n >= 40 else "ناشر مميّز" if n >= 15 else "مساهم نشط" if n >= 5 else "مساهم" if n >= 1 else ""

def contrib_counts(uid):
    with db() as con:
        rows = con.execute("SELECT COALESCE(NULLIF(mkind,''),'mulzama') k, COUNT(*) n FROM materials WHERE added_by=? AND src='student' AND active=1 GROUP BY k", (int(uid),)).fetchall()
    return {r["k"]: r["n"] for r in rows}

def player_ref(uid):
    """مرجع عام للّاعب (لا يكشف رقم تيليجرام)."""
    return hmac.new((BOT_TOKEN or "x").encode(), f"lb:{uid}".encode(), hashlib.sha256).hexdigest()[:12]

def leaderboard(me, limit=10):
    wk = (datetime.now() - timedelta(days=6)).strftime("%Y-%m-%d")
    with db() as con:
        rows = con.execute("SELECT a.user_id uid, SUM(a.points) p, u.first_name n FROM activity a JOIN users u ON u.user_id=a.user_id "
                           "WHERE a.day>=? AND u.hide_lb=0 AND u.banned=0 GROUP BY a.user_id HAVING p>0 ORDER BY p DESC LIMIT ?", (wk, limit)).fetchall()
        uids = [r["uid"] for r in rows]
        life = {}
        if uids:
            life = {r["user_id"]: r["p"] for r in con.execute(f"SELECT user_id, SUM(points) p FROM activity WHERE user_id IN ({','.join('?' * len(uids))}) GROUP BY user_id", uids)}
    return [{"name": (r["n"] or "طالب")[:18], "points": r["p"], "me": r["uid"] == me, "ref": player_ref(r["uid"]), "title": rank_of(life.get(r["uid"], 0))["title"]} for r in rows]

def player_by_ref(ref):
    wk = (datetime.now() - timedelta(days=6)).strftime("%Y-%m-%d")
    with db() as con:
        rows = con.execute("SELECT a.user_id uid FROM activity a JOIN users u ON u.user_id=a.user_id WHERE a.day>=? AND u.hide_lb=0 AND u.banned=0 "
                           "GROUP BY a.user_id HAVING SUM(a.points)>0 ORDER BY SUM(a.points) DESC LIMIT 50", (wk,)).fetchall()
    for pos, r in enumerate(rows, 1):
        if player_ref(r["uid"]) == ref: return r["uid"], pos
    return None, 0

def player_profile(uid):
    u = db_get_user(uid); a = activity_totals(uid); cc = contrib_counts(uid); rk = rank_of(a["total_points"])
    n = sum(cc.values())
    return {"name": (u.get("first_name") or "طالب")[:18], "title": rk["title"], "next": rk["next"], "next_at": rk["next_at"], "from": rk["from"],
            "total_points": a["total_points"], "week_points": a["week_points"], "mulzama": cc.get("mulzama", 0), "summary": cc.get("summary", 0),
            "contrib_title": contrib_title(n), "focus_total": a["focus_total"], "focus_today": a["focus_today"], "streak": streak_of(uid)}

def focus_log(uid, minutes, subject="", note=""):
    with db() as con:
        con.execute("INSERT INTO focus_sessions(user_id,subject,minutes,note,created_at) VALUES(?,?,?,?,?)", (int(uid), subject or "", int(minutes), (note or "")[:80], now_iso()))
        con.execute("DELETE FROM focus_sessions WHERE user_id=? AND id NOT IN (SELECT id FROM focus_sessions WHERE user_id=? ORDER BY id DESC LIMIT 200)", (int(uid), int(uid)))

def focus_stats(uid):
    days = [(datetime.now() - timedelta(days=i)).strftime("%Y-%m-%d") for i in range(6, -1, -1)]
    with db() as con:
        rows = {r["day"]: r["focus_min"] for r in con.execute("SELECT day, focus_min FROM activity WHERE user_id=? AND day>=?", (int(uid), days[0]))}
        best = con.execute("SELECT COALESCE(MAX(focus_min),0) m FROM activity WHERE user_id=?", (int(uid),)).fetchone()["m"]
        n = con.execute("SELECT COUNT(*) n FROM focus_sessions WHERE user_id=?", (int(uid),)).fetchone()["n"]
        rec = [dict(r) for r in con.execute("SELECT subject, minutes, created_at FROM focus_sessions WHERE user_id=? ORDER BY id DESC LIMIT 8", (int(uid),))]
        g = con.execute("SELECT focus_goal g FROM users WHERE user_id=?", (int(uid),)).fetchone()
    a = activity_totals(uid); week = [{"day": d, "min": rows.get(d, 0)} for d in days]
    return {"today": week[-1]["min"], "goal": (g["g"] if g and g["g"] else 60), "week": week, "week_total": sum(w["min"] for w in week), "total": a["focus_total"],
            "sessions": n, "best": best, "streak": streak_of(uid),
            "recent": [{"subject": subj_name(r["subject"]) if r["subject"] in SUBJECTS else "", "minutes": r["minutes"], "at": r["created_at"]} for r in rec]}

def set_ctx(uid, subject, chapter_id=0, title=""):
    if subject not in SUBJECTS or subject == "general": return
    try:
        with db() as con:
            con.execute("UPDATE users SET last_ctx=? WHERE user_id=?", (json.dumps({"subject": subject, "chapter_id": chapter_id or 0, "title": title[:80]}, ensure_ascii=False), int(uid)))
    except Exception: pass

def get_ctx(uid):
    try: return json.loads(db_get_user(uid).get("last_ctx") or "null")
    except Exception: return None

# ── المشرفون (يُديرهم المالك من التطبيق أو الأوامر) ──
_ADM = {"set": set(), "t": 0.0}

def admins_cached():
    if time.time() - _ADM["t"] > 60:
        try:
            with db() as con: _ADM["set"] = {r["user_id"] for r in con.execute("SELECT user_id FROM admins")}
            _ADM["t"] = time.time()
        except Exception as e:
            log.warning("admins: %s", e)
    return _ADM["set"]

def is_admin(uid):
    return uid in OWNER_IDS or uid in admins_cached()

def admin_add(uid, by=0):
    with db() as con:
        con.execute("INSERT OR IGNORE INTO admins(user_id,added_by,created_at) VALUES(?,?,?)", (int(uid), int(by), now_iso()))
    _ADM["t"] = 0.0

def admin_del(uid):
    with db() as con: n = con.execute("DELETE FROM admins WHERE user_id=?", (int(uid),)).rowcount
    _ADM["t"] = 0.0
    return n

def admin_list():
    with db() as con:
        rows = con.execute("SELECT a.user_id id, u.first_name n, u.username un FROM admins a LEFT JOIN users u ON u.user_id=a.user_id ORDER BY a.created_at").fetchall()
    return [{"id": r["id"], "name": r["n"] or r["un"] or "", "owner": False} for r in rows]

# ── الفصول والمدرّسون والأخطاء ──
def chapters_of(subject):
    with db() as con:
        return [dict(r) for r in con.execute("SELECT id,title,ord FROM chapters WHERE subject=? ORDER BY ord,id", (subject,))]

def chapter_get(cid):
    with db() as con:
        r = con.execute("SELECT id,subject,title FROM chapters WHERE id=?", (int(cid),)).fetchone()
    return dict(r) if r else None

def chapter_stats(uid, subject):
    """chapter_id → (score,total) من آخر 8 اختبارات لكل فصل (0 = كل المادة)."""
    with db() as con:
        rows = con.execute("SELECT COALESCE(chapter_id,0) c, score, total FROM quiz_results WHERE user_id=? AND subject=? AND COALESCE(kind,'')!='challenge' ORDER BY id DESC LIMIT 300", (int(uid), subject)).fetchall()
    agg, cnt = defaultdict(lambda: [0, 0]), defaultdict(int)
    for r in rows:
        if cnt[r["c"]] >= 8: continue
        cnt[r["c"]] += 1; agg[r["c"]][0] += r["score"]; agg[r["c"]][1] += r["total"]
    return {c: (a, b, cnt[c]) for c, (a, b) in agg.items() if b}

def teachers_of(subject=None):
    with db() as con:
        if subject: rows = con.execute("SELECT * FROM teachers WHERE active=1 AND subject=? ORDER BY id DESC LIMIT 40", (subject,)).fetchall()
        else: rows = con.execute("SELECT * FROM teachers WHERE active=1 ORDER BY id DESC LIMIT 100").fetchall()
    return [dict(r) for r in rows]

def material_add(title, subject, url, caption="", kind="link", by=0, channel_id=None, message_id=None,
                 mkind=None, teacher="", part="", batch="", mtype="", file_ids="", msg_ids="", src="", status="approved", active=1):
    mk = mkind or ("summary" if re.search("ملخص|مختصر", ar_norm(f"{title} {caption or ''}")) else "mulzama")
    norm = mat_norm(title, caption, teacher, part, batch, mtype, "ملخص" if mk == "summary" else "ملزمة")
    with db() as con:
        return con.execute("INSERT INTO materials(title,subject,kind,telegram_url,description,added_by,created_at,channel_id,message_id,norm,mkind,teacher,part,batch,mtype,file_ids,msg_ids,src,status,active) "
                           "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                           (title, subject, kind, url or "", (caption or "")[:400], by, now_iso(), channel_id, message_id, norm, mk, teacher or "", part or "", batch or "", mtype or "", file_ids or "", msg_ids or "", src or "", status, active)).lastrowid

def save_mistakes(uid, subject, chapter_id, items):
    clean = []
    for q in items[:12]:
        try:
            opts = [str(o)[:200] for o in q["options"]][:5]; ans = int(q["answer"])
            if 2 <= len(opts) and 0 <= ans < len(opts):
                clean.append(json.dumps({"q": str(q["q"])[:700], "options": opts, "answer": ans, "explain": str(q.get("explain", ""))[:500]}, ensure_ascii=False))
        except Exception: continue
    if not clean: return
    with db() as con:
        for c in clean:
            con.execute("INSERT INTO mistakes(user_id,subject,chapter_id,q,created_at) VALUES(?,?,?,?,?)", (int(uid), subject, int(chapter_id or 0), c, now_iso()))
        con.execute("DELETE FROM mistakes WHERE user_id=? AND id NOT IN (SELECT id FROM mistakes WHERE user_id=? ORDER BY id DESC LIMIT 80)", (int(uid), int(uid)))

def mistakes_quiz(uid, subject="", count=5):
    with db() as con:
        if subject: rows = con.execute("SELECT id,q FROM mistakes WHERE user_id=? AND subject=? ORDER BY random() LIMIT ?", (int(uid), subject, count)).fetchall()
        else: rows = con.execute("SELECT id,q FROM mistakes WHERE user_id=? ORDER BY random() LIMIT ?", (int(uid), count)).fetchall()
    out = []
    for r in rows:
        try: q = json.loads(r["q"]); q["mid"] = r["id"]; q["src"] = "من أخطائك السابقة"; out.append(q)
        except Exception: continue
    return out

def delete_mistakes(uid, ids):
    ids = [int(i) for i in ids if str(i).isdigit()][:20]
    if not ids: return
    with db() as con:
        con.execute(f"DELETE FROM mistakes WHERE user_id=? AND id IN ({','.join('?' * len(ids))})", [int(uid)] + ids)


# ═══════════════════════════════════════════════════════════════════
#             المصادر (كتب / وزاريات / أجوبة معتمدة) والبحث فيها
# ═══════════════════════════════════════════════════════════════════

KIND_LABEL = {"book": "كتاب", "wazari": "وزاري", "qa": "إجابة معتمدة", "mulzama": "ملزمة"}

def chunk_pages(pages, size=900, overlap=160):
    """pages: [(رقم_الصفحة, نص)] → [(page, chunk)] مع تداخل بسيط، ولا يعبر حدود الصفحة."""
    out = []
    for page, text in pages:
        lines = [ln.strip() for ln in (text or "").splitlines() if ln.strip()]
        buf, n = [], 0
        for ln in lines:
            while len(ln) > size * 1.5:
                out.append((page, ln[:size])); ln = ln[size - overlap:]
            buf.append(ln); n += len(ln) + 1
            if n >= size:
                out.append((page, "\n".join(buf)))
                keep, k = [], 0
                for b in reversed(buf):
                    if k + len(b) > overlap: break
                    keep.insert(0, b); k += len(b)
                buf, n = keep, k
        tail = "\n".join(buf).strip()
        if tail and (not out or out[-1][0] != page or len(tail) > overlap + 20):
            out.append((page, tail))
    return [(p, c) for p, c in out if len(c.strip()) >= 25]

def source_add(title, subject, kind, pages, added_by=0, year=None, rnd=None, append_to=None):
    """يضيف/يستبدل مصدراً نصياً ويفهرسه. append_to=معرّف مصدر موجود → يضيف الصفحات كجزء جديد منه. pages: [(page, text)]"""
    chunks = chunk_pages(pages)
    if not chunks:
        raise RuntimeError("لم أجد نصاً قابلاً للقراءة في الملف.")
    subject = subject if subject in SUBJECTS else "general"
    chars = sum(len(c) for _, c in chunks)
    rows_data = [(" ".join(tokens(c)), c, p) for p, c in chunks]        # الحساب خارج قفل القاعدة
    npages = len({p for p, _ in chunks if p}) or 0
    with db() as con:
        if append_to:
            src = con.execute("SELECT id,subject FROM sources WHERE id=?", (int(append_to),)).fetchone()
            if not src: raise RuntimeError("الكتاب الذي تريد الإضافة إليه غير موجود.")
            sid = src["id"]
            con.executemany("INSERT INTO chunks(norm,raw,source_id,subject,page) VALUES(?,?,?,?,?)", [(n, r, sid, src["subject"], p) for n, r, p in rows_data])
            con.execute("UPDATE sources SET pages=pages+?, chars=chars+?, chunks=chunks+? WHERE id=?", (npages, chars, len(chunks), sid))
            return {"id": sid, "chunks": len(chunks), "pages": npages, "chars": chars, "appended": True}
        for o in con.execute("SELECT id FROM sources WHERE title=? AND subject=? AND kind=?", (title, subject, kind)).fetchall():
            con.execute("DELETE FROM chunks WHERE source_id=?", (o["id"],))
            con.execute("DELETE FROM sources WHERE id=?", (o["id"],))
        cur = con.execute("INSERT INTO sources(title,subject,kind,pages,chars,chunks,added_by,created_at,year,round) VALUES(?,?,?,?,?,?,?,?,?,?)",
                          (title, subject, kind, npages, chars, len(chunks), int(added_by), now_iso(), year, rnd))
        sid = cur.lastrowid
        con.executemany("INSERT INTO chunks(norm,raw,source_id,subject,page) VALUES(?,?,?,?,?)", [(n, r, sid, subject, p) for n, r, p in rows_data])
    return {"id": sid, "chunks": len(chunks), "pages": npages, "chars": chars}

def source_delete(sid):
    with db() as con:
        con.execute("DELETE FROM chunks WHERE source_id=?", (int(sid),))
        r = con.execute("DELETE FROM sources WHERE id=?", (int(sid),)).rowcount
    return r

def sources_list(kinds=("book", "wazari", "qa"), subject=None):
    q = ",".join("?" * len(kinds)); params = list(kinds)
    sql = f"SELECT id,title,subject,kind,pages,chunks,year,round,created_at FROM sources WHERE kind IN ({q})"
    if subject: sql += " AND subject=?"; params.append(subject)
    with db() as con:
        rows = con.execute(sql + " ORDER BY subject, COALESCE(year,0) DESC, id", params).fetchall()
    return [dict(r) for r in rows]

def src_label(h):
    k = h.get("kind")
    if k == "wazari":
        base = "وزاري " + subj_name(subject_key(h.get("subject"))) + (f" {h['year']}" if h.get("year") else "") + (f" — {h['round']}" if h.get("round") else "")
    else:
        t = h.get('title', ''); kl = KIND_LABEL.get(k, 'مصدر')
        base = t if t.startswith(kl) else f"{kl} {t}"
    return base + (f" — ص{h['page']}" if h.get("page") else "")

def src_tag(h):
    """وسم قصير مثل: وزاري 2023 (يظهر على السؤال)."""
    if h.get("kind") == "wazari":
        return "وزاري" + (f" {h['year']}" if h.get("year") else "")
    return ""

_HIT_COLS = "c.raw raw, c.norm norm, c.page page, c.subject subject, c.source_id sid, s.title title, s.kind kind, s.year year, s.round round"

def search_chunks(question, subject="", limit=5, kinds=None, source_id=None):
    qt = list(dict.fromkeys(tokens(question)))[:10]
    if not qt: return []
    fts = " OR ".join(f'"{t}"*' if len(t) >= 3 else f'"{t}"' for t in qt)
    kinds = tuple(kinds) if kinds else None

    def run(sub):
        if not HAS_FTS:
            sql = f"SELECT c.rowid rid, {_HIT_COLS}, 0.0 sc FROM chunks c JOIN sources s ON s.id=c.source_id WHERE (" + " OR ".join(["c.norm LIKE ?"] * len(qt)) + ")"
            params = [f"%{t}%" for t in qt]
            if sub: sql += " AND c.subject=?"; params.append(sub)
            if kinds: sql += f" AND s.kind IN ({','.join('?' * len(kinds))})"; params += list(kinds)
            if source_id: sql += " AND c.source_id=?"; params.append(int(source_id))
            sql += " LIMIT 400"
        else:
            inner = "SELECT rowid rid, raw, norm, page, subject, source_id, bm25(chunks) sc FROM chunks WHERE chunks MATCH ?"
            params = [fts]
            if sub: inner += " AND subject=?"; params.append(sub)
            if kinds: inner += f" AND source_id IN (SELECT id FROM sources WHERE kind IN ({','.join('?' * len(kinds))}))"; params += list(kinds)
            if source_id: inner += " AND source_id=?"; params.append(int(source_id))
            inner += " ORDER BY sc LIMIT 14"
            sql = f"SELECT c.rid rid, c.raw raw, c.norm norm, c.page page, c.subject subject, c.source_id sid, s.title title, s.kind kind, s.year year, s.round round, c.sc sc FROM ({inner}) c JOIN sources s ON s.id=c.source_id ORDER BY c.sc"
        try:
            with db() as con: return [dict(r) for r in con.execute(sql, params)]
        except sqlite3.OperationalError as e:
            log.warning("search query failed: %s", e); return []

    sub = subject if subject in SUBJECTS and subject != "general" else ""
    rows = run(sub)
    if sub and len(rows) < 2 and not source_id:
        seen = {r["rid"] for r in rows}
        rows += [r for r in run("") if r["rid"] not in seen]
    need = 1 if len(qt) <= 2 else (len(qt) + 1) // 2
    good = []
    for r in rows:
        ctoks = set(r["norm"].split())
        cov = sum(1 for q in qt if q in ctoks or any(c.startswith(q) for c in ctoks if len(q) >= 3))
        if cov >= need:
            r["cov"] = cov; good.append(r)
    good.sort(key=lambda r: (-r["cov"], r["sc"]))
    return good[:limit]

def random_chunks(subject="", kinds=None, n=6, source_id=None):
    sql = f"SELECT c.rowid rid, {_HIT_COLS} FROM chunks c JOIN sources s ON s.id=c.source_id WHERE 1=1"
    params = []
    if subject in SUBJECTS and subject != "general": sql += " AND c.subject=?"; params.append(subject)
    if kinds: sql += f" AND s.kind IN ({','.join('?' * len(kinds))})"; params += list(kinds)
    if source_id: sql += " AND c.source_id=?"; params.append(int(source_id))
    sql += " ORDER BY random() LIMIT ?"; params.append(int(n))
    try:
        with db() as con: return [dict(r) for r in con.execute(sql, params)]
    except sqlite3.OperationalError as e:
        log.warning("random chunks failed: %s", e); return []

def fmt_sources(hits, limit=3):
    seen, out = set(), []
    for h in hits:
        label = src_label(h)
        if label in seen: continue
        seen.add(label); out.append(label)
        if len(out) >= limit: break
    return out


# ═══════════════════════════════════════════════════════════════════
#                  الملازم: البحث في قناة المالك (روابط فقط)
# ═══════════════════════════════════════════════════════════════════

_GEN = {_stem(ar_norm(w)) for w in """ملزمه ملازم ملزمة ملخص ملخصات مختصر رابط روابط ارسل ارسلي ادز دز اعطني عطني اريد ابي ابغى وين شنو شلون
رايك رايكم رأيك تنصح تنصحني تنصحوني افضل احسن اجمل مفيده مفيدة تستاهل يستاهل شكو اكو عندكم عندك سادس اعدادي الاعدادي للسادس
مدرس مدرسين يوتيوب يوتيوبر قناه قناة استاذ تقييم قيم زينه حلوه جيده ممتازه شلونها شلونه بالقناه للمدرس""".split()}


def _gen(t):
    """كلمة عامة (ملزمة/رأيك/مادة...) ولو التصق بها حرف جر مثل «بملزمة»."""
    if t in _GEN or t in _SUBJ_WORDS: return True
    return len(t) > 3 and t[0] in "بلوفك" and (t[1:] in _GEN or t[1:] in _SUBJ_WORDS)

def is_summary(row):
    if row.get("mkind") == "summary": return True
    t = (row.get("norm") or "").split()
    return any(w.startswith(("ملخص", "مختصر", "خلاصه")) for w in t)

def search_materials(query, limit=5):
    """يرجع أنسب الملازم فقط (لا يدمج كل الروابط): يطابق الاسم/المدرّس ضمن المادة المذكورة."""
    subj = guess_subject(query)
    terms = [t for t in dict.fromkeys(tokens(query)) if not _gen(t)]
    if not subj and not terms: return []
    with db() as con:
        rows = [dict(r) for r in con.execute("SELECT id,title,subject,telegram_url,description,norm,mkind,teacher,part,batch,mtype FROM materials WHERE active=1 ORDER BY id DESC LIMIT 3000")]
    out = []
    for r in rows:
        if subj and subject_key(r["subject"]) != subj: continue
        toks = set((r["norm"] or "").split())
        if terms:
            hit = sum(1 for t in terms if t in toks or (len(t) >= 3 and any(x.startswith(t) or (len(x) >= 4 and t.startswith(x)) for x in toks)))
            if hit < (len(terms) if len(terms) <= 2 else (len(terms) * 2 + 2) // 3): continue
            r["score"] = hit
        else:
            r["score"] = 0
        out.append(r)
    out.sort(key=lambda r: (-r["score"], -r["id"]))
    top = out[:limit]
    sums = [r for r in out if is_summary(r)]
    if sums and not any(is_summary(r) for r in top):      # حتى لو اكو ملخص، يظهر مع الروابط
        top = top[:limit - 1] + [sums[0]]
    return top

def mat_meta(r):
    bits = ["ملخص" if (r.get("mkind") == "summary" or (not r.get("mkind") and is_summary(r))) else "ملزمة"]
    if r.get("teacher"): bits.append("الأستاذ " + r["teacher"])
    if r.get("part"): bits.append("جزء " + r["part"])
    if r.get("batch"): bits.append("دفعة " + r["batch"])
    if r.get("mtype"): bits.append(r["mtype"])
    return " — ".join(bits)

def mat_url(r):
    u = r.get("telegram_url") or r.get("url") or ""
    if u: return u
    return f"https://t.me/{BOT_USERNAME}?start=mat_{r['id']}" if BOT_USERNAME and r.get("id") else ""

def materials_text(rows, head):
    lines = [head]
    for i, r in enumerate(rows, 1):
        desc = (r.get("description") or "").strip() if r.get("mkind") == "summary" else ""
        link = mat_url(r)
        lines.append(f"{i}. <b>{esc(r['title'])}</b>\n<i>{esc(mat_meta(r))}</i> — {esc(subj_name(subject_key(r['subject'])))}"
                     + (f"\n{esc(desc[:140])}" if desc else "") + (f"\n{esc(link)}" if link else f"\nالملف داخل البوت: /start mat_{r.get('id')}"))
    return "\n\n".join(lines)

def ai_rank_materials(query, cands, limit=5):
    """يفرز الذكاء الاصطناعي المرشّحين ويقدّم الأنسب لطلب الطالب (مع رجوع للترتيب العادي عند الفشل)."""
    items = "\n".join(f"{r['id']}: {mat_meta(r)} | {r['title']}" + (f" | {(r.get('description') or '')[:80]}" if r.get("description") else "") for r in cands[:30])
    ids = []
    try:
        out = llm("أنت أمين مكتبة ذكي لطلبة السادس الإعدادي. اختر الملازم/الملخصات الأنسب لطلب الطالب. أعد JSON فقط: مصفوفة أرقام (id) مرتبة من الأنسب للأقل، بحد أقصى 5، ولا تُدرج ما لا يناسب الطلب.",
                  [], f"طلب الطالب: {query[:200]}\n\nالمتاح:\n{items}", max_tokens=120, temperature=0)
        mm = re.search(r"\[[\d,\s]*\]", out)
        ids = [int(x) for x in json.loads(mm.group(0))] if mm else []
    except Exception as e:
        log.warning("ai rank: %s", e)
    by = {r["id"]: r for r in cands}
    ranked = [by[i] for i in dict.fromkeys(ids) if i in by][:limit]
    return ranked or cands[:limit]


# ═══════════════════════════════════════════════════════════════════
#                           قارئ الملفات
# ═══════════════════════════════════════════════════════════════════

def _fix_reversed(text):
    """بعض ملفات PDF العربية تُستخرج بترتيب معكوس؛ نكتشف ذلك بإحصاء «ال» مقابل «لا» في نهاية الكلمات."""
    fwd = len(re.findall(r"(?<!\w)ال\w", text))
    rev = len(re.findall(r"\wلا(?!\w)", text))
    if rev > 25 and rev > fwd * 2:
        return "\n".join(ln[::-1] for ln in text.splitlines()), True
    return text, False

def read_pdf_pages(path, max_pages=None):
    pages = []
    if fitz is not None:
        with fitz.open(path) as d:
            for i, p in enumerate(d):
                if max_pages and i >= max_pages: break
                try: pages.append((i + 1, p.get_text("text") or ""))
                except Exception: pages.append((i + 1, ""))
    elif PdfReader is not None:
        r = PdfReader(path)
        for i, p in enumerate(r.pages):
            if max_pages and i >= max_pages: break
            try: pages.append((i + 1, p.extract_text() or ""))
            except Exception: pages.append((i + 1, ""))
    else:
        raise RuntimeError("مكتبة قراءة PDF غير مثبتة (pymupdf أو pypdf)")
    return pages

def ocr_pdf_pages(path, progress=None):
    """للـ PDF الممسوح ضوئياً (شائع في الوزاريات): يحوّل كل صفحة لصورة ويقرؤها بالرؤية الذكية/OCR. يرجع (pages, total)."""
    if fitz is None: raise RuntimeError("قراءة الصفحات الممسوحة تحتاج مكتبة pymupdf (pip install pymupdf)")
    if not (GEMINI_KEYS or MISTRAL_KEYS or OCR_KEY): raise RuntimeError("الملف ممسوح ضوئياً ويحتاج GEMINI_API_KEY_1 أو MISTRAL_API_KEY_1 أو OCR_API_KEY لقراءته")
    pages, fails = [], 0
    with fitz.open(path) as d:
        total = len(d); n = min(total, OCR_MAX_PAGES)
        for i in range(n):
            try:
                img = d[i].get_pixmap(dpi=120).tobytes("jpeg")
                txt, _ = read_image(img)
                pages.append((i + 1, txt))
            except Exception as e:
                fails += 1; log.warning("ocr page %d: %s", i + 1, str(e)[:80])
                if fails >= 6 and not pages: raise RuntimeError("فشلت قراءة الصفحات (تحقق من مفاتيح Gemini/Mistral/OCR)")
            if progress and (i + 1) % 3 == 0: progress(i + 1, n)
            time.sleep(0.5)
    return pages, total

def read_docx_text(path):
    if DocxDocument is None: raise RuntimeError("python-docx غير مثبت")
    doc = DocxDocument(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)

def read_txt(path):
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1256"):
        try: return raw.decode(enc)
        except Exception: continue
    return raw.decode("utf-8", errors="replace")

def read_pages(path, suffix, max_pdf_pages=None):
    s = suffix.lower(); notes = []
    if s == ".pdf": pages = read_pdf_pages(path, max_pdf_pages)
    elif s == ".docx": pages = [(0, read_docx_text(path))]
    elif s == ".txt": pages = [(0, read_txt(path))]
    else: raise RuntimeError(f"نوع غير مدعوم: {suffix}")
    pages = [(p, unicodedata.normalize("NFKC", t or "")) for p, t in pages]
    _, flipped = _fix_reversed("\n".join(t for _, t in pages))
    if flipped:
        pages = [(p, "\n".join(ln[::-1] for ln in t.splitlines())) for p, t in pages]
        notes.append("تم إصلاح اتجاه النص العربي المعكوس")
    return pages, notes

def read_doc_text(path, suffix, limit=22000):
    pages, _ = read_pages(path, suffix, max_pdf_pages=60)
    return "\n\n".join(t for _, t in pages)[:limit]

_INGEST_LOCK = threading.Semaphore(1)       # قراءة كتاب واحد في وقت واحد (حماية الذاكرة)

def ingest_bytes(raw, filename, title, subject, kind, added_by, year=None, rnd=None, progress=None, append_to=None, start_page=None):
    """يقرأ ملفاً ويفهرسه. append_to: إضافته كجزء من كتاب موجود. start_page: رقم أول صفحة في هذا الجزء."""
    suffix = Path(filename or "").suffix.lower()
    if suffix not in (".pdf", ".docx", ".txt"): raise RuntimeError("المدعوم: PDF / DOCX / TXT")
    with _INGEST_LOCK:
        tmp = None
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as f:
                f.write(raw); tmp = f.name
            pages, notes = read_pages(tmp, suffix)
            if sum(len(t.strip()) for _, t in pages) < 200:
                if suffix != ".pdf": raise RuntimeError("الملف لا يحتوي نصاً قابلاً للاستخراج.")
                pages, total = ocr_pdf_pages(tmp, progress)
                notes.append("الملف ممسوح ضوئياً وقُرئ بالرؤية الذكية")
                if total > OCR_MAX_PAGES:
                    notes.append(f"قرأت أول {OCR_MAX_PAGES} صفحة من {total}. قسّم الملف وأضف الباقي كجزء ثانٍ من نفس الكتاب")
        finally:
            if tmp and os.path.exists(tmp): os.unlink(tmp)
    sample = "".join(t for _, t in pages)[:50000]
    if len(sample.strip()) < 200: raise RuntimeError("ما قدرت أستخرج نصاً كافياً من الملف.")
    letters = sum(1 for ch in sample if "\u0600" <= ch <= "\u06FF" or (ch.isascii() and ch.isalpha()))
    if letters / max(1, len(sample)) < 0.35: notes.append("جودة النص منخفضة، قد يكون الملف بترميز خاص")
    if append_to:
        with db() as con: mx = con.execute("SELECT MAX(CAST(page AS INTEGER)) m FROM chunks WHERE source_id=?", (int(append_to),)).fetchone()["m"] or 0
        sp = start_page or (mx + 1)
    else:
        sp = start_page or 1
    if sp != 1: pages = [((p + sp - 1) if p else 0, t) for p, t in pages]
    if kind == "wazari" and not append_to:
        first = " ".join(t for _, t in pages[:2])[:1500]
        year = year or detect_year(title, filename, first); rnd = rnd or detect_round(title, filename, first)
        notes.append(f"السنة: {year or 'غير محددة (اكتبها في العنوان)'}" + (f" — {rnd}" if rnd else ""))
    res = source_add(title, subject, kind, pages, added_by, year, rnd, append_to)
    if res.get("appended"): notes.append(f"أضيف كجزء من الكتاب (الصفحات من {sp})")
    res["notes"] = notes; res["year"] = year
    return res


# ═══════════════════════════════════════════════════════════════════
#                       الذكاء الاصطناعي (مع الحماية)
# ═══════════════════════════════════════════════════════════════════

SYSTEM = """أنت «الأستاذ» — مدرّس عراقي خبير لطلبة السادس الإعدادي (المنهج العراقي). تشرح بالعربية المبسّطة وبأسلوب ودود.

أهم قاعدة: الاختصار والبساطة. الطالب وقته ضيق ويريد الفهم بسرعة.
1. ابدأ بالجواب أو الفكرة الأساسية في سطر أو سطرين، ثم الخطوات المختصرة. الحد المعتاد 8 إلى 12 سطراً، ولا تطل إلا إذا طلب الطالب التفصيل.
2. استخدم أسهل طريقة في المنهج العراقي للسادس الإعدادي، ولا تستخدم طرقاً جامعية أو أدوات خارج المنهج. إن كان للمسألة حل بسيط فاعرضه هو فقط.
3. حل المسائل: 3 إلى 6 خطوات مرقّمة قصيرة، كل خطوة سطر أو سطران، ثم الناتج النهائي بوضوح.
4. جمل قصيرة وكلمات سهلة، وأي مصطلح علمي اشرحه بكلمات بسيطة مرة واحدة.
5. بلا مقدمات ولا إعادة للسؤال ولا خاتمة طويلة.
6. حين تُعطى «مقاطع من الكتب» فهي مصدرك الأول: اعتمد عليها ولا تخالفها، وأشر لها برقمها مثل [1].
7. إن كان المقطع من «وزاري» فاذكر سنته ودوره عند الإشارة إليه («ورد في وزاري 2023»). لا تذكر سنة وزاري إلا إذا كانت مكتوبة في عنوان المقطع، ولا تخترع أسئلة وزارية.
8. إن لم تكفِ المقاطع أو لم توجد، قل بصراحة «هذا غير موجود في الكتب التي عندي» ثم أجب من معرفتك العامة ونبّه الطالب لمراجعة كتابه.
9. لا تخترع معلومات ولا أرقام صفحات، وتجنّب الإنكليزية إلا للمصطلحات العلمية.
10. عند شرح درس (لا حل مسألة) اختم بسؤال قصير جداً يختبر الفهم.
11. لا تفترض موضوعاً سابقاً: ردّ على الرسالة الحالية. لا تكمل درساً سابقاً إلا إذا أشارت إليه الرسالة صراحة («كمل»، «ما فهمت»، «ليش»…). إن كانت الرسالة غامضة فعلاً اسأل سؤالاً توضيحياً واحداً قصيراً.
12. إن كان في الرسالة أكثر من طلب فأجب عنها بالترتيب بعناوين قصيرة. وإن لم تعرف قل «ما أعرف» ولا تخترع."""

FMT_TELEGRAM = "\n\nالتنسيق: الإجابة تظهر في تيليجرام؛ استخدم **للعناوين** وقوائم بسيطة، ولا تستخدم جداول ولا LaTeX ولا رموز $ — اكتب المعادلات نصاً واضحاً مثل F = m × a ومثل x² و √x و a/b."
FMT_APP = ("\n\nالتنسيق: الإجابة تظهر في تطبيق يعرض Markdown وLaTeX بشكل مرتّب، لذلك:\n"
           "• اكتب كل معادلة مهمة في سطر مستقل بين $$ و $$ ، والرموز داخل الجملة بين $ و $.\n"
           "• استخدم LaTeX صحيحاً: \\frac{a}{b} للكسور، \\sqrt{x} للجذور، ^ للأسس، _ للأدلة، \\times و \\cdot للضرب، \\rightarrow للأسهم، \\Delta، \\pi، والوحدات بـ \\text{m/s}.\n"
           "• في الكيمياء اكتب المركبات هكذا $\\mathrm{H_2SO_4}$ والتفاعلات $$2\\mathrm{H_2} + \\mathrm{O_2} \\rightarrow 2\\mathrm{H_2O}$$.\n"
           "• يُمنع كتابة المعادلات كنص عادي بحروف وأرقام الكيبورد مثل x^2/3 أو sqrt(x) أو 1/2 mv^2.\n"
           "• خطوات الحل: عنوان قصير لكل خطوة ثم المعادلة في سطرها.\n• اكتب المعادلة الواحدة كاملة في سطر واحد بين $$ و $$ ولا تقسّمها على عدة أسطر، ولا تستخدم \\left و\\right إلا عند الضرورة.\n• لا تستخدم الإيموجي ولا الرموز التعبيرية إطلاقاً.")
FMT_GROUP = "\n\nأنت ترد داخل مجموعة: كن مختصراً (حتى 10 أسطر) وخاطب الطالب باسمه."
SYSTEM_TIPS = "أنت مستشار دراسي للطلبة العراقيين. نصائحك عملية ومشجّعة ومرقّمة." + FMT_TELEGRAM
SYSTEM_STRICT = """أنت مدرّس عراقي. اشرح المحتوى المُرسل فقط.
1. اعتمد حصراً على النص المُرسل ولا تضف من عندك.
2. اشرح كل فكرة بترتيب ظهورها.
3. إذا فيه سؤال: حلّه خطوة بخطوة.
4. في النهاية: خلاصة من ٣ نقاط.""" + FMT_TELEGRAM

class Busy(RuntimeError):
    """الخدمة مشغولة (ضغط) — رسالتها تُعرض للمستخدم."""

DEAD_MODELS = set()
_LLM_SEM = threading.BoundedSemaphore(LLM_CONCURRENCY)
_BRK = {}      # قواطع المزوّدين: اسم → وقت إعادة المحاولة

@contextmanager
def llm_slot(wait=25):
    if not _LLM_SEM.acquire(timeout=wait): raise Busy("الضغط عالٍ حالياً، جرّب بعد لحظات 🙏")
    try: yield
    finally: _LLM_SEM.release()

def _brk_open(n): return _BRK.get(n, 0) > time.time()
def _brk_trip(n, secs): _BRK[n] = time.time() + max(3, min(float(secs), 90))
def _retry_after(r, default=20):
    try: return float(r.headers.get("retry-after") or default)
    except Exception: return default

def _err_msg(r):
    try:
        j = r.json(); e = j.get("error", j)
        return (e.get("message") if isinstance(e, dict) else str(e))[:110]
    except Exception:
        return (r.text or "")[:110]

_GM = {}
_GM_LIST = {}

def _gemini_candidates(key):
    if key in _GM_LIST: return _GM_LIST[key]
    r = requests.get("https://generativelanguage.googleapis.com/v1beta/models?pageSize=200", headers={"x-goog-api-key": key}, timeout=20)
    if r.status_code != 200: return []
    names = [m["name"].split("/")[-1] for m in r.json().get("models", []) if "generateContent" in m.get("supportedGenerationMethods", [])]
    bad = ("image", "tts", "live", "audio", "thinking", "exp", "8b", "robotics", "computer")
    full = sorted([n for n in names if "flash" in n and "lite" not in n and not any(x in n for x in bad) and "preview" not in n], reverse=True)
    lite = sorted([n for n in names if "flash" in n and "lite" in n and not any(x in n for x in bad) and "preview" not in n], reverse=True)
    _GM_LIST[key] = full + lite
    return _GM_LIST[key]

def gemini_call(key, payload, timeout):
    def post(model):
        return requests.post(f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent",
                             headers={"x-goog-api-key": key, "Content-Type": "application/json"}, json=payload, timeout=timeout)
    first = _GM.get(key, GEMINI_MODEL)
    r = post(first)
    if r.status_code in (404, 429):
        tried = {first}
        for m in _gemini_candidates(key)[:5]:
            if m in tried: continue
            tried.add(m)
            r2 = post(m)
            if r2.status_code == 200:
                log.info("Gemini: %s لم ينجح، أستخدم %s", first, m); _GM[key] = m
                return r2
            r = r2 if r2.status_code != 404 else r
    return r

def _gem_text(r):
    parts = r.json().get("candidates", [{}])[0].get("content", {}).get("parts", [])
    return "\n".join(p.get("text", "") for p in parts if p.get("text")).strip()

def _clean_ans(s):
    return re.sub(r"<think>.*?</think>", "", s or "", flags=re.S).strip()

def llm(system, history, user, max_tokens=2200, temperature=0.3):
    """Groq ثم Gemini ضمن ميزانية زمنية وحدّ تزامن، مع قواطع تتجاوز المزوّد المزدحم فوراً."""
    with llm_slot():
        return _llm(system, history, user, max_tokens, temperature)

_RR = itertools.count()

def _oai_chain(name, base, keys, models, msgs, max_tokens, temperature, deadline, errors, state, headers=None):
    """مزوّد متوافق مع OpenAI: يجرّب (موديل × مفتاح) بالتناوب مع قاطع لكل زوج. يرجع نصاً أو None."""
    if not keys: return None
    n = len(keys); start = next(_RR) % n
    for model in models:
        if (name, model) in DEAD_MODELS: continue
        for j in range(n):
            ki = (start + j) % n; bk = f"{name}:{ki}:{model}"
            if _brk_open(bk): state["skipped"] += 1; continue
            left = deadline - time.time()
            if left < 5: return None
            try:
                r = requests.post(f"{base}/chat/completions", headers={"Authorization": f"Bearer {keys[ki]}", **(headers or {})},
                                  json={"model": model, "messages": msgs, "temperature": temperature, "max_tokens": max_tokens}, timeout=min(45, left))
                if r.status_code == 200:
                    ans = _clean_ans((r.json()["choices"][0]["message"].get("content") or ""))
                    if ans:
                        log.info("OK %s/%s (key %d)", name, model, ki + 1); return ans
                    errors.append(f"{name} {model}: رد فارغ"); _brk_trip(bk, 20); continue
                errors.append(f"{name} {model}: HTTP {r.status_code} {_err_msg(r)}")
                body = r.text.lower()
                if r.status_code == 429: _brk_trip(bk, _retry_after(r))
                elif r.status_code >= 500: _brk_trip(bk, 15)
                elif r.status_code in (401, 402, 403): _brk_trip(bk, 900)
                elif r.status_code in (400, 404) and any(x in body for x in ("decommission", "does not exist", "model_not_found", "not found", "no endpoints")):
                    DEAD_MODELS.add((name, model)); break
            except Exception as e:
                _brk_trip(bk, 15); errors.append(f"{name} {model}: {str(e)[:60]}")
    return None

def _llm(system, history, user, max_tokens, temperature):
    """الترتيب: Groq ← Gemini ← Mistral ← OpenRouter (مجاني). تدوير على المفاتيح، قواطع، وميزانية زمنية."""
    errors, deadline, state = [], time.time() + LLM_BUDGET, {"skipped": 0}
    hist = [{"role": h["role"], "content": h["content"][:1500]} for h in history if h["role"] in ("user", "assistant")]
    msgs = [{"role": "system", "content": system}] + hist + [{"role": "user", "content": user}]

    ans = _oai_chain("groq", GROQ_BASE, GROQ_KEYS, GROQ_MODELS, msgs, max_tokens, temperature, deadline, errors, state)
    if ans: return ans

    contents = [{"role": "model" if h["role"] == "assistant" else "user", "parts": [{"text": h["content"]}]} for h in hist]
    contents.append({"role": "user", "parts": [{"text": user}]})
    nk = len(GEMINI_KEYS); rot = next(_RR) % nk if nk else 0
    for i in range(nk):
        idx = (rot + i) % nk + 1; key = GEMINI_KEYS[idx - 1]
        if _brk_open(f"gem{idx}"): state["skipped"] += 1; continue
        left = deadline - time.time()
        if left < 5: break
        try:
            r = gemini_call(key, {"systemInstruction": {"parts": [{"text": system}]}, "contents": contents,
                                  "generationConfig": {"temperature": temperature, "maxOutputTokens": max_tokens}}, min(55, left))
            if r.status_code == 200:
                ans = _clean_ans(_gem_text(r))
                if ans:
                    log.info("OK Gemini key %d", idx); return ans
                errors.append(f"Gemini {idx}: رد فارغ")
            else:
                errors.append(f"Gemini {idx}: HTTP {r.status_code} {_err_msg(r)}")
                if r.status_code == 429: _brk_trip(f"gem{idx}", _retry_after(r))
                elif r.status_code >= 500: _brk_trip(f"gem{idx}", 15)
                elif r.status_code in (400, 401, 403): _brk_trip(f"gem{idx}", 600)
        except Exception as e:
            _brk_trip(f"gem{idx}", 15); errors.append(f"Gemini {idx}: {str(e)[:60]}")

    ans = _oai_chain("mistral", MISTRAL_BASE, MISTRAL_KEYS, MISTRAL_MODELS, msgs, max_tokens, temperature, deadline, errors, state)
    if ans: return ans
    ans = _oai_chain("openrouter", OPENROUTER_BASE, OPENROUTER_KEYS, OPENROUTER_MODELS, msgs, max_tokens, temperature, deadline, errors, state,
                     headers={"X-Title": "Mudarris Bot", "HTTP-Referer": WEB["url"] or "https://t.me"})
    if ans: return ans
    if not errors and state["skipped"]: raise Busy("الخدمة مشغولة لحظياً، جرّب بعد ثوانٍ")
    raise RuntimeError("فشلت خدمات الذكاء الاصطناعي: " + " | ".join(errors[:8]))

def ai_whisper(audio, filename="audio.ogg"):
    if not GROQ_KEYS: raise RuntimeError("تحويل الصوت يحتاج GROQ_API_KEY")
    errors = []
    with llm_slot():
        n = len(GROQ_KEYS); start = next(_RR) % n
        for model in WHISPER_MODELS:
            for j in range(n):
                ki = (start + j) % n; bk = f"whisper:{ki}:{model}"
                if _brk_open(bk): continue
                try:
                    r = requests.post(f"{GROQ_BASE}/audio/transcriptions", headers={"Authorization": f"Bearer {GROQ_KEYS[ki]}"},
                                      files={"file": (filename, audio, "audio/ogg")}, data={"model": model, "language": "ar", "response_format": "text"}, timeout=90)
                    if r.status_code == 200: return r.text.strip()
                    errors.append(f"{model}: HTTP {r.status_code}")
                    if r.status_code == 429: _brk_trip(bk, _retry_after(r))
                    elif r.status_code in (401, 403): _brk_trip(bk, 900)
                except Exception as e:
                    errors.append(f"{model}: {str(e)[:60]}")
    raise RuntimeError("فشل تحويل الصوت: " + " | ".join(errors[:4]))

def gemini_vision(image, prompt, mime="image/jpeg"):
    last = None
    with llm_slot():
        for idx, key in enumerate(GEMINI_KEYS, 1):
            if _brk_open(f"gem{idx}"): continue
            try:
                r = gemini_call(key, {"contents": [{"role": "user", "parts": [{"text": prompt}, {"inlineData": {"mimeType": mime, "data": base64.b64encode(image).decode()}}]}],
                                      "generationConfig": {"temperature": 0.1, "maxOutputTokens": 2500}}, 55)
                if r.status_code == 200:
                    t = _gem_text(r)
                    if t: return t
                elif r.status_code == 429: _brk_trip(f"gem{idx}", _retry_after(r))
                last = f"HTTP {r.status_code}"
            except Exception as e:
                last = str(e)[:60]
    raise RuntimeError(f"Gemini vision: {last or 'مشغول'}")

def ocr_space(image):
    if not OCR_KEY: raise RuntimeError("OCR_API_KEY غير مضبوط")
    for engine in ("2", "3"):
        try:
            r = requests.post(OCR_URL, data={"apikey": OCR_KEY, "language": "auto" if engine == "2" else "ara", "OCREngine": engine,
                                             "scale": "true", "detectOrientation": "true"},
                              files={"file": ("image.jpg", image, "image/jpeg")}, timeout=60)
            d = r.json()
            if r.status_code == 200 and not d.get("IsErroredOnProcessing"):
                t = "\n".join(p.get("ParsedText", "") for p in (d.get("ParsedResults") or [])).strip()
                if len(t) > 15: return re.sub(r"\n{3,}", "\n\n", t)
        except Exception as e:
            log.warning("OCR-%s: %s", engine, e)
    raise RuntimeError("فشل OCR")

VISION_PROMPT = ("استخرج كل النص الظاهر في الصورة حرفياً كما هو (عربي أو إنكليزي). اكتب المعادلات بصيغة LaTeX داخل $...$. "
                 "إن كان فيها رسم أو شكل علمي أو جدول فصِفه باختصار بعد النص تحت عنوان [وصف الشكل]. لا تضف أي شرح آخر.")

def mistral_vision(image, prompt, mime="image/jpeg"):
    """قراءة الصور عبر Mistral (احتياط بعد Gemini)."""
    if not MISTRAL_KEYS: raise RuntimeError("لا مفاتيح Mistral")
    b64 = base64.b64encode(image).decode(); last = None
    with llm_slot():
        n = len(MISTRAL_KEYS); start = next(_RR) % n
        for j in range(n):
            ki = (start + j) % n; bk = f"mistralv:{ki}"
            if _brk_open(bk): continue
            try:
                r = requests.post(f"{MISTRAL_BASE}/chat/completions", headers={"Authorization": f"Bearer {MISTRAL_KEYS[ki]}"},
                                  json={"model": MISTRAL_VISION_MODEL, "temperature": 0.1, "max_tokens": 2500,
                                        "messages": [{"role": "user", "content": [{"type": "text", "text": prompt}, {"type": "image_url", "image_url": f"data:{mime};base64,{b64}"}]}]}, timeout=60)
                if r.status_code == 200:
                    t = _clean_ans(r.json()["choices"][0]["message"].get("content") or "")
                    if t: return t
                elif r.status_code == 429: _brk_trip(bk, _retry_after(r))
                elif r.status_code in (401, 403): _brk_trip(bk, 900)
                last = f"HTTP {r.status_code}"
            except Exception as e:
                last = str(e)[:60]
    raise RuntimeError(f"Mistral vision: {last or 'مشغول'}")

def read_image(image, mime="image/jpeg"):
    """يقرأ نص الصورة: Gemini ثم Mistral ثم OCR.space. يرجع (text, engine)."""
    if GEMINI_KEYS:
        try: return gemini_vision(image, VISION_PROMPT, mime), "رؤية ذكية"
        except Exception as e: log.warning("gemini vision: %s", e)
    if MISTRAL_KEYS:
        try: return mistral_vision(image, VISION_PROMPT, mime), "رؤية ذكية"
        except Exception as e: log.warning("mistral vision: %s", e)
    if OCR_KEY:
        return ocr_space(image), "OCR"
    raise RuntimeError("لا توجد خدمة لقراءة الصور (أضف GEMINI_API_KEY_1 أو MISTRAL_API_KEY_1 أو OCR_API_KEY)")


# ── رأي صريح من السوشيال ميديا (ملازم / مدرّسون) ──
SYSTEM_OPINION = """أنت مستشار دراسي صريح لطلبة السادس الإعدادي في العراق. تعطي رأياً صريحاً ومختصراً عن ملزمة أو مدرّس اعتماداً على ما يقوله الطلبة والمدرّسون عنه في مواقع التواصل (تيليجرام، فيسبوك، يوتيوب، إنستغرام، تيك توك) وعلى نتائج البحث.
قواعد صارمة:
- لا تخترع آراء ولا إحصاءات ولا أسماء ولا روابط. إن لم تجد معلومات كافية قل ذلك بوضوح واعطِ تقييماً مبدئياً مبنياً على المعروف فقط.
- السطر الأول: الحكم (ممتاز / جيد / متوسط / لا أنصح) مع سبب في جملة.
- ثم ثلاث نقاط قصيرة: القوة، الضعف، لمن يناسب.
- لا تتجاوز 12 سطراً، بلا روابط، وبلهجة بيضاء مفهومة."""

def opinion_get(key):
    with db() as con:
        r = con.execute("SELECT text,created_at,ttl_h FROM opinions WHERE key=?", (key,)).fetchone()
    if r and datetime.utcnow() - datetime.fromisoformat(r["created_at"]) < timedelta(hours=r["ttl_h"]): return r["text"]
    return None

def opinion_put(key, text, ttl_h):
    with db() as con:
        con.execute("INSERT INTO opinions(key,text,created_at,ttl_h) VALUES(?,?,?,?) ON CONFLICT(key) DO UPDATE SET text=excluded.text, created_at=excluded.created_at, ttl_h=excluded.ttl_h",
                    (key[:200], text, now_iso(), ttl_h))

def gemini_grounded(system, prompt, max_tokens=1400):
    """Gemini مع بحث Google الحيّ (يجلب ما يُقال فعلاً على الإنترنت). يرجع None إن لم يتوفر."""
    with llm_slot():
        for idx, key in enumerate(GEMINI_KEYS, 1):
            if _brk_open(f"gem{idx}"): continue
            try:
                r = gemini_call(key, {"systemInstruction": {"parts": [{"text": system}]}, "contents": [{"role": "user", "parts": [{"text": prompt}]}],
                                      "tools": [{"google_search": {}}], "generationConfig": {"temperature": 0.3, "maxOutputTokens": max_tokens}}, 55)
                if r.status_code == 200:
                    t = _clean_ans(_gem_text(r))
                    if len(t) > 40: return t
                elif r.status_code == 429: _brk_trip(f"gem{idx}", _retry_after(r))
            except Exception as e:
                log.warning("grounded: %s", str(e)[:80])
    return None

def web_opinion(key, subject, target, extra=""):
    """رأي صريح عن target. يُخزَّن (7 أيام للمبني على بحث حيّ، 6 ساعات للاحتياطي) لتخفيف الحمل."""
    cached = opinion_get(key)
    if cached: return cached
    ctx = f"المادة: {subj_name(subject)}\nالموضوع: {target}\n" + (f"معلومات من إدارة القناة: {extra}\n" if extra else "")
    ans = gemini_grounded(SYSTEM_OPINION, ctx + "ابحث عن آراء الطلبة والمدرّسين عنه في السوشيال ميديا وأعطني رأيك الصريح.") if GEMINI_KEYS else None
    if ans:
        opinion_put(key, ans, 168); return ans
    ans = llm(SYSTEM_OPINION, [], ctx + "أعطني رأيك الصريح بحسب ما تعرفه فقط. لا بحث متاح الآن.", max_tokens=900)
    ans += "\n\n⚠️ لم أستطع البحث المباشر الآن، فهذا تقييم مبدئي مما أعرفه وقد لا يكون كاملاً."
    opinion_put(key, ans, 6)
    return ans


# ═══════════════════════════════════════════════════════════════════
#                       المدرّس: الإجابة والاختبارات
# ═══════════════════════════════════════════════════════════════════

def build_system(uid, subject, channel, name="", group=False):
    u = db_get_user(uid)
    sys_ = SYSTEM + (FMT_APP if channel == "app" else FMT_TELEGRAM)
    if group: sys_ += FMT_GROUP
    ctx = []
    if name: ctx.append(f"اسم الطالب: {name}")
    if u.get("grade"): ctx.append(f"الصف: {u['grade']}")
    if subject in SUBJECT_PROMPTS: ctx.append(SUBJECT_PROMPTS[subject])
    elif u.get("preferred_subject") in SUBJECT_PROMPTS: ctx.append("المادة المفضلة للطالب: " + subj_name(u["preferred_subject"]))
    perf = db_student_summary(uid)
    if perf: ctx.append(f"أداؤه في اختباراته الأخيرة: {perf} (ركّز على نقاط ضعفه بلطف)")
    if ctx: sys_ += "\n\n[عن هذا الطالب تحديداً]\n" + "\n".join(ctx)
    return sys_

def refs_block(hits):
    if not hits: return ""
    lines = ["[مقاطع من الكتب]"]
    for i, h in enumerate(hits, 1):
        lines.append(f"[{i}] {src_label(h)}\n{h['raw'][:1100]}")
    lines.append("[نهاية المقاطع]\n")
    return "\n".join(lines)

# ── موجّه النيّة: دراسة / دردشة / متابعة قصيرة (بلا استدعاء نموذج إضافي) ──
SYSTEM_CHAT = """أنت «الأستاذ» — مساعد دراسي عراقي ودود لطلبة السادس الإعدادي، طوّرك صاحب البوت (المالك).
ردّ على رسالة الطالب الحالية فقط، بلهجة عراقية بسيطة وقصيرة (سطر إلى 3 أسطر). لا تكمل أي درس أو موضوع سابق، ولا تذكر وزاريات أو كتباً إلا إذا طلب.
إذا قال إنه هو من طوّرك أو صنعك فرحّب به وشكره وقل إنك جاهز لأي تعديل أو سؤال. إذا شكرك أو مدحك فاشكره بخفة. إذا كان متعباً أو متضايقاً فشجّعه بجملة صادقة واقترح استراحة قصيرة ثم الرجوع للدراسة.
إن لم تفهم قصده فاسأله سؤالاً واحداً قصيراً. لا تختلق معلومات عن نفسك."""
def _wre(words, suffix=True):
    w = "|".join(sorted({ar_norm(x) for x in words}, key=len, reverse=True))
    return re.compile(r"(?<![ء-ي])(?:و|ف|ب|ل|ال)?(?:" + w + r")" + (r"(?:ني|لي|لنا|ها|هم|ه|ك|نا)?" if suffix else "") + r"(?![ء-ي])")
_STUDY_CUES = _wre(["اشرح", "شرح", "حل", "حلل", "اعراب", "اعرب", "معادلة", "قانون", "درس", "فصل", "وزاري", "ملزمة", "سؤال", "اسئلة", "تمرين", "مسألة", "مثال", "اختبار", "كويز",
                    "لخص", "ملخص", "معنى", "عرف", "الفرق", "اثبت", "احسب", "ترجم", "قاعدة", "نحو", "فيزياء", "كيمياء", "احياء", "رياضيات", "عربي", "انكليزي", "انجليزي",
                    "اسلامية", "تاريخ", "جغرافية", "ادب", "كتاب", "ملزمه", "مادة", "امتحان", "نتيجة", "درجة", "جدول", "خطة"])
_CHAT_RE = _wre(["طورتك", "سويتك", "صممتك", "برمجتك", "صنعتك", "مطورك", "مبرمجك", "صانعك", "سواك", "صنعك", "طورك", "منو انت", "من انت", "شنو انت", "شسمك", "اسمك", "هلا", "هلو", "هاي",
                 "مرحبا", "اهلا", "السلام", "صباح", "مساء", "شكرا", "تسلم", "يسلمو", "عاش", "يعطيك", "احبك", "ممتاز", "حلو", "رائع", "مبدع", "جيد", "تمام", "شلونك", "شخبارك",
                 "اخبارك", "باي", "وداعا", "تصبح", "غبي", "ما تفهم", "نكتة", "نكته", "اضحك", "تعبان", "زهقان", "ملل", "متضايق", "حزين", "خايف", "زين", "كفو", "بطل", "حبيبي", "تدلل"], suffix=False)
_FOLLOW_RE = re.compile(r"^(?:" + "|".join(ar_norm(x) for x in ["كمل", "كملها", "زيد", "اكثر", "وبعد", "بعد", "ليش", "لماذا", "شلون", "كيف", "ما فهمت", "مافهمت", "اعد", "اعدها", "عيد", "مثال", "بسط", "بسطها", "وضح", "وضحها", "اختصر", "اختصرها", "لا مو هذا", "مو هذا", "هذا قصدي", "مو هيج", "غلط"]) + r")(?:\s|$)")
def classify_intent(text):
    low = re.sub(r"[^\w\s]", " ", ar_norm(text or "")).strip(); n = len(low.split())
    if not low or n > 10 or _STUDY_CUES.search(low): return "study"
    if n <= 5 and _FOLLOW_RE.match(low): return "follow"
    if _CHAT_RE.search(low): return "chat"
    return "study"

def chat_reply(uid, text, channel="telegram", name="", chat_id=0, save=True):
    system = SYSTEM_CHAT + (f"\nاسم الطالب: {name}" if name else "")
    ans = llm(system, db_history(uid, chat_id, 4), text, max_tokens=320, temperature=0.7)
    if save:
        src = "miniapp" if channel == "app" else "telegram"
        db_save_message(uid, "user", text, src, chat_id); db_save_message(uid, "assistant", ans, src, chat_id)
    return ans

def ask_teacher(uid, question, subject="", chat_id=0, channel="telegram", name="", group=False, save=True, max_tokens=1400, shown=None, from_image=False):
    """المسار الرئيسي: بحث في الكتب والوزاريات → ذاكرة هذا الطالب فقط → النموذج. يرجع (answer, hits)."""
    hits = []
    intent = "study" if from_image else classify_intent(question)
    if intent == "chat": return chat_reply(uid, question, channel, name, chat_id, save), []
    squery = question
    if intent == "follow":
        prev = [h for h in db_history(uid, chat_id, 6) if h.get("role") == "user"]
        if prev: squery = (prev[-1].get("content") or "")[:300] + " " + question
    try:
        hits = search_chunks(squery[:400], subject, 5)
        if not any(h["kind"] == "wazari" for h in hits):
            have = {h["raw"] for h in hits}
            hits += [h for h in search_chunks(squery[:400], subject, 2, kinds=("wazari",)) if h["raw"] not in have]
    except Exception as e:
        log.warning("search failed: %s", e)
    system = build_system(uid, subject, channel, name, group)
    history = db_history(uid, chat_id, MAX_HISTORY)
    label = "سؤال الطالب (نص مقروء من صورة أرسلها، حلّه أو اشرحه): " if from_image else "سؤال الطالب: "
    user_msg = refs_block(hits) + label + question + ("\n(رسالة متابعة قصيرة: يقصد بها آخر موضوع في المحادثة)" if intent == "follow" else "")
    ans = llm(system, history, user_msg, max_tokens=max_tokens)
    if channel == "app": ans = normalize_math(ans)
    if save:
        src = "miniapp" if channel == "app" else "telegram"
        db_save_message(uid, "user", shown or question, src, chat_id)
        db_save_message(uid, "assistant", ans, src, chat_id)
    touch_activity(uid, 1)
    return ans, hits

def parse_quiz_json(text):
    t = re.sub(r"```(?:json)?", "", text or "").strip()
    dec, data = json.JSONDecoder(), None
    for i, m in enumerate(re.finditer(r"[\{\[]", t)):
        if i >= 6: break
        try: data, _ = dec.raw_decode(t, m.start()); break
        except Exception: continue
    if data is None: return []
    qs = data.get("questions") if isinstance(data, dict) else data
    out = []
    for q in qs or []:
        try:
            opts = []
            for o in q["options"]:
                o = str(o).strip()[:200]
                if o and o not in opts: opts.append(o)
            ans = int(q["answer"])
            if 2 <= len(opts) <= 5 and 0 <= ans < len(q["options"]) and str(q["q"]).strip():
                correct = str(q["options"][ans]).strip()[:200]
                if correct not in opts: continue
                out.append({"q": str(q["q"]).strip()[:700], "options": opts, "answer": opts.index(correct),
                            "explain": str(q.get("explain", "")).strip()[:500], "ref": int(q.get("ref", 0) or 0)})
        except Exception: continue
    return out

class NoSource(RuntimeError):
    """لا توجد وزاريات للمادة."""

def make_quiz(subject, topic, count=5, qtype="mixed", wazari=False, source_id=None, topics=None):
    """يرجع (questions, info). في وضع الوزاري كل سؤال مربوط بمقطع حقيقي (سنة الوزاري من القاعدة لا من النموذج)."""
    count = max(1, min(10, int(count)))
    sub = subject if subject in SUBJECTS and subject != "general" else ""
    name = subj_name(sub) if sub else "أي مادة"
    hits, kind = [], "general"
    if source_id:
        hits = random_chunks("", None, 8, source_id=source_id); kind = "wazari"
        if not hits: raise NoSource("هذا الوزاري غير موجود.")
    elif wazari:
        hits = search_chunks(topic, sub, 8, kinds=("wazari",)) if topic else []
        hits = hits or random_chunks(sub, ("wazari",), 8)
        if not hits: raise NoSource("ما عندي وزاريات لهذه المادة بعد.")
        kind = "wazari"
    elif topic or sub:
        if topics and len(topics) > 1:                      # عدة مواضيع: نبحث في كل موضوع ونجمع
            seen, hits = set(), []
            for tp in topics[:6]:
                for h in search_chunks(tp, sub, 3):
                    if h["raw"] not in seen: seen.add(h["raw"]); hits.append(h)
            hits = hits[:9]
        else:
            hits = search_chunks(topic or name, sub, 6)
        if not hits and sub and not topic: hits = random_chunks(sub, ("book",), 6)
        kind = "book" if hits else "general"
    tkind = {"mcq": "اختيار من متعدد (4 خيارات)", "tf": "صح أو خطأ (الخياران: صح، خطأ)"}.get(qtype, "مزيج من الاختيار المتعدد وصح/خطأ")
    rules = ("حوّل الأسئلة الوزارية الواردة في المقاطع أعلاه فقط إلى أسئلة اختبار، ولا تخترع سؤالاً غير موجود فيها، وكل سؤال يحمل ref صحيحاً (رقم المقطع)." if kind == "wazari"
             else ("اعتمد على المقاطع أعلاه، و ref = رقم المقطع الذي بُني عليه السؤال أو 0." if hits else "التزم بمنهج السادس الإعدادي العراقي، و ref = 0."))
    prompt = (refs_block(hits) + f"أنشئ اختباراً من {count} أسئلة في مادة {name}" + (f" عن الموضوع: {topic}" if topic else " عن منهج المادة عموماً") + f".\nالنوع: {tkind}.\n"
              'أعد JSON فقط وبلا أي نص آخر بهذا الشكل:\n{"questions":[{"q":"نص السؤال","options":["...","...","...","..."],"answer":0,"explain":"شرح مختصر للجواب","ref":1}]}\n'
              "اجعل نص كل سؤال قصيراً (سطر أو سطران) بمستوى طالب السادس الإعدادي، والخيارات قصيرة، والشرح جملة أو جملتان. " + ("وزّع الأسئلة على المواضيع بالتساوي. " if topics and len(topics) > 1 else "") +
              "answer = رقم الخيار الصحيح ويبدأ من 0. لا تكرر الأسئلة ونوّع موضع الجواب الصحيح. اكتب أي معادلة بصيغة LaTeX داخل $...$. "
              "ممنوع منعاً باتاً الأسئلة السطحية: لا تسأل عن اسم المؤلف أو الناشر أو الطبعة أو الغلاف أو المقدمة أو الفهرس أو أرقام الصفحات أو السنوات والأسماء الجانبية. اسأل فقط عن مضمون الدرس: مفهوم، قانون، تعليل، مقارنة، تطبيق أو حل مسألة، بمستوى يقيس الفهم لا الحفظ الأعمى. " + rules)
    system = "أنت مدرّس عراقي تُعدّ أسئلة اختبار دقيقة باللغة العربية. أعد JSON صالحاً فقط."
    qs = []
    for _ in range(2):
        qs = parse_quiz_json(llm(system, [], prompt, max_tokens=3000, temperature=0.4))
        if qs: break
    out = []
    _junk = ("المؤلف", "مؤلف الكتاب", "الناشر", "الطبعة", "الغلاف", "الفهرس", "رقم الصفحة", "دار النشر", "حقوق الطبع")
    qs = [q for q in qs if not any(j in str(q.get("q", "")) for j in _junk)]
    for q in qs:
        ref = q.pop("ref", 0)
        h = hits[ref - 1] if 1 <= ref <= len(hits) else None
        if kind == "wazari" and not h: continue
        if h: q["src"] = src_label(h); q["tag"] = src_tag(h)
        if len(q["options"]) > 2 and not {"صح", "خطأ"} >= set(q["options"]):
            correct = q["options"][q["answer"]]; random.shuffle(q["options"]); q["answer"] = q["options"].index(correct)
        out.append(q)
    return out[:count], {"kind": kind, "grounded": bool(hits)}


# ═══════════════════════════════════════════════════════════════════
#                  حالة المستخدم، الحدود، وطوابير العمل
# ═══════════════════════════════════════════════════════════════════

USER_LAST, USER_MODE, USER_SUBJECT, USER_QTYPE, OWNER_STATE = {}, {}, {}, {}, {}
STATS = {"started": datetime.now(), "requests": 0, "photos": 0, "docs": 0, "voices": 0, "quizzes": 0, "errors": 0, "webapp_visits": 0}

def is_owner(uid):
    return uid in OWNER_IDS

def check_rate(uid):
    if is_owner(uid):
        USER_LAST[uid] = time.time(); return True, 0
    now = time.time(); last = USER_LAST.get(uid, 0)
    if now - last < RATE_LIMIT: return False, int(RATE_LIMIT - (now - last)) + 1
    USER_LAST[uid] = now
    return True, 0

_LIM = {}

def daily_limit_for(uid):
    """80 للجميع، ويزيد 1 لكل 4 نقاط (استخدام + مساهمات + تركيز) حتى 200. المالك بلا حد، والمشرف 200."""
    if is_owner(uid): return 10**6
    if is_admin(uid): return DAILY_MAX
    c = _LIM.get(uid)
    if c and time.time() - c[1] < 60: return c[0]
    try: pts = activity_totals(uid)["total_points"]
    except Exception: pts = 0
    lim = min(DAILY_MAX, DAILY_LIMIT + pts // 4); _LIM[uid] = (lim, time.time())
    return lim

def check_daily(uid):
    limit = daily_limit_for(uid)
    day = today_key()
    with db() as con:
        row = con.execute("SELECT requests FROM daily_usage WHERE user_id=? AND day=?", (int(uid), day)).fetchone()
        cnt = int(row["requests"]) if row else 0
        if cnt >= limit: return False, cnt
        con.execute("INSERT INTO daily_usage(user_id,day,requests) VALUES(?,?,1) ON CONFLICT(user_id,day) DO UPDATE SET requests=requests+1", (int(uid), day))
    return True, cnt + 1

def refund_daily(uid):
    try:
        with db() as con:
            con.execute("UPDATE daily_usage SET requests=MAX(0,requests-1) WHERE user_id=? AND day=?", (int(uid), today_key()))
    except Exception: pass

# ── طابور الأعمال الثقيلة: لا يجمّد الاستقبال مهما كثر الطلب ──
JOBS = ThreadPoolExecutor(max_workers=JOB_WORKERS, thread_name_prefix="job")
_JL = threading.Lock()
INFLIGHT = set()
_SEQ = itertools.count()

def claim(uid, tag="chat"):
    """(0,key) تم الحجز | (1,None) المستخدم عنده طلب جارٍ | (2,None) النظام ممتلئ."""
    with _JL:
        if len(INFLIGHT) >= JOB_QUEUE_MAX: return 2, None
        key = (uid, tag, next(_SEQ)) if is_owner(uid) else (uid, tag)
        if key in INFLIGHT: return 1, None
        INFLIGHT.add(key); return 0, key

def release(key):
    with _JL: INFLIGHT.discard(key)

@contextmanager
def slot(uid, tag="chat"):
    st, key = claim(uid, tag)
    try: yield st
    finally:
        if st == 0: release(key)

def submit(key, fn):
    def w():
        try: fn()
        except Exception: log.exception("job failed")
        finally: release(key)
    try: JOBS.submit(w)
    except Exception:
        release(key); raise

_CHAT_HITS = defaultdict(deque)
def chat_flooded(cid, limit=30, per=60):
    now = time.time(); dq = _CHAT_HITS[cid]
    while dq and now - dq[0] > per: dq.popleft()
    if len(dq) >= limit: return True
    dq.append(now); return False

# ═══════════════════════════════════════════════════════════════════
#                         أدوات إرسال تيليجرام
# ═══════════════════════════════════════════════════════════════════

bot = telebot.TeleBot(BOT_TOKEN or "0:placeholder", parse_mode=None, threaded=True, num_threads=8)

def esc(t):
    return html.escape(str(t if t is not None else ""), quote=False)

def _plain(t):
    return html.unescape(re.sub(r"</?[a-zA-Z][^>]*>", "", t))

def _flood(e):
    mm = re.search(r"retry after (\d+)", str(e), re.I)
    return min(int(mm.group(1)), 6) if mm else 0

def _tg(fn_html, fn_plain):
    try: return fn_html()
    except Exception as e:
        w = _flood(e)
        if w:
            time.sleep(w)
            try: return fn_html()
            except Exception: pass
        log.info("HTML send failed (%s) → plain", str(e)[:80])
        return fn_plain()

def say(chat_id, text, **kw):
    return _tg(lambda: bot.send_message(chat_id, text, parse_mode="HTML", **kw), lambda: bot.send_message(chat_id, _plain(text), **kw))

def reply(m, text, **kw):
    return _tg(lambda: bot.reply_to(m, text, parse_mode="HTML", **kw), lambda: bot.reply_to(m, _plain(text), **kw))

def edit_msg(chat_id, msg_id, text, kb=None):
    try:
        bot.edit_message_text(text, chat_id, msg_id, parse_mode="HTML", reply_markup=kb)
    except Exception as e:
        if "not modified" in str(e): return
        try: bot.edit_message_text(_plain(text), chat_id, msg_id, reply_markup=kb)
        except Exception as e2:
            if "not modified" not in str(e2): log.warning("edit: %s", e2)

def del_msg(msg):
    try: bot.delete_message(msg.chat.id, msg.message_id)
    except Exception: pass

def err_text(e, uid=0):
    """رسالة خطأ للطالب بلا تفاصيل داخلية؛ المالك يرى التفاصيل."""
    if isinstance(e, (Busy, NoSource)): return str(e)
    if is_owner(uid): return "❌ " + esc(str(e)[:250])
    return "❌ تعذّر الرد الآن، جرّب بعد قليل."

def md_to_html(text):
    t = esc(text)
    t = re.sub(r"```(?:\w+)?\n?(.*?)```", lambda m: "<pre>" + m.group(1).strip("\n") + "</pre>", t, flags=re.S)
    t = re.sub(r"^\s{0,3}#{1,6}\s*(.+?)\s*$", r"<b>\1</b>", t, flags=re.M)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t, flags=re.S)
    t = re.sub(r"(?<![\w*])\*(?!\s)([^*\n]+?)(?<!\s)\*(?![\w*])", r"<i>\1</i>", t)
    t = re.sub(r"`([^`\n]+)`", r"<code>\1</code>", t)
    t = re.sub(r"^\s*[-*]\s+", "• ", t, flags=re.M)
    return t

def tg_text(text):
    """تيليجرام لا يعرض LaTeX: نحوّله لرموز مقروءة (x² ، √ ، ÷ ...)."""
    if "$" in text or "\\" in text: return latex_to_unicode(normalize_math(text))
    return text

def split_long(text, limit=MAX_LEN):
    if len(text) <= limit: return [text]
    parts, cur = [], ""
    for para in text.split("\n\n"):
        if len(cur) + len(para) + 2 <= limit: cur += ("\n\n" if cur else "") + para
        else:
            if cur: parts.append(cur)
            cur = para
    if cur: parts.append(cur)
    final = []
    for p in parts:
        if len(p) <= limit: final.append(p)
        else: final += [p[i:i + limit] for i in range(0, len(p), limit)]
    return final or [text[:limit]]

def send_answer(chat_id, text, m=None, kb=None, hits=None):
    text = tg_text(text)
    parts = split_long(text)
    srcs = fmt_sources(hits or [])
    foot = ("\n\n📖 <i>من المصادر: " + " • ".join(esc(s) for s in srcs) + "</i>") if srcs else ""
    for i, part in enumerate(parts):
        last = i == len(parts) - 1
        body = md_to_html(part) + (foot if last else "")
        k = kb if last else None
        plain = part + ((("\n\n📖 من المصادر: " + " • ".join(srcs)) if srcs else "") if last else "")
        first = m is not None and i == 0
        _tg(lambda: (bot.reply_to(m, body, parse_mode="HTML", reply_markup=k) if first else bot.send_message(chat_id, body, parse_mode="HTML", reply_markup=k)),
            lambda: (bot.reply_to(m, plain, reply_markup=k) if first else bot.send_message(chat_id, plain, reply_markup=k)))

def download(file_id, limit=MAX_FILE):
    info = bot.get_file(file_id)
    if info.file_size and info.file_size > limit:
        raise RuntimeError(f"الملف أكبر من {limit // (1024 * 1024)}MB")
    return bot.download_file(info.file_path)

def is_group(m):
    return m.chat.type in ("group", "supergroup")

def display_name(user):
    return (getattr(user, "first_name", "") or "").strip() or "طالب"

def stale(m):
    return time.time() - (getattr(m, "date", None) or time.time()) > STALE_SECS

def safe(fn):
    @functools.wraps(fn)
    def w(*a, **k):
        try: return fn(*a, **k)
        except Exception: log.exception("handler %s", fn.__name__)
    return w

def gate(m, group, count=True):
    """حظر + معدّل + حدّ يومي لكل طالب على حدة. في المجموعات يتجاهل بصمت لتفادي الإزعاج."""
    uid = m.from_user.id
    if db_is_banned(uid):
        if not group: reply(m, "🚫 تم تقييد وصولك إلى البوت.")
        return False
    ok, wait = check_rate(uid)
    if not ok:
        if not group: reply(m, f"⏳ انتظر {wait} ثانية يا {esc(display_name(m.from_user))}.")
        return False
    if count:
        ok, _ = check_daily(uid)
        if not ok:
            if not group: reply(m, f"⚠️ وصلت حدّك اليومي ({daily_limit_for(uid)} طلب). ساهم بملزمة أو ملخص واستخدم البوت أكثر ليزيد حدّك حتى {DAILY_MAX}. نكمل باجر.")
            return False
    return True

def run_heavy(m, fn, group=False, count=True, tag="chat"):
    """يمرّ بالبوّابة ثم يحجز مكاناً في الطابور؛ الاستقبال لا يتجمّد أبداً."""
    uid = m.from_user.id
    if group and chat_flooded(m.chat.id): return False
    if not gate(m, group, count): return False
    st, key = claim(uid, tag)
    if st:
        if count: refund_daily(uid)
        if not group: reply(m, "⏳ لسا أجاوب على سؤالك السابق، انتظر لحظة." if st == 1 else "🚦 الضغط عالٍ حالياً، حاول بعد لحظات.")
        return False
    submit(key, fn)
    return True

def run_owner_job(chat_id, uid, fn):
    st, key = claim(uid, "ingest")
    if st: say(chat_id, "⏳ النظام مشغول، أعد المحاولة بعد قليل."); return False
    submit(key, fn); return True


# ═══════════════════════════════════════════════════════════════════
#                           لوحات المفاتيح
# ═══════════════════════════════════════════════════════════════════

def webapp_btn(text="🎨 افتح التطبيق"):
    if not WEB["url"]: return None
    try: return types.InlineKeyboardButton(text, web_app=types.WebAppInfo(url=WEB["url"]))
    except Exception as e:
        log.warning("webapp_btn: %s", e); return None

def kb_home():
    return types.InlineKeyboardMarkup([[types.InlineKeyboardButton("🏠 الرئيسية", callback_data="menu:home")]])

MENU_DEFAULTS = [   # (الاسم, النوع, القيمة, صف كامل)
    ("🎨 افتح التطبيق المصغّر", "webapp", "", 1),
    ("📚 اسأل الأستاذ", "cb", "menu:ask", 0), ("📝 اختبرني", "cb", "menu:quiz", 0),
    ("🏛 أسئلة وزارية", "cb", "menu:wazari", 0), ("🎯 اختر مادة", "cb", "menu:subjects", 0),
    ("📄 ملزمة من القناة", "cb", "menu:handout", 0), ("🎬 مدرّسو يوتيوب", "cb", "menu:teacher", 0),
    ("📸 ملف / صورة", "cb", "menu:files", 0), ("🎤 سؤال صوتي", "cb", "menu:voice", 0),
    ("📤 أضف ملزمة / ملخص", "cb", "menu:contrib", 0), ("💡 نصائح الدراسة", "cb", "menu:tips", 0),
    ("📊 إحصائياتي", "cb", "menu:stats", 0), ("ℹ️ مساعدة", "cb", "menu:help", 0)]
_MENU = {"rows": None}
_STYLE_EMO = {"primary": "🔵", "success": "🟢", "danger": "🔴"}

def menu_load(force=False):
    """يرجع أزرار القائمة الرئيسية من القاعدة (وينشئها بالقيم الافتراضية أول مرة)."""
    if _MENU["rows"] is not None and not force: return _MENU["rows"]
    with db() as con:
        con.execute("CREATE TABLE IF NOT EXISTS menu_btns(id INTEGER PRIMARY KEY AUTOINCREMENT, label TEXT NOT NULL, kind TEXT NOT NULL DEFAULT 'cb', value TEXT NOT NULL DEFAULT '', style TEXT NOT NULL DEFAULT '', pos INTEGER NOT NULL DEFAULT 0, active INTEGER NOT NULL DEFAULT 1, full INTEGER NOT NULL DEFAULT 0, clicks INTEGER NOT NULL DEFAULT 0)")
        if not con.execute("SELECT 1 FROM menu_btns LIMIT 1").fetchone():
            for i, (lb, kd, vl, fl) in enumerate(MENU_DEFAULTS):
                con.execute("INSERT INTO menu_btns(label,kind,value,pos,full) VALUES(?,?,?,?,?)", (lb, kd, vl, i, fl))
        rows = [dict(r) for r in con.execute("SELECT * FROM menu_btns ORDER BY pos,id")]
    _MENU["rows"] = rows
    return rows

def _mkbtn(label, style, **kw):
    B = types.InlineKeyboardButton
    if style in _STYLE_EMO:
        try: return B(label, style=style, **kw)
        except TypeError: label = _STYLE_EMO[style] + " " + label
    return B(label, **kw)

def _menu_button(r):
    k = r["kind"]
    try:
        if k == "webapp":
            if not WEB["url"]: return None
            return _mkbtn(r["label"], r["style"], web_app=types.WebAppInfo(url=WEB["url"]))
        if k == "url": return _mkbtn(r["label"], r["style"], url=r["value"])
        if k == "text": return _mkbtn(r["label"], r["style"], callback_data=f"cbtn:{r['id']}")
        return _mkbtn(r["label"], r["style"], callback_data=r["value"])
    except Exception as e:
        log.warning("menu_button %s: %s", r.get("id"), e); return None

def kb_main():
    k = types.InlineKeyboardMarkup(row_width=2)
    try: rows = [r for r in menu_load() if r["active"]]
    except Exception as e:
        log.warning("menu_load: %s", e); rows = [dict(id=0, label=lb, kind=kd, value=vl, style="", full=fl, active=1) for lb, kd, vl, fl in MENU_DEFAULTS]
    pend = []
    for r in rows:
        b = _menu_button(r)
        if not b: continue
        if r["full"]:
            if pend: k.row(*pend); pend = []
            k.row(b)
        else:
            pend.append(b)
            if len(pend) == 2: k.row(*pend); pend = []
    if pend: k.row(*pend)
    return k

def menu_click(d):
    try:
        with db() as con: con.execute("UPDATE menu_btns SET clicks=clicks+1 WHERE kind='cb' AND value=?", (d,))
    except Exception: pass

def menu_text_click(c, d):
    try: bid = int(d[5:])
    except ValueError: return
    with db() as con:
        r = con.execute("SELECT value FROM menu_btns WHERE id=? AND kind='text' AND active=1", (bid,)).fetchone()
        if r: con.execute("UPDATE menu_btns SET clicks=clicks+1 WHERE id=?", (bid,))
    if r: edit_msg(c.message.chat.id, c.message.message_id, r["value"], kb_home())

def kb_subjects(prefix="subj", back=True):
    k = types.InlineKeyboardMarkup(row_width=2)
    items = list(SUBJECTS.items())
    for i in range(0, len(items), 2):
        k.row(*[types.InlineKeyboardButton(f"{s['emoji']} {s['name']}", callback_data=f"{prefix}:{key}") for key, s in items[i:i + 2]])
    if back: k.row(types.InlineKeyboardButton("🏠 الرئيسية", callback_data="menu:home"))
    return k

def kb_subject(subj):
    B = types.InlineKeyboardButton
    return types.InlineKeyboardMarkup([
        [B("📖 اشرح موضوع", callback_data=f"subj_ask:{subj}")], [B("📝 اختبرني", callback_data=f"subj_quiz:{subj}")],
        [B("🏛 أسئلة وزارية", callback_data=f"subj_waz:{subj}")],
        [B("⬅️ المواد", callback_data="menu:subjects"), B("🏠 الرئيسية", callback_data="menu:home")]])

def kb_quiz_types(subj=""):
    p = f"qtype:{subj}:"; B = types.InlineKeyboardButton
    return types.InlineKeyboardMarkup([[B("✅ اختيار متعدد", callback_data=p + "mcq"), B("❓ صح وخطأ", callback_data=p + "tf")],
                                       [B("🔀 مختلط", callback_data=p + "mixed")], [B("🏠 الرئيسية", callback_data="menu:home")]])

def kb_tips():
    rows = [("🗓 جدول مراجعة", "schedule"), ("🧠 طرق الحفظ", "memory"), ("✍️ حل الوزاري", "past_exams"),
            ("😰 قلق الامتحان", "stress"), ("⏰ تنظيم الوقت", "time"), ("🍎 تغذية ورياضة", "health")]
    k = types.InlineKeyboardMarkup(row_width=2)
    k.add(*[types.InlineKeyboardButton(t, callback_data=f"tip:{c}") for t, c in rows])
    k.row(types.InlineKeyboardButton("🏠 الرئيسية", callback_data="menu:home"))
    return k

def kb_admin():
    B = types.InlineKeyboardButton
    return types.InlineKeyboardMarkup([
        [B("📊 إحصائيات", callback_data="admin:stats"), B("👥 المستخدمون", callback_data="admin:users")],
        [B("➕ إضافة كتاب", callback_data="admin:addbook"), B("🏛 إضافة وزاري", callback_data="admin:addwazari")],
        [B("📖 الكتب والوزاريات", callback_data="admin:books"), B("📢 القنوات", callback_data="admin:channels")],
        [B("📚 الملازم", callback_data="admin:materials"), B("🌐 رابط التطبيق", callback_data="admin:webapp")],
        [B("🧪 اختبار AI", callback_data="admin:test"), B("🧰 دليل الأوامر", callback_data="admin:help")]])


# ═══════════════════════════════════════════════════════════════════
#                              النصوص
# ═══════════════════════════════════════════════════════════════════

def home_text(name="طالب"):
    return (f"🎓 <b>أهلاً بيك يا {esc(name)}!</b>\n\nأنا <b>الأستاذ</b>، مدرّسك الذكي للسادس الإعدادي 👨‍🏫\n"
            "أدرس كتب المنهج والوزاريات وأجاوبك منها مع ذكر المصدر.\n\n"
            "📚 <b>أكدر أساعدك:</b>\n• شرح أي درس وحل المسائل خطوة بخطوة\n• أسئلة وزارية مع سنة كل سؤال\n• البحث عن ملزمتك بالقناة + رأي صريح بيها\n"
            "• نصائح عن مدرّسي اليوتيوب\n• قراءة صور الأسئلة 📸 والرسائل الصوتية 🎤\n\n🎯 <b>اختر من القائمة:</b>")

def help_text():
    return ("ℹ️ <b>دليل الاستخدام</b>\n━━━━━━━━━━━━━━━━\n\n"
            "📚 <b>اسأل</b> — اكتب سؤالك مباشرة\n📸 <b>صورة</b> — صوّر السؤال وأنا أقرأه وأحلّه\n🎤 <b>صوت</b> — تكلم وأنا أفهم\n"
            "🏛 <b>وزاري</b> — «اعطني أسئلة وزارية كيمياء» وأذكر لك سنة كل سؤال\n"
            "📄 <b>ملزمة</b> — «ملزمة فيزياء» أو «رأيك بملزمة الأستاذ فلان كيمياء»\n"
            "🎬 <b>مدرّس</b> — «رأيك بمدرّس يوتيوب للرياضيات»\n"
            "📝 <b>اختبار</b> — أسئلة تفاعلية من الكتاب\n🎨 <b>التطبيق</b> — واجهة كاملة داخل تيليجرام\n📤 <b>ساهم</b> — أضف ملزمة أو ملخصاً، تكسب نقاطاً تزيد حدّك اليومي حتى 200\n/contrib — إضافة ملزمة/ملخص\n\n"
            "👥 <b>في المجموعات:</b> ابدأ رسالتك بكلمة <b>استاذ</b> مثل: «استاذ شنو الميتوكوندريا؟»\n\n"
            "━━━━━━━━━━━━━━━━\n⌨️ <b>الأوامر:</b>\n/start — البداية\n/help — المساعدة\n/stats — إحصائياتي\n/wazari — أسئلة وزارية\n"
            "/clear — مسح ذاكرتي\n/webapp — التطبيق\n/materials — آخر الملازم\n/id — معرّفي\n\n"
            "⚠️ الإجابات من الذكاء الاصطناعي — راجع كتابك دائماً.")

def owner_help_text():
    return ("🛠 <b>أوامر المالك</b>\n━━━━━━━━━━━━━━━━\n"
            "/admin — لوحة الإدارة\n/addadmin ID — إضافة مشرف (للمالك فقط)\n/pending — إضافات الطلاب بانتظار الموافقة\n/deladmin ID\n/admins\n"
            "<b>الكتب والوزاريات</b>\n/addbook — إضافة كتب لمادة\n/addwazari — إضافة وزاريات (السنة والدور تُستخرج من اسم الملف أو التعليق، مثل «وزاري كيمياء 2023 دور ثاني»)\n"
            "/done — إنهاء وضع الإضافة\n/books — المصادر المخزّنة\n/delbook ID — حذف مصدر\n"
            "<b>الفصول</b>\n/chapter مادة | عنوان الفصل | ترتيب(اختياري)\n/autochapters مادة — يستخرجها من فهرس الكتاب\n/chapters [مادة]\n/delchapter ID\n"
            "<b>الملازم (روابط فقط)</b>\n/setchannel @القناة أو -100…\n/channels\n/material عنوان | المادة | رابط تيليجرام\n/editmat ID | العنوان | المادة\n/delmaterial ID\n/materials\n"
            "<b>مدرّسو اليوتيوب</b>\n/teacher الاسم | المادة | رابط القناة | ملاحظات\n/teachers\n/delteacher ID\n"
            "<b>أخرى</b>\n/saveqa السؤال | الجواب\n/users\n/ban ID\n/unban ID\n/broadcast نص\n/setwebapp https://…\n\n"
            "💡 <b>الملازم:</b> أضف البوت مشرفاً في قناتك فيلتقط أي ملف جديد (روابط فقط، لا يقرأ محتواه). للقديمة: حوّلها (Forward) من القناة إلى البوت.")


# ═══════════════════════════════════════════════════════════════════
#                          أوامر الطلاب
# ═══════════════════════════════════════════════════════════════════

@bot.message_handler(commands=["start"])
@safe
def cmd_start(m):
    uid = m.from_user.id
    if m.chat.type != "private":
        return reply(m, "👋 افتح محادثتي الخاصة واضغط /start لتفعيل حسابك. وهنا في المجموعة ابدأ رسالتك بكلمة <b>استاذ</b>.")
    with db() as con:
        existed = con.execute("SELECT 1 FROM users WHERE user_id=?", (uid,)).fetchone() is not None
    db_upsert_user(m.from_user, "telegram_start")
    if not existed:
        db_event(uid, "new_user", "telegram_start")
        for oid in OWNER_IDS:
            if oid != uid:
                try: say(oid, f"👤 طالب جديد\nالاسم: {esc(m.from_user.first_name or '—')}\nالمعرّف: <code>{uid}</code>\nاليوزر: @{esc(m.from_user.username or 'بدون')}")
                except Exception as e: log.warning("notify owner: %s", e)
    if db_is_banned(uid): return say(m.chat.id, "🚫 تم تقييد وصولك إلى البوت.")
    payload = (m.text or "").partition(" ")[2].strip()
    if payload == "contrib": return contrib_begin(uid, m.from_user, m.chat.id)
    if payload.startswith("login_"): return login_begin_bot(uid, m.from_user, payload[6:], m.chat.id)
    if payload.startswith("mat_") and payload[4:].isdigit(): return send_material_files(uid, int(payload[4:]))
    USER_MODE[uid] = "general"
    if WEB["url"]:
        try: bot.set_chat_menu_button(chat_id=uid, menu_button=types.MenuButtonWebApp(text="🎨 التطبيق", web_app=types.WebAppInfo(url=WEB["url"])))
        except Exception as e: log.warning("menu button: %s", e)
    say(uid, home_text(display_name(m.from_user)), reply_markup=kb_main())

@bot.message_handler(commands=["ping", "version"])
@safe
def cmd_ping(m): reply(m, "✅ المدرّس الذكي <b>v3</b> يعمل")

@bot.message_handler(commands=["id"])
@safe
def cmd_id(m): reply(m, f"معرّفك: <code>{m.from_user.id}</code>")

@bot.message_handler(commands=["help"])
@safe
def cmd_help(m): say(m.chat.id, help_text(), reply_markup=kb_home())

@bot.message_handler(commands=["stats"])
@safe
def cmd_stats(m):
    uid = m.from_user.id
    say(m.chat.id, stats_text(uid), reply_markup=kb_home())

def stats_text(uid):
    perf = db_student_summary(uid); a = activity_totals(uid)
    return ("📊 <b>إحصائياتك</b>\n━━━━━━━━━━━━━━━━\n"
            f"📅 طلبات اليوم: {db_usage_today(uid)} / {daily_limit_for(uid)}\n🔥 أيام متتالية: {streak_of(uid)}\n⭐ نقاط الأسبوع: {a['week_points']}\n"
            f"⏱ دقائق التركيز: {a['focus_total']}\n💬 رسائل محفوظة: {len(db_history(uid, 0, 400))}\n" + (f"🎯 اختباراتك: {esc(perf)}\n" if perf else ""))

@bot.message_handler(commands=["clear"])
@safe
def cmd_clear(m):
    db_clear_history(m.from_user.id, 0 if m.chat.type == "private" else m.chat.id)
    db_event(m.from_user.id, "history_cleared")
    reply(m, "🧹 تم مسح ذاكرة محادثتك أنت فقط.")

@bot.message_handler(commands=["webapp"])
@safe
def cmd_webapp(m):
    if not WEB["url"]: return reply(m, "⚠️ التطبيق غير مهيّأ بعد.")
    k = types.InlineKeyboardMarkup(); wb = webapp_btn("🎨 افتح التطبيق")
    if wb: k.row(wb)
    say(m.from_user.id, "🎨 <b>التطبيق المصغّر</b>", reply_markup=k)

@bot.message_handler(commands=["wazari"])
@safe
def cmd_wazari(m):
    USER_MODE[m.from_user.id] = "wazari"
    reply(m, "🏛 <b>أسئلة وزارية</b>\nاكتب المادة والموضوع وعدد الأسئلة.\n<b>مثال:</b> كيمياء الاتزان الكيميائي 5 أسئلة")

@bot.message_handler(commands=["materials"])
@safe
def cmd_materials(m):
    with db() as con: rows = con.execute("SELECT id,title,subject,telegram_url,description,norm,mkind,teacher,part,batch,mtype FROM materials WHERE active=1 ORDER BY id DESC LIMIT 15").fetchall()
    if not rows: return reply(m, "📭 المكتبة فارغة حالياً.")
    say(m.chat.id, materials_text([dict(r) for r in rows], "📚 <b>آخر الملازم</b>"), disable_web_page_preview=True)


# ═══════════════════════════════════════════════════════════════════
#                          أوامر المالك
# ═══════════════════════════════════════════════════════════════════

def admin_only(fn):
    @functools.wraps(fn)
    def w(m, *a, **k):
        if not is_admin(m.from_user.id): return
        return fn(m, *a, **k)
    return w

def owner_only(fn):
    @functools.wraps(fn)
    def w(m, *a, **k):
        if not is_owner(m.from_user.id): return
        return fn(m, *a, **k)
    return w

def _arg(m): return (m.text or "").partition(" ")[2].strip()
def _parts(m, n): return [x.strip() for x in _arg(m).split("|", n - 1)]

@bot.message_handler(commands=["ownerhelp"])
@admin_only
@safe
def cmd_ownerhelp(m): say(m.chat.id, owner_help_text())

@bot.message_handler(commands=["update"])
@safe
def cmd_update(m):
    if not is_owner(m.from_user.id): return
    if not GH_REPO: return say(m.chat.id, "ما مربوط GitHub. ضيف GITHUB_REPO=owner/repo بملف الإعدادات.")
    try: ch = selfupdate()
    except Exception as e: return say(m.chat.id, "❌ فشل التحديث: " + esc(str(e)[:300]))
    st = []
    for n in ("world.js", "atlas.png", "atlas.json"):
        p = os.path.join(STATIC_DIR, n)
        st.append(("✅ " if os.path.isfile(p) else "❌ ") + n + (f" ({os.path.getsize(p) // 1024}KB)" if os.path.isfile(p) else " ناقص"))
    try: nw = "fielddirt" in open(os.path.join(STATIC_DIR, "atlas.json"), encoding="utf-8").read()
    except OSError: nw = False
    st.append("الرسم الجديد: " + ("✅" if nw else "❌ ارفع atlas.json و atlas.png الجديدين"))
    stl = "\n" + "\n".join(st)
    if not ch: return say(m.chat.id, "✅ أنت على آخر نسخة." + stl)
    say(m.chat.id, "✅ تحدّث: " + ", ".join(ch) + stl + "\nجاري إعادة التشغيل…"); threading.Thread(target=_restart, daemon=True).start()

@bot.message_handler(commands=["admin"])
@admin_only
@safe
def cmd_admin(m): say(m.chat.id, admin_summary(), reply_markup=kb_admin())

def admin_summary():
    with db() as con:
        users = con.execute("SELECT COUNT(*) n FROM users").fetchone()["n"]
        b = con.execute("SELECT COUNT(*) n, COALESCE(SUM(chunks),0) c FROM sources WHERE kind='book'").fetchone()
        w = con.execute("SELECT COUNT(*) n FROM sources WHERE kind='wazari'").fetchone()["n"]
        mats = con.execute("SELECT COUNT(*) n FROM materials WHERE active=1").fetchone()["n"]
        ch = con.execute("SELECT COUNT(*) n FROM chapters").fetchone()["n"]
        tc = con.execute("SELECT COUNT(*) n FROM teachers WHERE active=1").fetchone()["n"]
    up = int((datetime.now() - STATS["started"]).total_seconds())
    return ("🛠 <b>لوحة المالك</b>\n━━━━━━━━━━━━━━━━\n"
            f"⏱️ التشغيل: {up // 60} دقيقة\n👥 المستخدمون: {users}\n📖 الكتب: {b['n']} ({b['c']} مقطع)\n🏛 الوزاريات: {w}\n"
            f"📚 الملازم: {mats}\n🗂 الفصول: {ch}\n🎬 المدرّسون: {tc}\n"
            f"💬 طلبات الجلسة: {STATS['requests']}  |  ❌ أخطاء: {STATS['errors']}\n⚙️ أعمال جارية: {len(INFLIGHT)}/{JOB_QUEUE_MAX}\n"
            f"🌐 التطبيق: <code>{esc(WEB['url'] or 'غير مهيّأ')}</code>")

@bot.message_handler(commands=["users"])
@owner_only
@safe
def cmd_users(m):
    with db() as con:
        total = con.execute("SELECT COUNT(*) n FROM users").fetchone()["n"]
        banned = con.execute("SELECT COUNT(*) n FROM users WHERE banned=1").fetchone()["n"]
        day = con.execute("SELECT COUNT(*) n FROM users WHERE last_seen>=?", ((datetime.utcnow() - timedelta(days=1)).isoformat(),)).fetchone()["n"]
    say(m.chat.id, f"👥 المستخدمون: {total}\n🟢 نشطون خلال 24 ساعة: {day}\n🚫 المحظورون: {banned}")

@bot.message_handler(commands=["ban", "unban"])
@owner_only
@safe
def cmd_ban(m):
    parts = (m.text or "").split(maxsplit=1)
    if len(parts) != 2 or not parts[1].strip().isdigit(): return reply(m, "الاستخدام: /ban ID أو /unban ID")
    target = int(parts[1]); flag = 1 if m.text.startswith("/ban") else 0
    if flag and target in OWNER_IDS: return reply(m, "⚠️ لا يمكن حظر مالك البوت.")
    with db() as con:
        con.execute("INSERT INTO users(user_id,joined_at,last_seen,banned) VALUES(?,?,?,?) ON CONFLICT(user_id) DO UPDATE SET banned=excluded.banned", (target, now_iso(), now_iso(), flag))
    _BAN.pop(target, None)
    db_event(target, "banned" if flag else "unbanned", m.from_user.id)
    reply(m, f"{'🚫 تم الحظر' if flag else '✅ تم رفع الحظر'}: <code>{target}</code>")

@bot.message_handler(commands=["broadcast"])
@owner_only
@safe
def cmd_broadcast(m):
    msg = _arg(m)
    if not msg: return reply(m, "الاستخدام: /broadcast نص الإعلان")
    with db() as con: ids = [r["user_id"] for r in con.execute("SELECT user_id FROM users WHERE banned=0 AND private_started=1")]
    def job():
        ok = fail = 0
        for t in ids:
            try: bot.send_message(t, msg); ok += 1
            except Exception as e:
                fail += 1
                w = _flood(e)
                if w: time.sleep(w)
            time.sleep(0.05)
        say(m.chat.id, f"📣 نجح: {ok} — فشل: {fail}")
    if run_owner_job(m.chat.id, m.from_user.id, job): reply(m, f"📣 بدأ الإرسال إلى {len(ids)} مستخدم…")

@bot.message_handler(commands=["setwebapp"])
@owner_only
@safe
def cmd_setwebapp(m):
    parts = m.text.split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip().startswith("https://"): return reply(m, "الاستخدام:\n<code>/setwebapp https://example.com</code>")
    WEB["url"] = parts[1].strip().rstrip("/")
    try: bot.set_chat_menu_button(menu_button=types.MenuButtonWebApp(text="🎨 التطبيق", web_app=types.WebAppInfo(url=WEB["url"])))
    except Exception as e: log.warning("menu button: %s", e)
    reply(m, f"✅ تم تعيين الرابط:\n<code>{esc(WEB['url'])}</code>")

# ── إدارة المشرفين (للمالك) ──
@bot.message_handler(commands=["addadmin", "deladmin", "admins"])
@owner_only
@safe
def cmd_admins(m):
    cmd = (m.text or "").split()[0].lstrip("/").split("@")[0]; p = _arg(m)
    if cmd == "admins":
        rows = admin_list()
        return reply(m, "المشرفون:\n" + ("\n".join(f"<code>{r['id']}</code> {esc(r['name'])}" for r in rows) if rows else "لا يوجد مشرفون غير المالك."))
    if not p.isdigit(): return reply(m, f"الاستخدام: /{cmd} ID")
    t = int(p)
    if cmd == "addadmin":
        if t in OWNER_IDS: return reply(m, "هذا مالك أصلاً.")
        admin_add(t, m.from_user.id); db_event(t, "admin_added", m.from_user.id)
        try: say(t, "تمت إضافتك مشرفاً في المدرّس الذكي. افتح التطبيق من حسابي ثم لوحة الإدارة.")
        except Exception: pass
        reply(m, f"تمت إضافة المشرف <code>{t}</code>.")
    else:
        reply(m, "تم الحذف." if admin_del(t) else "هذا المعرّف ليس مشرفاً.")

# ── الكتب والوزاريات ──
def start_addsource(chat_id, uid, kind="book"):
    OWNER_STATE.pop(uid, None)
    what = "الوزاريات" if kind == "wazari" else "الكتب"
    extra = "\n💡 اكتب السنة والدور في اسم الملف أو التعليق: «وزاري كيمياء 2023 دور ثاني»." if kind == "wazari" else ""
    say(chat_id, f"{'🏛' if kind == 'wazari' else '📖'} <b>إضافة {what}</b>\nاختر المادة، ثم أرسل الملفات (PDF / DOCX / TXT) واحداً واحداً:{extra}",
        reply_markup=kb_subjects("aw" if kind == "wazari" else "ab", back=False))

@bot.message_handler(commands=["addbook"])
@admin_only
@safe
def cmd_addbook(m): start_addsource(m.chat.id, m.from_user.id, "book")

@bot.message_handler(commands=["addwazari"])
@admin_only
@safe
def cmd_addwazari(m): start_addsource(m.chat.id, m.from_user.id, "wazari")

@bot.message_handler(commands=["done"])
@admin_only
@safe
def cmd_done(m):
    st = OWNER_STATE.pop(m.from_user.id, None)
    reply(m, "✅ انتهى وضع الإضافة." if st else "لا يوجد وضع إضافة نشط.")

def books_text():
    rows = sources_list()
    if not rows: return "📖 لا توجد مصادر مخزّنة بعد. استخدم /addbook أو /addwazari"
    lines = ["📖 <b>الكتب والوزاريات</b>"]; cur = None
    for r in rows:
        if r["subject"] != cur:
            cur = r["subject"]; lines.append(f"\n{SUBJECTS.get(cur, {}).get('emoji', '📚')} <b>{esc(subj_name(cur))}</b>")
        yr = f" · {r['year']}" if r.get("year") else ""
        lines.append(f"  <code>{r['id']}</code> · {esc(KIND_LABEL.get(r['kind'], r['kind']))}{yr} · {esc(r['title'])} ({r['pages']} ص، {r['chunks']} مقطع)")
    return "\n".join(lines) + "\n\nللحذف: /delbook ID"

@bot.message_handler(commands=["books"])
@admin_only
@safe
def cmd_books(m):
    for part in split_long(books_text()): say(m.chat.id, part)

@bot.message_handler(commands=["delbook"])
@admin_only
@safe
def cmd_delbook(m):
    p = _arg(m)
    if not p.isdigit(): return reply(m, "الاستخدام: /delbook ID")
    reply(m, "🗑 تم الحذف." if source_delete(int(p)) else "لم أجد هذا المعرّف.")

@bot.message_handler(commands=["saveqa"])
@admin_only
@safe
def cmd_saveqa(m):
    parts = _arg(m).split("|", 1)
    if len(parts) != 2 or not all(x.strip() for x in parts): return reply(m, "الاستخدام: /saveqa السؤال | الجواب")
    q, a = parts[0].strip(), parts[1].strip()
    res = source_add("إجابة: " + q[:60], guess_subject(q) or "general", "qa", [(0, f"{q}\n{a}")], m.from_user.id)
    reply(m, f"✅ حُفظ الجواب. ({res['chunks']} مقطع)")

# ── الفصول ──
@bot.message_handler(commands=["chapter"])
@admin_only
@safe
def cmd_chapter(m):
    p = _parts(m, 3)
    if len(p) < 2 or not p[1]: return reply(m, "الاستخدام: /chapter مادة | عنوان الفصل | ترتيب(اختياري)")
    s = subject_key(p[0])
    if s == "general": return reply(m, "اكتب اسم مادة صحيح (مثل: كيمياء).")
    order = int(p[2]) if len(p) > 2 and p[2].isdigit() else len(chapters_of(s)) + 1
    with db() as con:
        con.execute("INSERT INTO chapters(subject,title,ord) VALUES(?,?,?) ON CONFLICT(subject,title) DO UPDATE SET ord=excluded.ord", (s, p[1][:120], order))
    reply(m, f"✅ أضيف الفصل في {esc(subj_name(s))}: {esc(p[1])}")

@bot.message_handler(commands=["chapters"])
@admin_only
@safe
def cmd_chapters(m):
    a = _arg(m); subs = [subject_key(a)] if a else [k for k in SUBJECTS if k != "general"]
    lines = []
    for s in subs:
        cs = chapters_of(s)
        if cs: lines.append(f"\n{SUBJECTS[s]['emoji']} <b>{esc(subj_name(s))}</b>\n" + "\n".join(f"  <code>{c['id']}</code> · {esc(c['title'])}" for c in cs))
    say(m.chat.id, ("🗂 <b>الفصول</b>" + "".join(lines) + "\n\nللحذف: /delchapter ID") if lines else "لا توجد فصول. استخدم /chapter أو /autochapters مادة")

@bot.message_handler(commands=["delchapter"])
@admin_only
@safe
def cmd_delchapter(m):
    p = _arg(m)
    if not p.isdigit(): return reply(m, "الاستخدام: /delchapter ID")
    with db() as con: n = con.execute("DELETE FROM chapters WHERE id=?", (int(p),)).rowcount
    reply(m, "🗑 تم الحذف." if n else "لم أجد هذا المعرّف.")

def autochapters(subject):
    with db() as con:
        src = con.execute("SELECT id FROM sources WHERE kind='book' AND subject=? ORDER BY id LIMIT 1", (subject,)).fetchone()
        if not src: raise RuntimeError("ما عندي كتاب لهذه المادة. أضفه بـ /addbook أولاً.")
        rows = con.execute("SELECT raw FROM chunks WHERE source_id=? ORDER BY CAST(page AS INTEGER), rowid LIMIT 14", (src["id"],)).fetchall()
    text = "\n".join(r["raw"] for r in rows)[:7000]
    out = llm("أنت تستخرج فهارس الكتب. أعد JSON فقط: مصفوفة نصوص.", [],
              "هذا نص من بداية كتاب (غالباً الفهرس). استخرج عناوين الفصول الرئيسية بالترتيب بلا أرقام الصفحات، كمصفوفة JSON من النصوص فقط.\n\n" + text, max_tokens=800, temperature=0.1)
    mm = re.search(r"\[.*\]", out, re.S)
    arr = json.loads(mm.group(0)) if mm else []
    titles = [str(t).strip()[:100] for t in arr if 3 <= len(str(t).strip()) <= 100][:25]
    if not titles: raise RuntimeError("ما قدرت أستخرج الفصول. أضفها يدوياً بـ /chapter")
    with db() as con:
        for i, t in enumerate(titles, 1):
            con.execute("INSERT INTO chapters(subject,title,ord) VALUES(?,?,?) ON CONFLICT(subject,title) DO UPDATE SET ord=excluded.ord", (subject, t, i))
    return titles

@bot.message_handler(commands=["autochapters"])
@admin_only
@safe
def cmd_autochapters(m):
    s = subject_key(_arg(m))
    if s == "general": return reply(m, "الاستخدام: /autochapters كيمياء")
    def job():
        try:
            t = autochapters(s)
            say(m.chat.id, f"✅ أضفت {len(t)} فصلاً في {esc(subj_name(s))}:\n" + "\n".join(f"• {esc(x)}" for x in t) + "\n\nراجعها بـ /chapters وعدّل بـ /chapter أو /delchapter")
        except Exception as e: say(m.chat.id, f"❌ {esc(str(e)[:250])}")
    if run_owner_job(m.chat.id, m.from_user.id, job): reply(m, "🗂 أستخرج الفصول من فهرس الكتاب…")

# ── المدرّسون ──
@bot.message_handler(commands=["teacher"])
@admin_only
@safe
def cmd_teacher(m):
    p = _parts(m, 4)
    if len(p) < 2 or not p[0]: return reply(m, "الاستخدام: /teacher الاسم | المادة | رابط القناة | ملاحظات")
    url = p[2] if len(p) > 2 else ""
    if url and not url.startswith("https://"): return reply(m, "الرابط يجب أن يبدأ بـ https://")
    with db() as con:
        tid = con.execute("INSERT INTO teachers(name,subject,youtube_url,notes,created_at) VALUES(?,?,?,?,?)",
                          (p[0][:80], subject_key(p[1]), url, (p[3] if len(p) > 3 else "")[:400], now_iso())).lastrowid
    reply(m, f"✅ أضيف المدرّس. ID: {tid}")

@bot.message_handler(commands=["teachers"])
@admin_only
@safe
def cmd_teachers(m):
    rows = teachers_of()
    if not rows: return reply(m, "لا يوجد مدرّسون. استخدم /teacher")
    say(m.chat.id, "🎬 <b>المدرّسون</b>\n" + "\n".join(f"<code>{r['id']}</code> · {esc(r['name'])} — {esc(subj_name(subject_key(r['subject'])))}" for r in rows) + "\n\nللحذف: /delteacher ID")

@bot.message_handler(commands=["delteacher"])
@admin_only
@safe
def cmd_delteacher(m):
    p = _arg(m)
    if not p.isdigit(): return reply(m, "الاستخدام: /delteacher ID")
    with db() as con: n = con.execute("UPDATE teachers SET active=0 WHERE id=?", (int(p),)).rowcount
    reply(m, "🗑 تم الحذف." if n else "لم أجد هذا المعرّف.")

# ── الملازم والقنوات (روابط فقط، بلا قراءة الملف) ──
@bot.message_handler(commands=["material"])
@admin_only
@safe
def cmd_material(m):
    p = _parts(m, 3)
    if len(p) < 3: return reply(m, "الاستخدام: /material عنوان | المادة | https://t.me/...")
    title, subject, url = p
    if not url.startswith(("https://t.me/", "https://telegram.me/")): return reply(m, "الرابط يجب أن يكون رابط تيليجرام https://t.me/...")
    reply(m, f"✅ أضيفت الملزمة. ID: {material_add(title[:120], subject_key(subject), url, by=m.from_user.id)}")

@bot.message_handler(commands=["editmat"])
@admin_only
@safe
def cmd_editmat(m):
    p = _parts(m, 3)
    if not p or not p[0].isdigit(): return reply(m, "الاستخدام: /editmat ID | العنوان | المادة (اترك خانة فارغة لإبقائها)")
    with db() as con:
        r = con.execute("SELECT title,subject,description FROM materials WHERE id=?", (int(p[0]),)).fetchone()
        if not r: return reply(m, "لم أجد هذا المعرّف.")
        title = (p[1] if len(p) > 1 and p[1] else r["title"])[:120]
        subj = subject_key(p[2]) if len(p) > 2 and p[2] else r["subject"]
        con.execute("UPDATE materials SET title=?, subject=?, norm=? WHERE id=?", (title, subj, mat_norm(title, r["description"]), int(p[0])))
    reply(m, f"✅ تم التعديل: {esc(title)} — {esc(subj_name(subject_key(subj)))}")

@bot.message_handler(commands=["delmaterial"])
@admin_only
@safe
def cmd_delmaterial(m):
    p = _arg(m)
    if not p.isdigit(): return reply(m, "الاستخدام: /delmaterial ID")
    with db() as con: n = con.execute("UPDATE materials SET active=0 WHERE id=?", (int(p),)).rowcount
    reply(m, "🗑 تم حذفها من المكتبة." if n else "لم أجد هذا المعرّف.")

def _env_channel_ids():
    return {int(c) for c in CHANNEL_ENV if re.fullmatch(r"-?\d+", c)}

def channel_allowed(chat):
    if chat.id in _env_channel_ids(): return True
    if chat.username and f"@{chat.username}".lower() in {c.lower() for c in CHANNEL_ENV}: return True
    with db() as con: r = con.execute("SELECT approved FROM channels WHERE chat_id=?", (chat.id,)).fetchone()
    return bool(r and r["approved"])

def channel_register(chat, approved):
    with db() as con:
        con.execute("INSERT INTO channels(chat_id,title,username,approved,added_at) VALUES(?,?,?,?,?) "
                    "ON CONFLICT(chat_id) DO UPDATE SET title=excluded.title, username=excluded.username, approved=excluded.approved",
                    (chat.id, getattr(chat, "title", "") or "", getattr(chat, "username", None), 1 if approved else 0, now_iso()))

@bot.message_handler(commands=["setchannel"])
@admin_only
@safe
def cmd_setchannel(m):
    ref = _arg(m)
    if not ref: return reply(m, "الاستخدام: /setchannel @اسم_القناة\nتأكد أن البوت مشرف في القناة.")
    try: chat = bot.get_chat(int(ref) if re.fullmatch(r"-?\d+", ref) else ref)
    except Exception as e: return reply(m, f"❌ ما كدرت أوصل للقناة: {esc(str(e)[:120])}\nأضف البوت مشرفاً فيها أولاً.")
    channel_register(chat, True)
    reply(m, f"✅ اعتمدت القناة: <b>{esc(chat.title)}</b>\nأي ملف جديد تنشره فيها يُضاف للمكتبة كرابط.")

def channels_text():
    with db() as con: rows = con.execute("SELECT * FROM channels ORDER BY added_at DESC").fetchall()
    if not rows and not CHANNEL_ENV: return "📢 لا توجد قنوات. أضف البوت مشرفاً في قناتك أو استخدم /setchannel"
    lines = ["📢 <b>القنوات</b>"]
    for r in rows: lines.append(f"{'✅' if r['approved'] else '⏳'} {esc(r['title'] or r['chat_id'])} <code>{r['chat_id']}</code>")
    for c in CHANNEL_ENV: lines.append(f"✅ (من الإعدادات) <code>{esc(c)}</code>")
    return "\n".join(lines)

@bot.message_handler(commands=["channels"])
@admin_only
@safe
def cmd_channels(m): say(m.chat.id, channels_text())

def material_link(chat, message_id):
    if getattr(chat, "username", None): return f"https://t.me/{chat.username}/{message_id}"
    cid = str(chat.id)
    return f"https://t.me/c/{cid[4:] if cid.startswith('-100') else cid.lstrip('-')}/{message_id}"

def register_channel_doc(chat, message_id, doc, caption, by=0):
    """يسجّل ملف القناة كرابط في المكتبة (لا يقرأ محتوى الملزمة). يرجع نص ملخّص أو None إن كانت مسجّلة."""
    with db() as con:
        if con.execute("SELECT 1 FROM materials WHERE channel_id=? AND message_id=?", (chat.id, message_id)).fetchone(): return None
    fname = doc.file_name or "ملزمة"
    cap = (caption or "").strip()
    title = (cap.splitlines()[0].strip() if cap else Path(fname).stem.replace("_", " ").replace("-", " ")).strip()[:120]
    subject = guess_subject(f"{cap} {fname}") or "general"
    mid = material_add(title, subject, material_link(chat, message_id), f"{cap} {Path(fname).stem}", "channel", by, chat.id, message_id)
    db_event(by, "material_from_channel", f"{mid}:{title}")
    return (f"📥 ملزمة جديدة <code>{mid}</code>: <b>{esc(title)}</b> — {esc(subj_name(subject))}"
            + ("\n⚠️ ما عرفت المادة — صحّحها بـ <code>/editmat " + str(mid) + " | | كيمياء</code>" if subject == "general" else ""))

@bot.channel_post_handler(content_types=["document"])
@safe
def on_channel_doc(m):
    chat = m.chat
    if not channel_allowed(chat):
        with db() as con: seen = con.execute("SELECT 1 FROM channels WHERE chat_id=?", (chat.id,)).fetchone()
        if not seen:
            channel_register(chat, False)
            k = types.InlineKeyboardMarkup([[types.InlineKeyboardButton("✅ اعتماد", callback_data=f"chan:ok:{chat.id}"), types.InlineKeyboardButton("❌ رفض", callback_data=f"chan:no:{chat.id}")]])
            for oid in OWNER_IDS:
                try: say(oid, f"📢 وصلت رسالة من قناة جديدة:\n<b>{esc(chat.title)}</b> <code>{chat.id}</code>\nهل أعتمدها كمصدر ملازم؟", reply_markup=k)
                except Exception: pass
        return
    info = register_channel_doc(chat, m.message_id, m.document, m.caption)
    if info:
        for oid in OWNER_IDS:
            try: say(oid, info, disable_web_page_preview=True)
            except Exception: pass


# ═══════════════════════════════════════════════════════════════════
#                              الأزرار
# ═══════════════════════════════════════════════════════════════════

TIP_PROMPTS = {
    "schedule": "أنشئ جدول مراجعة يومي لطالب سادس إعدادي عراقي على 30 يوم قبل الامتحانات.",
    "memory": "أعطني 8 تقنيات علمية فعّالة للحفظ والفهم. اشرح كل واحدة بمثال.",
    "past_exams": "اشرح طريقة الاستفادة القصوى من الأسئلة الوزارية. خطة أسبوعية.",
    "stress": "أعطني 8 نصائح لقلق الامتحان + تقنية تنفّس خطوة بخطوة.",
    "time": "استراتيجيات إدارة الوقت: بومودورو، التوازن، تجنّب التسويف.",
    "health": "نصائح غذائية: أطعمة للذاكرة، ساعات النوم، رياضة خفيفة.",
}

MODE_PROMPTS = {
    "ask": ("general", "📚 <b>اسأل الأستاذ</b>\n\nاكتب أي سؤال أو موضوع وسأشرحه من كتبك.\n\n<b>أمثلة:</b>\n• اشرح جهاز كولجي\n• شنو قانون أوم؟\n• وضّح التمثيل الضوئي\n\nأو أرسل صورة / رسالة صوتية."),
    "handout": ("handout", "📄 <b>ملزمة من القناة</b>\n\nاكتب المادة أو اسم الملزمة/المدرّس وأرسل لك الروابط المناسبة.\nوإذا تريد رأيي الصريح بيها اكتب «رأيك»:\n\n<b>أمثلة:</b>\n• ملزمة فيزياء\n• رأيك بملزمة الأستاذ فلان كيمياء"),
    "teacher": ("teacher", "🎬 <b>مدرّسو اليوتيوب</b>\n\nاكتب المادة أو اسم المدرّس وأعطيك رأيي الصريح مما يُقال عنه بالسوشيال ميديا.\n\n<b>أمثلة:</b>\n• أفضل مدرّس رياضيات يوتيوب\n• رأيك بالمدرّس فلان فيزياء"),
    "wazari": ("wazari", "🏛 <b>أسئلة وزارية</b>\n\nاكتب المادة والموضوع وعدد الأسئلة، وأذكر لك سنة كل سؤال.\n\n<b>مثال:</b> كيمياء الاتزان الكيميائي 5 أسئلة"),
}

@bot.callback_query_handler(func=lambda c: True)
@safe
def on_callback(c):
    try: bot.answer_callback_query(c.id)
    except Exception: pass
    if c.message is None: return
    d = c.data or ""; uid = c.from_user.id
    if db_is_banned(uid): return
    chat_id = c.message.chat.id; msg_id = c.message.message_id
    ed = lambda text, kb=None: edit_msg(chat_id, msg_id, text, kb)
    if d.startswith("cbtn:"): return menu_text_click(c, d)
    if d.startswith("menu:"): menu_click(d)

    if d == "menu:home":
        USER_MODE[uid] = "general"; ed(home_text(display_name(c.from_user)), kb_main())
    elif d.startswith("menu:") and d[5:] in MODE_PROMPTS:
        mode, txt = MODE_PROMPTS[d[5:]]; USER_MODE[uid] = mode; ed(txt, kb_home())
    elif d == "menu:contrib": contrib_begin(uid, c.from_user, chat_id)
    elif d in ("cfinish", "ccancel") or d.startswith(("ck:", "cb:", "ct:", "cskip:")): contrib_cb(c, d)
    elif d.startswith(("apv:", "rej:")) and is_admin(uid): contrib_decide(c, d)
    elif d.startswith("view:") and is_admin(uid): send_material_files(uid, int(d[5:]) if d[5:].isdigit() else 0, True)
    elif d.startswith("delmat:") and is_admin(uid):
        mid = int(d.split(":")[1]) if d.split(":")[1].isdigit() else 0
        with db() as con:
            r = con.execute("SELECT channel_id,msg_ids FROM materials WHERE id=?", (mid,)).fetchone()
            con.execute("UPDATE materials SET active=0 WHERE id=?", (mid,))
        try:
            for x in json.loads((r["msg_ids"] if r else "") or "[]"):
                try: bot.delete_message(r["channel_id"], x)
                except Exception: pass
        except Exception: pass
        ed("🗑 تم حذف الإضافة من المكتبة والقناة.")
    elif d == "menu:quiz":
        USER_MODE[uid] = "quiz"; ed("📝 <b>اختبرني</b>\n\nاختر النوع:", kb_quiz_types(USER_SUBJECT.get(uid, "")))
    elif d == "menu:subjects": ed("🎯 <b>اختر المادة:</b>", kb_subjects())
    elif d == "menu:files": ed("📸 <b>ملف أو صورة</b>\n\nأرسل صورة سؤال وأحلّه لك مباشرة، أو PDF / DOCX / TXT.\n🔹 بدون تعليق: شرح كامل\n🔹 <code>اختبرني</code>: اختبار من المحتوى\n🔹 <code>لخّص</code>: ملخص قصير", kb_home())
    elif d == "menu:voice": ed("🎤 <b>سؤال صوتي</b>\n\nأرسل رسالة صوتية وسأحوّلها إلى نص وأجيبك.", kb_home())
    elif d == "menu:tips": ed("💡 <b>نصائح الدراسة</b>\n\nاختر:", kb_tips())
    elif d == "menu:stats": ed(stats_text(uid), kb_home())
    elif d == "menu:help": ed(help_text(), kb_home())
    elif d.startswith("subj:"):
        s = d.split(":")[1]; USER_SUBJECT[uid] = s; USER_MODE[uid] = "subject"
        ed(f"{SUBJECTS.get(s, {}).get('emoji', '')} <b>{esc(subj_name(s))}</b>\n\nأنا مدرّسك في {esc(subj_name(s))}. اكتب سؤالك أو اختر:", kb_subject(s))
    elif d.startswith("subj_ask:"):
        s = d.split(":")[1]; USER_SUBJECT[uid] = s; USER_MODE[uid] = "subject"
        ed(f"{SUBJECTS.get(s, {}).get('emoji', '')} <b>{esc(subj_name(s))}</b>\n\nاكتب الموضوع اللي تريد شرحه:", kb_home())
    elif d.startswith("subj_quiz:"):
        s = d.split(":")[1]; USER_SUBJECT[uid] = s; USER_MODE[uid] = "quiz"
        ed(f"📝 <b>اختبار {esc(subj_name(s))}</b>\n\nاختر النوع:", kb_quiz_types(s))
    elif d.startswith("subj_waz:"):
        s = d.split(":")[1]; USER_SUBJECT[uid] = s; USER_MODE[uid] = "wazari"
        ed(f"🏛 <b>وزاري {esc(subj_name(s))}</b>\n\nاكتب الموضوع وعدد الأسئلة (أو أرسل رقماً فقط لأسئلة عامة).\n<b>مثال:</b> 5 أسئلة", kb_home())
    elif d.startswith("qtype:"):
        p = d.split(":"); s = p[1] if len(p) > 1 else ""; t = p[2] if len(p) > 2 else "mixed"
        USER_MODE[uid] = "quiz"; USER_SUBJECT[uid] = s; USER_QTYPE[uid] = t
        names = {"mcq": "اختيار متعدد", "tf": "صح وخطأ", "mixed": "مختلط"}
        ed(f"✅ <b>النوع:</b> {names.get(t, 'مختلط')}\n📚 <b>المادة:</b> {esc(subj_name(s)) if s else 'أي مادة'}\n\nاكتب الموضوع وعدد الأسئلة.\n<b>مثال:</b> جهاز كولجي، 5 أسئلة\n(سأرسلها كاختبار تفاعلي تضغط فيه على الجواب)", kb_home())
    elif d.startswith("tip:"):
        ed("⏳ جاري التحضير...")
        ok, _ = check_daily(uid)
        if not ok: return ed(f"⚠️ وصلت حدّك اليومي ({daily_limit_for(uid)}).", kb_home())
        st, key = claim(uid)
        if st:
            refund_daily(uid); return ed("⏳ لسا أجاوب على طلبك السابق، حاول بعد لحظة.", kb_home())
        def job():
            try: send_answer(chat_id, llm(SYSTEM_TIPS, [], TIP_PROMPTS.get(d.split(":")[1], "نصيحة دراسية."), max_tokens=1800), kb=kb_home())
            except Exception as e:
                STATS["errors"] += 1; refund_daily(uid); ed(err_text(e, uid), kb_home())
        submit(key, job)

    # ── المالك ──
    elif d.startswith(("ab:", "aw:")) and is_admin(uid):
        s = d.split(":")[1]; kind = "wazari" if d.startswith("aw:") else "book"
        OWNER_STATE[uid] = {"subject": s, "kind": kind}
        ed(f"{'🏛' if kind == 'wazari' else '📖'} <b>وضع إضافة {'وزاريات' if kind == 'wazari' else 'كتب'} — {esc(subj_name(s))}</b>\n\nأرسل الملفات الآن (PDF / DOCX / TXT) <b>واحداً واحداً</b>. "
           "العنوان = اسم الملف (أو تعليق الرسالة)." + ("\nالسنة والدور من العنوان: «وزاري كيمياء 2023 دور ثاني»." if kind == "wazari" else "") +
           "\nعند الانتهاء أرسل /done\n\nℹ️ حدّ تيليجرام 20MB للملف؛ قسّم الكبير. الملفات الممسوحة ضوئياً تُقرأ بالرؤية الذكية (حتى " + str(OCR_MAX_PAGES) + " صفحة).")
    elif d.startswith("chan:") and is_admin(uid):
        _, act, cid = d.split(":"); cid = int(cid)
        try: chat = bot.get_chat(cid)
        except Exception: chat = types.Chat(cid, "channel")
        if act == "ok": channel_register(chat, True); ed(f"✅ اعتمدت القناة <b>{esc(getattr(chat, 'title', cid))}</b>. حوّل لي آخر ملف منها لأسجّله.")
        else:
            with db() as con: con.execute("DELETE FROM channels WHERE chat_id=?", (cid,))
            ed("❌ رُفضت القناة.")
    elif d.startswith("admin:") and is_admin(uid):
        sub = d.split(":")[1]
        if sub == "stats": ed(admin_summary(), kb_admin())
        elif sub == "users":
            with db() as con:
                total = con.execute("SELECT COUNT(*) n FROM users").fetchone()["n"]
                banned = con.execute("SELECT COUNT(*) n FROM users WHERE banned=1").fetchone()["n"]
                day = con.execute("SELECT COUNT(*) n FROM users WHERE last_seen>=?", ((datetime.utcnow() - timedelta(days=1)).isoformat(),)).fetchone()["n"]
            ed(f"👥 <b>المستخدمون</b>\nالإجمالي: {total}\nنشطون خلال 24 ساعة: {day}\nالمحظورون: {banned}", kb_admin())
        elif sub == "books":
            for part in split_long(books_text()): say(chat_id, part)
        elif sub == "addbook": start_addsource(chat_id, uid, "book")
        elif sub == "addwazari": start_addsource(chat_id, uid, "wazari")
        elif sub == "channels": ed(channels_text(), kb_admin())
        elif sub == "materials":
            with db() as con: rows = [dict(r) for r in con.execute("SELECT id,title,subject,telegram_url,description,norm,mkind,teacher,part,batch,mtype FROM materials WHERE active=1 ORDER BY id DESC LIMIT 10")]
            ed((materials_text(rows, "📚 <b>آخر الملازم</b>") if rows else "المكتبة فارغة.")[:3900], kb_admin())
        elif sub == "help": ed(owner_help_text(), kb_admin())
        elif sub == "webapp": ed(f"🌐 <b>رابط التطبيق</b>\n<code>{esc(WEB['url'] or 'غير مهيّأ')}</code>\n\nلتغييره: <code>/setwebapp https://…</code>", kb_admin())
        elif sub == "test":
            ed("🧪 جاري الاختبار...")
            def mask(k): return (k[:4] + "…" + k[-4:]) if k else "غير موجود"
            diag = f"\n\n🔑 Groq: <code>{esc(', '.join(mask(k) for k in GROQ_KEYS) or 'غير موجود')}</code>\n🔑 Gemini: <code>{esc(', '.join(mask(k) for k in GEMINI_KEYS) or 'غير موجود')}</code>"
            def job():
                try: ed(f"✅ الذكاء الاصطناعي يعمل:\n{esc(llm('أجب بكلمة واحدة.', [], 'قل: تم', max_tokens=20))}" + diag, kb_admin())
                except Exception as e: ed(f"❌ {esc(str(e)[:700])}" + diag, kb_admin())
            run_owner_job(chat_id, uid, job)


# ═══════════════════════════════════════════════════════════════════
#          مساهمات الطلاب: إضافة ملزمة / ملخص (تُنشر بالقناة)
# ═══════════════════════════════════════════════════════════════════

CONTRIB = {}
CONTRIB_TTL = 40 * 60
CONTRIB_POINTS = {"mulzama": 25, "summary": 20}
CONTRIB_PUBLISH_BONUS = 10
CONTRIB_DAILY_MAX = 5
CONTRIB_MAX_FILES = 10
CONTRIB_EXTS = {".pdf", ".doc", ".docx", ".ppt", ".pptx", ".txt", ".zip", ".rar", ".jpg", ".jpeg", ".png", ".webp"}
MTYPES = {"wazari": "وزاريات", "manhaj": "منهج", "other": "أخرى"}
_PUB = {"id": None, "t": 0.0}

def _contrib_state(uid):
    st = CONTRIB.get(uid)
    if st and time.time() - st["t"] > CONTRIB_TTL:
        CONTRIB.pop(uid, None); return None
    return st

def kb_contrib(skip=None, rows=None):
    B = types.InlineKeyboardButton; k = types.InlineKeyboardMarkup()
    for r in rows or []: k.row(*[B(t, callback_data=d) for t, d in r])
    if skip: k.row(B("تخطي", callback_data=f"cskip:{skip}"))
    k.row(B("إلغاء", callback_data="ccancel"))
    return k

def contrib_ask(chat_id, uid):
    st = CONTRIB[uid]; step = st["step"]; B = types.InlineKeyboardButton
    if step == "kind":
        say(chat_id, "📤 <b>إضافة للمكتبة</b>\nشنو تريد تضيف؟", reply_markup=kb_contrib(rows=[[("📘 ملزمة", "ck:mulzama"), ("📝 ملخص", "ck:summary")]]))
    elif step == "subject":
        k = kb_subjects("cb", back=False); k.row(B("إلغاء", callback_data="ccancel")); say(chat_id, "اختر المادة:", reply_markup=k)
    elif step == "teacher": say(chat_id, "شسم المدرس صاحب الملزمة؟ (اكتب الاسم)", reply_markup=kb_contrib())
    elif step == "part": say(chat_id, "يا جزء؟ مثال: الأول، الثاني، كامل", reply_markup=kb_contrib(skip="part"))
    elif step == "batch": say(chat_id, "يا دفعة؟ مثال: 2025", reply_markup=kb_contrib(skip="batch"))
    elif step == "mtype": say(chat_id, "شنو نوع الملزمة؟", reply_markup=kb_contrib(rows=[[("وزاريات", "ct:wazari"), ("منهج", "ct:manhaj")], [("أخرى", "ct:other")]]))
    elif step == "desc":
        if st["kind"] == "mulzama": say(chat_id, "اكتب وصفاً قصيراً للملزمة (شنو تحتوي؟ لأي فصول؟) وينزل مع المنشور، أو اضغط تخطي.", reply_markup=kb_contrib(skip="desc"))
        else: say(chat_id, "اكتب وصف الملخص بسطر أو سطرين (شنو يحتوي؟ لأي فصل؟)", reply_markup=kb_contrib())
    elif step == "files":
        what = "ملف الملزمة (PDF أو Word أو PowerPoint)" if st["kind"] == "mulzama" else "ملف الملخص أو صوره (تكدر ترسل أكثر من صورة)"
        say(chat_id, f"أرسل الآن {what}. بعد ما تخلص اضغط «انتهيت ونشر».", reply_markup=kb_contrib())

def contrib_begin(uid, user, chat_id):
    if db_is_banned(uid): return say(chat_id, "🚫 تم تقييد وصولك إلى البوت.")
    CONTRIB[uid] = {"step": "kind", "files": [], "t": time.time()}
    contrib_ask(chat_id, uid)

def contrib_cb(c, d):
    uid = c.from_user.id; chat_id = c.message.chat.id
    if d == "ccancel":
        CONTRIB.pop(uid, None); return say(chat_id, "تم الإلغاء.", reply_markup=kb_home())
    st = _contrib_state(uid)
    if not st: return say(chat_id, "انتهت جلسة الإضافة. ابدأ من جديد من القائمة.", reply_markup=kb_home())
    st["t"] = time.time()
    if d.startswith("ck:"): st["kind"] = d[3:] if d[3:] in ("mulzama", "summary") else "mulzama"; st["step"] = "subject"
    elif d.startswith("cb:"): st["subject"] = d[3:] if d[3:] in SUBJECTS else "general"; st["step"] = "teacher" if st.get("kind") == "mulzama" else "desc"
    elif d.startswith("ct:"): st["mtype"] = MTYPES.get(d[3:], "أخرى"); st["step"] = "desc"
    elif d.startswith("cskip:"):
        f = d[6:]
        if f == "part": st["part"] = ""; st["step"] = "batch"
        elif f == "batch": st["batch"] = ""; st["step"] = "mtype"
        elif f == "desc": st["desc"] = ""; st["step"] = "files"
    elif d == "cfinish": return contrib_finish(chat_id, c.from_user)
    contrib_ask(chat_id, uid)

@bot.message_handler(commands=["contrib", "add"])
@safe
def cmd_contrib(m):
    if m.chat.type != "private": return reply(m, "افتح محادثتي الخاصة لإضافة ملزمة أو ملخص.")
    contrib_begin(m.from_user.id, m.from_user, m.chat.id)

@bot.message_handler(commands=["cancel"])
@safe
def cmd_cancel(m):
    CONTRIB.pop(m.from_user.id, None); OWNER_STATE.pop(m.from_user.id, None); reply(m, "تم الإلغاء.")

@bot.message_handler(content_types=["text"], func=lambda m: m.chat.type == "private" and m.from_user.id in CONTRIB and not (m.text or "").startswith("/"))
@safe
def on_contrib_text(m):
    uid = m.from_user.id; st = _contrib_state(uid)
    if not st: return say(m.chat.id, "انتهت جلسة الإضافة. ابدأ من جديد من القائمة.", reply_markup=kb_home())
    t = (m.text or "").strip()[:300]; step = st["step"]; st["t"] = time.time()
    if step == "teacher": st["teacher"] = t[:60]; st["step"] = "part"
    elif step == "part": st["part"] = t[:30]; st["step"] = "batch"
    elif step == "batch": st["batch"] = t[:20]; st["step"] = "mtype"
    elif step == "desc": st["desc"] = t; st["step"] = "files"
    else: return reply(m, "استخدم الأزرار الظاهرة أو أرسل الملف المطلوب.")
    contrib_ask(m.chat.id, uid)

def _file_exists(u):
    with db() as con:
        return con.execute("SELECT 1 FROM materials WHERE (active=1 OR status='pending') AND file_ids LIKE ? LIMIT 1", (f'%"u": "{u}"%',)).fetchone() is not None

@bot.message_handler(content_types=["photo", "document"], func=lambda m: m.chat.type == "private" and m.from_user.id in CONTRIB and CONTRIB[m.from_user.id].get("step") == "files")
@safe
def on_contrib_file(m):
    uid = m.from_user.id; st = _contrib_state(uid)
    if not st: return
    if len(st["files"]) >= CONTRIB_MAX_FILES: return reply(m, f"وصلت الحد الأقصى ({CONTRIB_MAX_FILES} ملفات). اضغط «انتهيت ونشر».")
    if m.content_type == "photo":
        p = m.photo[-1]; f = {"t": "photo", "id": p.file_id, "u": p.file_unique_id, "name": "صورة"}
    else:
        d = m.document; ext = Path(d.file_name or "").suffix.lower()
        if ext not in CONTRIB_EXTS: return reply(m, "نوع الملف غير مدعوم. المدعوم: PDF و Word و PowerPoint و ZIP والصور.")
        f = {"t": "document", "id": d.file_id, "u": d.file_unique_id, "name": d.file_name or "ملف"}
    if any(x["u"] == f["u"] for x in st["files"]): return
    if _file_exists(f["u"]): return reply(m, "هذا الملف موجود مسبقاً في المكتبة. شكراً لك!")
    st["files"].append(f); st["t"] = time.time()
    if time.time() - st.get("ack", 0) > 2.5:                 # لا نردّ على كل صورة بألبوم
        st["ack"] = time.time()
        B = types.InlineKeyboardButton
        reply(m, "استلمت. أرسل المزيد أو اضغط انتهيت.", reply_markup=types.InlineKeyboardMarkup([[B("✅ انتهيت ونشر", callback_data="cfinish")], [B("إلغاء", callback_data="ccancel")]]))

def pick_publish_channel():
    if time.time() - _PUB["t"] < 600 and _PUB["id"]: return _PUB["id"]
    cid = None
    try:
        if PUBLISH_CHANNEL: cid = bot.get_chat(int(PUBLISH_CHANNEL) if re.fullmatch(r"-?\d+", PUBLISH_CHANNEL) else PUBLISH_CHANNEL).id
        elif _env_channel_ids(): cid = sorted(_env_channel_ids())[0]
        else:
            named = [c for c in CHANNEL_ENV if not re.fullmatch(r"-?\d+", c)]
            if named: cid = bot.get_chat(named[0]).id
            else:
                with db() as con: r = con.execute("SELECT chat_id FROM channels WHERE approved=1 ORDER BY added_at LIMIT 1").fetchone()
                cid = r["chat_id"] if r else None
    except Exception as e:
        log.warning("publish channel: %s", e)
    _PUB.update(id=cid, t=time.time())
    return cid

def publish_files(chan, files, caption):
    ids, first = [], True
    docs = [f for f in files if f["t"] != "photo"]; photos = [f for f in files if f["t"] == "photo"]
    for f in docs:
        ids.append(bot.send_document(chan, f["id"], caption=caption if first else None).message_id); first = False; time.sleep(0.4)
    for i in range(0, len(photos), 10):
        grp = photos[i:i + 10]
        if len(grp) == 1:
            ids.append(bot.send_photo(chan, grp[0]["id"], caption=caption if first else None).message_id)
        else:
            media = [types.InputMediaPhoto(g["id"], caption=(caption if (first and j == 0) else None)) for j, g in enumerate(grp)]
            ids += [mm.message_id for mm in bot.send_media_group(chan, media)]
        first = False; time.sleep(0.5)
    return ids

def contrib_finish(chat_id, user):
    uid = user.id; st = _contrib_state(uid)
    if not st or not st["files"]: return say(chat_id, "ما أرسلت أي ملف بعد. أرسل الملف أولاً.")
    with db() as con:
        n_today = con.execute("SELECT COUNT(*) n FROM materials WHERE added_by=? AND src='student' AND created_at>=?", (uid, datetime.utcnow().strftime("%Y-%m-%d"))).fetchone()["n"]
    if n_today >= CONTRIB_DAILY_MAX and not is_admin(uid):
        CONTRIB.pop(uid, None); return say(chat_id, f"وصلت حدّ الإضافات اليومي ({CONTRIB_DAILY_MAX}). كمّل باجر.", reply_markup=kb_home())
    code, key = claim(uid, "contrib")
    if code: return say(chat_id, "لسا أنشر إضافتك السابقة، انتظر لحظة.")
    CONTRIB.pop(uid, None)
    submit(key, lambda: _contrib_publish(chat_id, user, st))

def _contrib_texts(st):
    kind = st["kind"]; sname = subj_name(st.get("subject", "general")); descr = (st.get("desc") or "").strip()
    if kind == "mulzama":
        bits = [f"ملزمة {sname}", f"الأستاذ {st.get('teacher', '')}"] + ([f"جزء {st['part']}"] if st.get("part") else []) + ([f"دفعة {st['batch']}"] if st.get("batch") else []) + ([st["mtype"]] if st.get("mtype") else [])
        return " — ".join(bits), descr
    return f"ملخص {sname} — {descr[:60]}", descr

def publish_material(mid):
    """ينشر الإضافة بالقناة (بمعرّف الملف دون تنزيل أو ذكاء اصطناعي، فلا حد لحجمه) ثم يفعّلها ويمنح الطالب نقاطه."""
    with db() as con: r = con.execute("SELECT * FROM materials WHERE id=?", (int(mid),)).fetchone()
    if not r: raise RuntimeError("الإضافة غير موجودة")
    r = dict(r)
    try: files = json.loads(r.get("file_ids") or "[]")
    except Exception: files = []
    uname = db_get_user(r["added_by"]).get("first_name") or "طالب"
    descr = r["description"] if r.get("description") and r["description"] != r["title"] else ""
    caption = (r["title"] + (f"\n{descr}" if descr else "") + f"\nبواسطة: {uname}")[:1000]
    chan = pick_publish_channel(); msg_ids, url = [], ""
    if chan and files:
        try: msg_ids = publish_files(chan, files, caption)
        except Exception as e: log.warning("publish: %s", e); msg_ids = []
        if msg_ids:
            try: url = material_link(bot.get_chat(chan), msg_ids[0])
            except Exception as e: log.warning("link: %s", e)
    with db() as con:
        con.execute("UPDATE materials SET active=1, status='approved', telegram_url=?, channel_id=?, message_id=?, msg_ids=? WHERE id=?",
                    (url, chan if msg_ids else None, msg_ids[0] if msg_ids else None, json.dumps(msg_ids), int(mid)))
    pts = CONTRIB_POINTS.get(r.get("mkind") or "mulzama", 20) + (CONTRIB_PUBLISH_BONUS if msg_ids else 0)
    touch_activity(r["added_by"], pts); _LIM.pop(r["added_by"], None); db_event(r["added_by"], "contrib_published", f"{mid}")
    return {"url": url, "pts": pts, "published": bool(msg_ids), "title": r["title"], "by": r["added_by"]}

def notify_admins_pending(mid, user, title, descr, files):
    B = types.InlineKeyboardButton
    cap = (f"إضافة جديدة بانتظار موافقتك\n{title}" + (f"\n{descr}" if descr else "") + f"\nمن: {display_name(user)} ({user.id})")[:1000]
    kb = types.InlineKeyboardMarkup([[B("✅ نشر", callback_data=f"apv:{mid}"), B("❌ رفض", callback_data=f"rej:{mid}")]])
    for oid in set(OWNER_IDS) | set(admins_cached()):
        try:
            try: publish_files(oid, files[:3], cap)                  # معاينة للمشرف
            except Exception as e: log.warning("preview: %s", e)
            say(oid, f"📥 <b>إضافة #{mid} بانتظار الموافقة</b>\n{esc(title)}\nمن: {esc(display_name(user))} (<code>{user.id}</code>)", reply_markup=kb)
        except Exception as e: log.warning("notify admin %s: %s", oid, e)

def _contrib_publish(chat_id, user, st):
    """إضافة العضو تُحفظ بانتظار موافقة الإدارة؛ إضافة المشرف تُنشر مباشرة."""
    uid = user.id; kind = st["kind"]; subj = st.get("subject", "general")
    wait = say(chat_id, "جاري حفظ إضافتك...")
    try:
        title, descr = _contrib_texts(st); direct = is_admin(uid)
        mid = material_add(title[:120], subj, "", descr or title, "student", uid, None, None, mkind=kind, teacher=st.get("teacher", ""), part=st.get("part", ""),
                           batch=st.get("batch", ""), mtype=st.get("mtype", ""), file_ids=json.dumps(st["files"], ensure_ascii=False), msg_ids="", src="student",
                           status="approved" if direct else "pending", active=1 if direct else 0)
        db_event(uid, "contrib_submit", f"{kind}:{mid}")
        if direct:
            res = publish_material(mid)
            edit_msg(wait.chat.id, wait.message_id, f"تم النشر {'بالقناة' if res['published'] else 'بالمكتبة'}.\n<b>{esc(title)}</b>\n+{res['pts']} نقطة.", kb_home())
        else:
            edit_msg(wait.chat.id, wait.message_id, f"تم استلام إضافتك ✅\n<b>{esc(title)}</b>\nهي الآن بانتظار موافقة الإدارة. إذا وافقوا تُنشر بالقناة وتنضاف نقاطك.", kb_home())
            notify_admins_pending(mid, user, title, descr, st["files"])
    except Exception:
        STATS["errors"] += 1; log.exception("contrib submit")
        edit_msg(wait.chat.id, wait.message_id, "❌ تعذّر حفظ الإضافة، جرّب مرة ثانية.")

def contrib_decide(c, d):
    act, sid = d.split(":", 1); mid = int(sid) if sid.isdigit() else 0
    chat_id, msg_id = c.message.chat.id, c.message.message_id
    with db() as con:
        r = con.execute("SELECT title,added_by,status FROM materials WHERE id=? AND src='student'", (mid,)).fetchone()
        if not r: return edit_msg(chat_id, msg_id, "هذه الإضافة غير موجودة.")
        won = con.execute("UPDATE materials SET status='processing' WHERE id=? AND status='pending'", (mid,)).rowcount     # أول مشرف يضغط يحسمها
    if not won:
        st = {"approved": "نُشرت", "rejected": "رُفضت"}.get(r["status"], "قيد المعالجة")
        return edit_msg(chat_id, msg_id, f"تم البتّ بهذه الإضافة مسبقاً ({st}).")
    title = r["title"]
    if act == "apv":
        edit_msg(chat_id, msg_id, "جاري النشر...")
        try:
            res = publish_material(mid)
        except Exception as e:
            with db() as con: con.execute("UPDATE materials SET status='pending' WHERE id=?", (mid,))
            log.exception("approve"); return edit_msg(chat_id, msg_id, "❌ فشل النشر، جرّب مرة ثانية: " + esc(str(e)[:150]))
        edit_msg(chat_id, msg_id, f"✅ نُشرت: <b>{esc(title)}</b>" + (f"\n{esc(res['url'])}" if res["url"] else "\n(حُفظت بالمكتبة؛ لم أجد قناة للنشر. اعتمد قناة بأمر /setchannel)"))
        try: say(r["added_by"], f"وافقت الإدارة على إضافتك ونُشرت.\n<b>{esc(title)}</b>\n+{res['pts']} نقطة. شكراً لمساهمتك!")
        except Exception: pass
    else:
        with db() as con: con.execute("UPDATE materials SET status='rejected', active=0 WHERE id=?", (mid,))
        edit_msg(chat_id, msg_id, f"❌ رُفضت: <b>{esc(title)}</b>")
        try: say(r["added_by"], f"نعتذر، لم تتم الموافقة على إضافتك «{esc(title)}». تأكد من وضوح الملف ومطابقة المادة وجرّب مرة ثانية.")
        except Exception: pass

@bot.message_handler(commands=["pending"])
@admin_only
@safe
def cmd_pending(m):
    with db() as con:
        rows = con.execute("SELECT id,title,added_by FROM materials WHERE src='student' AND status='pending' ORDER BY id LIMIT 10").fetchall()
    if not rows: return reply(m, "لا توجد إضافات بانتظار الموافقة.")
    B = types.InlineKeyboardButton
    for r in rows:
        say(m.chat.id, f"📥 <b>#{r['id']}</b> {esc(r['title'])}\nمن: <code>{r['added_by']}</code>",
            reply_markup=types.InlineKeyboardMarkup([[B("✅ نشر", callback_data=f"apv:{r['id']}"), B("❌ رفض", callback_data=f"rej:{r['id']}")], [B("👁 عرض الملفات", callback_data=f"view:{r['id']}")]]))

def send_material_files(uid, mid, any_status=False):
    ok, _ = check_rate(uid)
    if not ok: return
    with db() as con: r = con.execute("SELECT * FROM materials WHERE id=?" + ("" if any_status else " AND active=1"), (int(mid),)).fetchone()
    if not r: return say(uid, "هذه الملزمة غير موجودة أو حُذفت.")
    r = dict(r)
    try: files = json.loads(r.get("file_ids") or "[]")
    except Exception: files = []
    say(uid, f"<b>{esc(r['title'])}</b>\n<i>{esc(mat_meta(r))}</i>" + (f"\n{esc(r['telegram_url'])}" if r.get("telegram_url") else ""), disable_web_page_preview=True)
    for f in files[:12]:
        try:
            (bot.send_photo if f.get("t") == "photo" else bot.send_document)(uid, f["id"]); time.sleep(0.3)
        except Exception as e: log.warning("send material: %s", e)


# ── تأكيد دخول التطبيق عبر البوت (رمز من 4 أرقام يمنع سرقة الحساب برابط مفخّخ) ──
LOGIN_WAIT = {}

def login_begin_bot(uid, user, tok, chat_id):
    r = _login_row(tok)
    if not r or r["status"] != "pending": return say(chat_id, "انتهت صلاحية طلب الدخول أو استُخدم مسبقاً. افتح التطبيق واطلب رمزاً جديداً.")
    LOGIN_WAIT[uid] = {"token": tok, "tries": 0, "t": time.time()}
    say(chat_id, f"🔐 <b>طلب تسجيل دخول للتطبيق</b>\nالجهاز: {esc(device_name(r['ua']))}\n\nاكتب هنا <b>الرمز المكوّن من 4 أرقام</b> الظاهر على شاشة التطبيق لتأكيد أنك أنت من طلب الدخول.\nإذا ما طلبت هذا الدخول فتجاهل الرسالة.")

@bot.message_handler(content_types=["text"], func=lambda m: m.chat.type == "private" and m.from_user.id in LOGIN_WAIT and re.fullmatch(r"\s*[0-9\u0660-\u0669]{4}\s*", m.text or "") is not None)
@safe
def on_login_code(m):
    uid = m.from_user.id; w = LOGIN_WAIT.get(uid)
    if not w or time.time() - w["t"] > LOGIN_TTL:
        LOGIN_WAIT.pop(uid, None); return reply(m, "انتهت صلاحية الطلب. افتح التطبيق واطلب رمزاً جديداً.")
    r = _login_row(w["token"])
    if not r:
        LOGIN_WAIT.pop(uid, None); return reply(m, "انتهت صلاحية الطلب. افتح التطبيق واطلب رمزاً جديداً.")
    code = (m.text or "").strip().translate(_AR_DIGITS)
    if not hmac.compare_digest(code, r["code"]):
        w["tries"] += 1
        if w["tries"] >= 3:
            LOGIN_WAIT.pop(uid, None)
            with db() as con: con.execute("DELETE FROM login_tokens WHERE token=?", (w["token"],))
            return reply(m, "الرمز غلط 3 مرات، أُلغي الطلب. افتح التطبيق واطلب رمزاً جديداً.")
        return reply(m, f"الرمز غير صحيح. باقي لك {3 - w['tries']} محاولة.")
    with db() as con: con.execute("UPDATE login_tokens SET status='approved', user_id=? WHERE token=?", (uid, w["token"]))
    LOGIN_WAIT.pop(uid, None); db_upsert_user(m.from_user, "telegram_message"); db_event(uid, "app_login", device_name(r["ua"]))
    reply(m, "✅ تم تسجيل دخولك للتطبيق. ارجع له الآن.")


# ═══════════════════════════════════════════════════════════════════
#                       الرسائل النصية وتوجيه النية
# ═══════════════════════════════════════════════════════════════════

MATERIAL_WORDS = tuple(ar_norm(w) for w in ("ملزمة", "ملازم", "ملخص", "ملخصات"))
OPINION_RE = re.compile(r"رايك|رايكم|تنصح|ترشح|تستاهل|يستاهل|احسن|افضل|شلون|مفيد|تقييم|قيم|زينه|حلوه")
TEACHER_RE = re.compile(r"مدرس|يوتيوب|يوتيوبر")
WAZARI_RE = re.compile(r"وزاري")
WAZARI_ASK_RE = re.compile(r"اسئله|سؤال|اختبر|اختبار|نماذج|اعطني|عطني|اريد|ورد|جاب|اجت|سنه|مثال|دز")

def _terms(text):
    return [t for t in dict.fromkeys(tokens(text)) if not _gen(t)]

def send_quiz_polls(chat_id, qs):
    for i, q in enumerate(qs, 1):
        head = f"{i}. " + (f"🏛 {q['tag']}\n" if q.get("tag") else "")
        question = (head + latex_to_unicode(q["q"]))[:300]
        opts = [latex_to_unicode(o)[:100] for o in q["options"]]
        expl = latex_to_unicode(q.get("explain") or "")
        if q.get("src") and q["src"] not in expl: expl = (expl[:120] + "\n📖 " + q["src"]) if expl else "📖 " + q["src"]
        try:
            bot.send_poll(chat_id, question, opts, is_anonymous=False, type="quiz", correct_option_id=q["answer"], explanation=expl[:195] or None)
            time.sleep(0.4)
        except Exception as e:
            log.warning("poll failed: %s", e)
            say(chat_id, f"{i}. {esc(q['q'])}\n" + "\n".join(f"  {chr(0x61 + j)}) {esc(o)}" for j, o in enumerate(q["options"])) + f"\n✔ الجواب: {esc(q['options'][q['answer']])}")

def parse_quiz_request(text):
    n = 5
    mm = re.search(r"(\d+)\s*(?:سؤال|اسئلة|أسئلة|اسئله)?", ar_norm(text))
    if mm: n = max(1, min(10, int(mm.group(1))))
    topic = re.sub(r"[\d٠-٩]+\s*(?:سؤال|أسئلة|اسئلة|اسئله)?", "", text).strip(" ،,.-")
    return topic, n

def opinion_job(m, key, subj, target, extra, head):
    uid = m.from_user.id
    def job():
        wait = reply(m, "🔎 أدوّر شنو يگولون عنها بالسوشيال ميديا...")
        try:
            txt = web_opinion(key, subj or "general", target, extra)
            del_msg(wait)
            send_answer(m.chat.id, head + txt, m=m)
        except Exception as e:
            STATS["errors"] += 1; refund_daily(uid); log.warning("opinion: %s", e)
            edit_msg(wait.chat.id, wait.message_id, err_text(e, uid))
    return job

def ai_rank_job(m, text, cands, op=None):
    uid = m.from_user.id
    def job():
        wait = reply(m, "ثواني، أرتّب لك الأنسب...")
        try:
            rows = ai_rank_materials(text, cands)
            del_msg(wait)
            reply(m, materials_text(rows, f"📚 أنسب {len(rows)} حسب طلبك:"), disable_web_page_preview=True)
        except Exception as e:
            STATS["errors"] += 1; refund_daily(uid); log.warning("rank: %s", e)
            return edit_msg(wait.chat.id, wait.message_id, err_text(e, uid))
        if op:
            try: opinion_job(m, *op)()
            except Exception: log.exception("opinion after rank")
    return job

def handle_handout(m, text, group):
    """يبحث عن الملزمة بالمكتبة؛ عند كثرة النتائج يفرزها الذكاء الاصطناعي ويقدّم الأنسب، ثم رأي صريح إن طُلب."""
    if not gate(m, group, count=False): return
    cands = search_materials(text, 30); rows = cands[:5]
    subj = guess_subject(text); terms = _terms(text)
    wants_opinion = bool(OPINION_RE.search(ar_norm(text)))
    op = None
    if wants_opinion and terms:
        op = ("h:" + " ".join(sorted(terms))[:90] + ":" + (subj or ""), subj, f"ملزمة/ملخص (كما كتبها الطالب): {text[:200]}", "", "🗣 **رأيي الصريح:**\n\n")
    if len(cands) > 5:
        run_heavy(m, ai_rank_job(m, text, cands, op), group=group); return
    if rows:
        reply(m, materials_text(rows, f"📚 لقيت لك {len(rows)} بالمكتبة:"), disable_web_page_preview=True)
    else:
        with db() as con: total = con.execute("SELECT COUNT(*) n FROM materials WHERE active=1").fetchone()["n"]
        if not total: reply(m, "📭 المكتبة فارغة حالياً — ما انضافت ملازم بعد. تكدر تضيف أنت من زر «أضف ملزمة / ملخص».")
        elif not subj and not terms: return reply(m, "اكتب اسم المادة أو الملزمة/المدرّس، مثلاً: «ملزمة فيزياء» أو «ملزمة الأستاذ فلان كيمياء».")
        else: reply(m, "🔎 ما لقيت ملزمة بهذا الاسم بالمكتبة." + (" بس أكدر أعطيك رأيي بيها 👇" if op else ""))
    if wants_opinion:
        if not op: return reply(m, "حتى أعطيك رأيي الصريح، اكتب اسم الملزمة أو المدرّس، مثلاً: «رأيك بملزمة ... كيمياء».")
        run_heavy(m, opinion_job(m, *op), group=group)

def handle_teacher(m, text, group):
    if not gate(m, group, count=False): return
    subj = guess_subject(text) or (USER_SUBJECT.get(m.from_user.id, "") if not group else "") or None
    terms = _terms(text)
    rows = teachers_of(subj) if subj else teachers_of()
    if terms:
        rows = [r for r in rows if all(any(x.startswith(t) or t.startswith(x) for x in tokens(r["name"])) for t in terms)]
    if rows:
        lines = ["🎬 <b>مدرّسون عندي:</b>"]
        for i, r in enumerate(rows[:5], 1):
            lines.append(f"{i}. <b>{esc(r['name'])}</b> — {esc(subj_name(subject_key(r['subject'])))}" + (f"\n{esc(r['notes'])}" if r["notes"] else "") + (f"\n{esc(r['youtube_url'])}" if r["youtube_url"] else ""))
        reply(m, "\n\n".join(lines), disable_web_page_preview=True)
    if terms:
        extra = "؛ ".join(f"{r['name']}: {r['notes']}" for r in rows[:2] if r["notes"])
        key = "t:" + " ".join(sorted(terms))[:90] + ":" + (subj or "")
        run_heavy(m, opinion_job(m, key, subj, f"مدرّس يوتيوب (كما كتبه الطالب): {text[:200]}", extra, "🗣 **رأيي الصريح:**\n\n"), group=group)
    elif subj:
        names = "، ".join(r["name"] for r in rows[:8])
        key = "y:" + subj
        run_heavy(m, opinion_job(m, key, subj, f"أفضل مدرّسي يوتيوب لمادة {subj_name(subj)} للسادس الإعدادي في العراق وكيف يختار الطالب بينهم",
                                 f"مدرّسون مسجّلون عندنا: {names}" if names else "", "🎬 **نصيحتي لمدرّسي اليوتيوب:**\n\n"), group=group)
    elif not rows:
        reply(m, "اكتب المادة، مثلاً: «أفضل مدرّس رياضيات يوتيوب» أو «رأيك بالمدرّس فلان».")

def handle_wazari(m, text, group):
    uid = m.from_user.id
    subject = guess_subject(text) or (USER_SUBJECT.get(uid, "") if not group else "")
    topic, n = parse_quiz_request(re.sub(r"وزاري\w*|أسئلة|اسئلة|اسئله|سؤال|اعطني|عطني|اريد|أريد|دز|لي", " ", text))
    topic = " ".join(t for t in tokens(topic) if not _gen(t))
    def job():
        wait = reply(m, "🏛 أجهّز لك أسئلة وزارية...")
        try:
            note = ""
            try: qs, info = make_quiz(subject, topic, n, "mixed", wazari=True)
            except NoSource:
                qs, info = make_quiz(subject, topic, n, USER_QTYPE.get(uid, "mixed"))
                note = "\nℹ️ ما عندي وزاريات لهالمادة بعد، فهذا اختبار من المنهج."
            if not qs: raise RuntimeError("ما قدرت أكوّن أسئلة، جرّب موضوعاً أوضح")
            STATS["quizzes"] += 1; touch_activity(uid, 1)
            del_msg(wait)
            say(m.chat.id, f"🏛 <b>أسئلة {'وزارية' if info['kind'] == 'wazari' else 'اختبار'} {esc(subj_name(subject)) if subject else ''}</b> — {len(qs)} أسئلة" + (" (السنة مكتوبة على كل سؤال)" if info["kind"] == "wazari" else "") + note)
            send_quiz_polls(m.chat.id, qs)
        except Exception as e:
            STATS["errors"] += 1; refund_daily(uid); log.warning("wazari: %s", e)
            edit_msg(wait.chat.id, wait.message_id, err_text(e, uid))
    run_heavy(m, job, group=group)

@bot.message_handler(content_types=["text"], func=lambda m: not (m.text or "").startswith("/"))
@safe
def on_text(m):
    if stale(m): return
    uid = m.from_user.id; group = is_group(m)
    text = (m.text or "").strip()
    if not text: return
    if group:
        q = strip_call(text)
        if q is None: return
        db_upsert_user(m.from_user, "telegram_group")
        if db_is_banned(uid): return
        if not q: return reply(m, f"نعم يا {esc(display_name(m.from_user))}؟ اكتب سؤالك بعد كلمة «استاذ» 🙂")
        text = q
    else:
        if m.chat.type != "private": return
        db_upsert_user(m.from_user, "telegram_message")
    text = text[:2000]
    low = ar_norm(text)
    mode = USER_MODE.get(uid, "general") if not group else "general"
    if mode != "general" and classify_intent(text) == "chat": mode = "general"      # دردشة داخل وضع مقيّد: لا نفرضها كطلب ملزمة/وزاري

    if mode == "handout" or any(w in low for w in MATERIAL_WORDS): return handle_handout(m, text, group)
    if mode == "teacher" or (TEACHER_RE.search(low) and OPINION_RE.search(low)): return handle_teacher(m, text, group)
    if mode == "wazari" or (WAZARI_RE.search(low) and WAZARI_ASK_RE.search(low)): return handle_wazari(m, text, group)

    subject = USER_SUBJECT.get(uid, "") if (not group and mode in ("subject", "quiz")) else ""
    cid = m.chat.id if group else 0
    name = display_name(m.from_user)

    def job():
        wait = reply(m, "⏳ لحظة...")
        try:
            if mode == "quiz" and not group:
                topic, n = parse_quiz_request(text)
                qs, info = make_quiz(subject, topic or text, n, USER_QTYPE.get(uid, "mixed"))
                STATS["quizzes"] += 1
                if not qs: raise RuntimeError("ما قدرت أكوّن اختباراً، جرّب موضوعاً أوضح")
                touch_activity(uid, 1); del_msg(wait)
                say(m.chat.id, f"📝 <b>اختبار {esc(subj_name(subject)) if subject else ''}</b> — {len(qs)} أسئلة" + (" (من الكتب 📖)" if info["grounded"] else ""))
                return send_quiz_polls(m.chat.id, qs)
            ans, hits = ask_teacher(uid, text, subject, cid, "telegram", name, group)
            del_msg(wait)
            send_answer(m.chat.id, ans, m=m, hits=hits)
        except Exception as e:
            STATS["errors"] += 1; refund_daily(uid); log.warning("text fail: %s", e)
            edit_msg(wait.chat.id, wait.message_id, err_text(e, uid))
    STATS["requests"] += 1
    run_heavy(m, job, group=group)


# ═══════════════════════════════════════════════════════════════════
#                       الصور (OCR → النموذج مباشرة)
# ═══════════════════════════════════════════════════════════════════

def caption_instruction(caption, what):
    c = ar_norm(caption or "")
    if "اختبر" in c or "quiz" in c: return f"أنشئ اختباراً من {what}: 5 اختيار متعدد + 3 صح/خطأ + 2 مقالي. مفتاح الإجابة في النهاية.", SYSTEM_STRICT
    if "لخص" in c: return f"لخّص {what} في 5-7 نقاط.", SYSTEM_STRICT
    if (caption or "").strip(): return f"{caption.strip()}\n(اعتمد على {what} فقط)", SYSTEM_STRICT
    return f"اشرح {what} كاملاً بالتفصيل. في النهاية: خلاصة 3 نقاط.", SYSTEM_STRICT

@bot.message_handler(content_types=["photo"])
@safe
def on_photo(m):
    if stale(m): return
    uid = m.from_user.id; group = is_group(m)
    caption = (m.caption or "").strip()
    if group:
        q = strip_call(caption)
        if q is None: return
        caption = q
    db_upsert_user(m.from_user, "telegram_media")
    cid = m.chat.id if group else 0
    def job():
        wait = reply(m, "🔎 أقرأ الصورة وأحلّها...")
        try:
            text, _ = read_image(download(m.photo[-1].file_id))
            if len(text.strip()) < 8: raise RuntimeError("ما قرأت نصاً واضحاً بالصورة، جرّب صورة أوضح وأقرب.")
            STATS["photos"] += 1
            q = (caption + "\n" if caption else "") + text[:6000]
            ans, hits = ask_teacher(uid, q, guess_subject(q) or "", cid, "telegram", display_name(m.from_user), group, shown=("[صورة] " + caption).strip(), from_image=True, max_tokens=1800)
            del_msg(wait)
            send_answer(m.chat.id, ans, m=m, hits=hits)     # النص المقروء لا يُعرض؛ الجواب فقط
        except Exception as e:
            STATS["errors"] += 1; refund_daily(uid); log.warning("photo: %s", e)
            edit_msg(wait.chat.id, wait.message_id, err_text(e, uid) if isinstance(e, (Busy, RuntimeError)) and (is_owner(uid) or "صورة" in str(e) or isinstance(e, Busy)) else "❌ ما كدرت أقرأ الصورة، جرّب صورة أوضح.")
    STATS["requests"] += 1
    run_heavy(m, job, group=group)


# ═══════════════════════════════════════════════════════════════════
#                              الصوت
# ═══════════════════════════════════════════════════════════════════

@bot.message_handler(content_types=["voice", "audio"])
@safe
def on_voice(m):
    if stale(m) or is_group(m): return
    uid = m.from_user.id
    db_upsert_user(m.from_user, "telegram_media")
    voice = m.voice or m.audio
    if not voice: return
    def job():
        wait = reply(m, "🎤 جاري التحويل...")
        try:
            text = ai_whisper(download(voice.file_id, MAX_AUDIO))
            if not text.strip(): raise RuntimeError("ما فهمت شي من الصوت")
            STATS["voices"] += 1
            edit_msg(wait.chat.id, wait.message_id, f"📝 <b>سمعتك تقول:</b>\n{esc(text)}")
            subject = USER_SUBJECT.get(uid, "") if USER_MODE.get(uid) == "subject" else ""
            ans, hits = ask_teacher(uid, text, subject, 0, "telegram", display_name(m.from_user))
            send_answer(m.chat.id, ans, hits=hits)
        except Exception as e:
            STATS["errors"] += 1; refund_daily(uid); log.warning("voice: %s", e)
            edit_msg(wait.chat.id, wait.message_id, err_text(e, uid))
    STATS["requests"] += 1
    run_heavy(m, job)


# ═══════════════════════════════════════════════════════════════════
#                              الملفات
# ═══════════════════════════════════════════════════════════════════

def forwarded_channel(m):
    origin = getattr(m, "forward_origin", None)
    if origin is not None and getattr(origin, "type", "") == "channel": return origin.chat, origin.message_id
    fc = getattr(m, "forward_from_chat", None)
    if fc is not None and getattr(fc, "type", "") == "channel": return fc, getattr(m, "forward_from_message_id", None)
    return None, None

@bot.message_handler(content_types=["document"])
@safe
def on_document(m):
    if stale(m): return
    uid = m.from_user.id; group = is_group(m)
    caption = (m.caption or "").strip(); doc = m.document
    suffix = Path(doc.file_name or "").suffix.lower()
    if group:
        q = strip_call(caption)
        if q is None: return
        caption = q
    db_upsert_user(m.from_user, "telegram_media")

    if is_admin(uid) and not group:
        ch, mid = forwarded_channel(m)
        if ch is not None and mid:                       # ملف محوَّل من القناة → رابط فقط
            channel_register(ch, True)
            try: info = register_channel_doc(ch, mid, doc, m.caption, uid)
            except Exception as e: info = f"❌ فشل: {esc(str(e)[:200])}"
            return reply(m, info or "ℹ️ هذه الملزمة مسجّلة مسبقاً.", disable_web_page_preview=True)
        st = OWNER_STATE.get(uid)
        if st or caption.startswith("/ingest"):          # وضع إضافة كتب/وزاريات
            subject = (st or {}).get("subject", "general"); kind = (st or {}).get("kind", "book")
            title = (caption.partition(" ")[2].strip() if caption.startswith("/ingest") else caption) or Path(doc.file_name or "كتاب").stem.replace("_", " ")
            def job():
                wait = reply(m, f"📖 أقرأ «{esc(title[:80])}»… (قد تستغرق دقائق للكتب الكبيرة أو الممسوحة)")
                try:
                    prog = lambda i, n: edit_msg(wait.chat.id, wait.message_id, f"📷 أقرأ الصفحات الممسوحة… {i}/{n}")
                    res = ingest_bytes(download(doc.file_id), doc.file_name or "x.pdf", title[:120], subject, kind, uid, progress=prog)
                    db_event(uid, f"{kind}_added", f"{title}:{res['chunks']}")
                    msg = (f"✅ <b>{esc(title[:100])}</b>\n{'🏛' if kind == 'wazari' else '📚'} {esc(subj_name(subject))} · {res['pages']} صفحة · {res['chunks']} مقطع"
                           + "".join(f"\n{esc(n)}" for n in res["notes"]) + ("\n\nأرسل الملف التالي أو /done" if st else ""))
                except Exception as e:
                    STATS["errors"] += 1; msg = f"❌ فشل: {esc(str(e)[:250])}"
                edit_msg(wait.chat.id, wait.message_id, msg)
            run_owner_job(m.chat.id, uid, job)
            return

    if suffix not in (".pdf", ".docx", ".txt"):
        if not group: reply(m, "📄 المدعوم: PDF / DOCX / TXT")
        return
    if doc.file_size and doc.file_size > MAX_STUDENT_FILE:
        return reply(m, "📄 الملف كبير (الحد 8MB). أرسل الجزء المطلوب فقط.") if not group else None
    def job():
        wait = reply(m, "📄 جاري التحليل...")
        tmp = None
        try:
            raw = download(doc.file_id, MAX_STUDENT_FILE)
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as f:
                f.write(raw); tmp = f.name
            content = read_doc_text(tmp, suffix, 22000)
            if not content.strip(): raise RuntimeError("ما ظهر نص في الملف")
            instr, system = caption_instruction(caption, "الملف")
            ans = llm(system, [], instr + "\n\nالمحتوى:\n" + content, max_tokens=3200)
            STATS["docs"] += 1; del_msg(wait)
            send_answer(m.chat.id, ans, m=m)
        except Exception as e:
            STATS["errors"] += 1; refund_daily(uid); log.warning("doc: %s", e)
            edit_msg(wait.chat.id, wait.message_id, err_text(e, uid))
        finally:
            if tmp and os.path.exists(tmp): os.unlink(tmp)
    run_heavy(m, job, group=group)


# ═══════════════════════════════════════════════════════════════════
#                       واجهات التطبيق المصغّر (API)
# ═══════════════════════════════════════════════════════════════════

# ═══════════════════════════════════════════════════════════════════
#     تطبيق مستقل (PWA): دخول عبر البوت + أيقونات + تحديث تلقائي
# ═══════════════════════════════════════════════════════════════════

MAIN_HASH = hashlib.sha1(Path(__file__).read_bytes()).hexdigest()[:8]     # يتغيّر مع كل تحديث لملف البوت
LOGIN_TTL = 600
SESSION_DAYS = 90

def login_create(ua):
    tok = secrets.token_urlsafe(12); code = f"{secrets.randbelow(10000):04d}"
    with db() as con:
        con.execute("DELETE FROM login_tokens WHERE created_at < ?", ((datetime.utcnow() - timedelta(seconds=LOGIN_TTL * 2)).isoformat(timespec="seconds"),))
        con.execute("INSERT INTO login_tokens(token,code,status,ua,created_at) VALUES(?,?,?,?,?)", (tok, code, "pending", (ua or "")[:200], now_iso()))
    return tok, code

def _login_row(tok):
    with db() as con: r = con.execute("SELECT * FROM login_tokens WHERE token=?", (tok,)).fetchone()
    if not r: return None
    r = dict(r)
    if datetime.utcnow() - datetime.fromisoformat(r["created_at"]) > timedelta(seconds=LOGIN_TTL): return None
    return r

def device_name(ua):
    ua = ua or ""
    os_ = "Android" if "Android" in ua else "iPhone" if "iPhone" in ua else "iPad" if "iPad" in ua else "Windows" if "Windows" in ua else "Mac" if "Macintosh" in ua else "Linux" if "Linux" in ua else "جهاز"
    br = "Firefox" if "Firefox" in ua else "Edge" if "Edg" in ua else "Chrome" if "Chrome" in ua else "Safari" if "Safari" in ua else "متصفح"
    return f"{br} على {os_}"

def session_make(uid, ua=""):
    raw = secrets.token_urlsafe(32); h = hashlib.sha256(raw.encode()).hexdigest()
    with db() as con:
        con.execute("INSERT INTO sessions(token_hash,user_id,created_at,last_seen,ua) VALUES(?,?,?,?,?)", (h, int(uid), now_iso(), now_iso(), (ua or "")[:200]))
        con.execute("DELETE FROM sessions WHERE user_id=? AND token_hash NOT IN (SELECT token_hash FROM sessions WHERE user_id=? ORDER BY created_at DESC LIMIT 5)", (int(uid), int(uid)))
    return raw

def login_poll(tok, ua):
    r = _login_row(tok)
    if not r: return {"status": "expired"}
    if r["status"] != "approved" or not r["user_id"]: return {"status": "pending"}
    with db() as con: con.execute("DELETE FROM login_tokens WHERE token=?", (tok,))
    return {"status": "approved", "session": session_make(r["user_id"], ua)}

_SESS_CACHE = {}
def session_user(raw):
    if not raw or len(raw) > 100: return None
    h = hashlib.sha256(raw.encode()).hexdigest(); c = _SESS_CACHE.get(h)
    if c and time.time() - c[1] < 300: return c[0]
    with db() as con:
        r = con.execute("SELECT s.user_id uid, s.last_seen ls, u.first_name fn, u.username un FROM sessions s JOIN users u ON u.user_id=s.user_id WHERE s.token_hash=?", (h,)).fetchone()
        if not r: return None
        age = datetime.utcnow() - datetime.fromisoformat(r["ls"])
        if age > timedelta(days=SESSION_DAYS):
            con.execute("DELETE FROM sessions WHERE token_hash=?", (h,)); return None
        if age > timedelta(hours=6): con.execute("UPDATE sessions SET last_seen=? WHERE token_hash=?", (now_iso(), h))
    user = {"id": r["uid"], "first_name": r["fn"] or "", "username": r["un"] or ""}
    _SESS_CACHE[h] = (user, time.time()); return user

def session_revoke(raw):
    if not raw: return
    h = hashlib.sha256(raw.encode()).hexdigest(); _SESS_CACHE.pop(h, None)
    with db() as con: con.execute("DELETE FROM sessions WHERE token_hash=?", (h,))

def _png(size, rows):
    raw = b"".join(b"\x00" + bytes(r) for r in rows)
    def chunk(t, d): return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw, 6)) + chunk(b"IEND", b"")

def _fill_poly(rows, size, poly, color):
    ys = [p[1] for p in poly]; n = len(poly); col = bytes(color)
    for y in range(max(0, int(min(ys))), min(size - 1, int(max(ys))) + 1):
        yc = y + 0.5; xs = []
        for i in range(n):
            (xa, ya), (xb, yb) = poly[i], poly[(i + 1) % n]
            if (ya <= yc < yb) or (yb <= yc < ya): xs.append(xa + (yc - ya) * (xb - xa) / (yb - ya))
        xs.sort()
        for j in range(0, len(xs) - 1, 2):
            a, b = max(0, int(round(xs[j]))), min(size, int(round(xs[j + 1])))
            if b > a: rows[y][3 * a:3 * b] = col * (b - a)

def make_icon(size, maskable=False):
    """أيقونة التطبيق (قبعة تخرّج) تُرسم برمجياً، فلا حاجة لرفع ملفات صور."""
    top, bot = (124, 92, 255), (84, 58, 224); rows = []
    for y in range(size):
        t = y / (size - 1); rows.append(bytearray(bytes(int(top[i] + (bot[i] - top[i]) * t) for i in range(3)) * size))
    k = (0.62 if maskable else 0.82) * size; cx, cy = size / 2, size / 2 + 0.02 * size
    P = lambda dx, dy: (cx + dx * k, cy + dy * k)
    _fill_poly(rows, size, [P(-0.24, 0.02), P(0, 0.13), P(0.24, 0.02), P(0.24, 0.24), P(0, 0.36), P(-0.24, 0.24)], (226, 220, 255))
    _fill_poly(rows, size, [P(0, -0.30), P(0.42, -0.08), P(0, 0.14), P(-0.42, -0.08)], (255, 255, 255))
    _fill_poly(rows, size, [P(0.395, -0.075), P(0.425, -0.075), P(0.425, 0.20), P(0.395, 0.20)], (244, 182, 61))
    _fill_poly(rows, size, [P(0.38, 0.20), P(0.44, 0.20), P(0.46, 0.27), P(0.36, 0.27)], (244, 182, 61))
    return _png(size, rows)

ICONS = {"/icon-192.png": (192, False), "/icon-512.png": (512, False), "/icon-maskable.png": (512, True)}
_ICON_CACHE = {}
def icon_bytes(path):
    if path not in _ICON_CACHE: _ICON_CACHE[path] = make_icon(*ICONS[path])
    return _ICON_CACHE[path]

def manifest_dict():
    return {"name": "الأستاذ — مدرّسك الذكي", "short_name": "الأستاذ", "description": "مدرّسك الذكي للسادس الإعدادي: شرح وأسئلة وزارية وملازم واختبارات", "lang": "ar", "dir": "rtl",
            "id": "/", "start_url": "/", "scope": "/", "display": "standalone", "orientation": "any", "background_color": "#0f1628", "theme_color": "#0f1628",
            "icons": [{"src": "/icon-192.png", "sizes": "192x192", "type": "image/png", "purpose": "any"}, {"src": "/icon-512.png", "sizes": "512x512", "type": "image/png", "purpose": "any"},
                      {"src": "/icon-maskable.png", "sizes": "512x512", "type": "image/png", "purpose": "maskable"}]}

def sw_js():
    return "const V='" + app_version() + "';\n" + r"""
self.addEventListener('install', e => { self.skipWaiting(); e.waitUntil(caches.open('app-' + V).then(c => c.add('/')).catch(() => {})); });
self.addEventListener('activate', e => e.waitUntil((async () => { for (const k of await caches.keys()) if (k !== 'app-' + V) await caches.delete(k); await self.clients.claim(); })()));
self.addEventListener('fetch', e => {
  const r = e.request; if (r.method !== 'GET') return;
  const u = new URL(r.url); if (u.origin !== location.origin || u.pathname.startsWith('/api/')) return;
  e.respondWith(fetch(r).then(res => { if (res.ok && (r.mode === 'navigate' || u.pathname.startsWith('/static/'))) { const c = res.clone(); caches.open('app-' + V).then(x => x.put(r, c)); } return res; })
    .catch(() => caches.match(r).then(m => m || caches.match('/'))));
});
"""

def validate_init_data(init_data):
    """التحقق من توقيع تيليجرام (HMAC) وعمر الجلسة. يرجع dict المستخدم أو None."""
    if not init_data or len(init_data) > 4096 or not BOT_TOKEN: return None
    try: pairs = dict(parse_qsl(init_data, keep_blank_values=True))
    except Exception: return None
    got = pairs.pop("hash", None)
    if not got: return None
    check = "\n".join(f"{k}={v}" for k, v in sorted(pairs.items()))
    secret = hmac.new(b"WebAppData", BOT_TOKEN.encode(), hashlib.sha256).digest()
    if not hmac.compare_digest(hmac.new(secret, check.encode(), hashlib.sha256).hexdigest(), got): return None
    try:
        if time.time() - int(pairs.get("auth_date", "0")) > INITDATA_MAX_AGE: return None
        user = json.loads(pairs.get("user", "{}"))
        return user if isinstance(user, dict) and user.get("id") else None
    except Exception:
        return None

_IP_HITS, _API_HITS = defaultdict(deque), defaultdict(deque)

def hit_limited(store, key, limit, per=60):
    now = time.time(); dq = store[key]
    while dq and now - dq[0] > per: dq.popleft()
    if len(dq) >= limit: return True
    dq.append(now); return False

def err(msg, code=400):
    return {"error": msg, "_code": code}

def api_gate(uid, count=True):
    if db_is_banned(uid): return err("🚫 تم تقييد وصولك إلى البوت.", 403)
    ok, wait = check_rate(uid)
    if not ok: return err(f"انتظر {wait} ثانية.", 429)
    if count:
        ok, _ = check_daily(uid)
        if not ok: return err(f"وصلت حدّك اليومي ({daily_limit_for(uid)} طلب). ساهم بملزمة أو ملخص واستخدم البوت أكثر ليزيد حدّك حتى {DAILY_MAX}.", 429)
    return None

API_HEAVY = threading.BoundedSemaphore(max(4, HTTP_WORKERS // 2))

_HEAVY_HTTP = threading.BoundedSemaphore(max(4, HTTP_WORKERS // 2))

def heavy_api(uid, fn, count=True):
    """يشغّل عملاً ثقيلاً (نموذج ذكاء) بعد البوّابة وحجز مكان؛ أي فشل يُرجع رسالة نظيفة ويستردّ الطلب اليومي."""
    if not API_HEAVY.acquire(blocking=False):          # لا نسمح للطلبات الثقيلة بحجز كل خيوط الخادم
        return err("الضغط عالٍ حالياً، جرّب بعد لحظات.", 503)
    try: return _heavy_api(uid, fn, count)
    finally: API_HEAVY.release()

def _heavy_api(uid, fn, count):
    g = api_gate(uid, count)
    if g: return g
    if not _HEAVY_HTTP.acquire(blocking=False):
        if count: refund_daily(uid)
        return err("الضغط عالٍ حالياً على الأسئلة الذكية، جرّب بعد لحظات.", 503)
    try: return _heavy_inner(uid, fn, count)
    finally: _HEAVY_HTTP.release()

def _heavy_inner(uid, fn, count):
    with slot(uid, "api") as st:
        if st:
            if count: refund_daily(uid)
            return err("لسا عندك طلب جارٍ، انتظر لحظة." if st == 1 else "الضغط عالٍ حالياً، جرّب بعد لحظات.", 429 if st == 1 else 503)
        try: return fn()
        except NoSource as e:
            if count: refund_daily(uid)
            return err(str(e), 404)
        except Busy as e:
            if count: refund_daily(uid)
            return err(str(e), 503)
        except Exception as e:
            STATS["errors"] += 1
            if count: refund_daily(uid)
            log.warning("api heavy failed: %s", e)
            return err("تعذّر تنفيذ الطلب الآن، جرّب بعد قليل.", 502)

def _subj(body, allow_general=False):
    s = str(body.get("subject") or "")
    return s if s in SUBJECTS and (allow_general or s != "general") else ""

def _pct(a, b): return round(100 * a / b) if b else None

def api_init(uid, user, body):
    db_upsert_user(user, "miniapp"); STATS["webapp_visits"] += 1
    u = db_get_user(uid); a = activity_totals(uid)
    hist = [{"role": h["role"], "content": h["content"]} for h in db_history(uid, 0, 30)]
    subs = [{"key": k, "name": v["name"], "emoji": v["emoji"]} for k, v in SUBJECTS.items() if k != "general"]
    return {"role": ("owner" if is_owner(uid) else "admin" if is_admin(uid) else "student"), "user": {"id": uid, "name": (user.get("first_name") or "طالب"), "grade": u.get("grade") or "", "preferred_subject": u.get("preferred_subject") or "", "hide_lb": bool(u.get("hide_lb"))},
            "stats": {"used_today": db_usage_today(uid), "daily_limit": daily_limit_for(uid), "unlimited": is_owner(uid), "limit_max": DAILY_MAX, "limit_base": DAILY_LIMIT,
                      "streak": streak_of(uid), "title": rank_of(a["total_points"])["title"], "next_title": rank_of(a["total_points"])["next"],
                      "next_at": rank_of(a["total_points"])["next_at"], "from_at": rank_of(a["total_points"])["from"], "contrib": contrib_counts(uid), **a},
            "bot_username": BOT_USERNAME, "subjects": subs, "resume": get_ctx(uid), "perf": db_student_summary(uid), "history": hist,
            "insta": {"url": IG_URL, "handle": IG_HANDLE} if IG_URL else None}

def api_subject(uid, user, body):
    key = _subj(body)
    if not key: return err("مادة غير معروفة")
    name = subj_name(key); cid = int(body.get("chapter_id") or 0) if str(body.get("chapter_id") or "0").lstrip("-").isdigit() else 0
    stats = chapter_stats(uid, key)
    def info(i, title):
        s = stats.get(i)
        return {"id": i, "title": title, "pct": _pct(s[0], s[1]) if s else None, "expected": round(20 * s[0] / s[1]) if s else None, "tests": s[2] if s else 0}
    chs = chapters_of(key)
    chapters = [info(c["id"], c["title"]) for c in chs] or [info(0, "كل منهج " + name)]
    tot = [sum(v[0] for v in stats.values()), sum(v[1] for v in stats.values())]
    with db() as con:
        mats = [dict(r) for r in con.execute("SELECT id,title,telegram_url url,norm,description,mkind,teacher,part,batch,mtype FROM materials WHERE active=1 AND subject=? ORDER BY id DESC LIMIT 80", (key,))]
    filtered = False
    ch = chapter_get(cid) if cid else None
    if ch and ch["subject"] == key:
        ct = [t for t in tokens(ch["title"]) if not _gen(t)]
        sub = [r for r in mats if any(t in set(r["norm"].split()) or any(x.startswith(t) for x in r["norm"].split()) for t in ct)] if ct else []
        if sub: mats, filtered = sub, True
    mats = [{"id": r["id"], "title": r["title"], "url": mat_url({**r, "telegram_url": r["url"]}), "summary": is_summary(r), "meta": mat_meta(r), "desc": ((r.get("description") or "")[:140] if r.get("mkind") == "summary" else "")} for r in mats]
    src = sources_list(("book", "wazari"), key)
    return {"subject": key, "name": name, "emoji": SUBJECTS[key]["emoji"], "mastery": _pct(*tot), "chapters": chapters,
            "materials": mats, "filtered": filtered,
            "books": [{"id": s["id"], "title": s["title"], "pages": s["pages"]} for s in src if s["kind"] == "book"],
            "wazari": [{"id": s["id"], "title": s["title"], "year": s["year"], "round": s["round"]} for s in src if s["kind"] == "wazari"],
            "teachers": [{"id": t["id"], "name": t["name"], "youtube": t["youtube_url"], "notes": t["notes"]} for t in teachers_of(key)]}

def api_library(uid, user, body):
    q = str(body.get("q") or "").strip()[:100]
    out = {"materials": [], "books": [], "teachers": [], "chapters": []}
    qt = tokens(q)
    if len(q) < 2 or not qt: return out
    def has(text):
        t = tokens(text)
        return all(any(x.startswith(w) or w.startswith(x) for x in t) for w in qt)
    out["materials"] = [{"id": r["id"], "title": r["title"], "url": mat_url(r), "subject": subj_name(subject_key(r["subject"])), "summary": is_summary(r), "meta": mat_meta(r)} for r in search_materials(q, 12)]
    out["books"] = [{"id": s["id"], "title": s["title"], "kind": s["kind"], "year": s["year"], "subject": subj_name(s["subject"])} for s in sources_list(("book", "wazari")) if has(s["title"] + " " + subj_name(s["subject"]))][:8]
    out["teachers"] = [{"id": t["id"], "name": t["name"], "youtube": t["youtube_url"], "subject": subj_name(subject_key(t["subject"])), "key": subject_key(t["subject"])} for t in teachers_of() if has(t["name"])][:8]
    with db() as con: chs = [dict(r) for r in con.execute("SELECT id,subject,title FROM chapters ORDER BY subject,ord LIMIT 400")]
    out["chapters"] = [{"id": c["id"], "title": c["title"], "subject": c["subject"], "subject_name": subj_name(c["subject"])} for c in chs if has(c["title"])][:8]
    return out

def api_chat(uid, user, body):
    msg = str(body.get("message") or "").strip()[:2000]
    if not msg: return err("اكتب سؤالك أولاً")
    subject = _subj(body)
    def run():
        ans, hits = ask_teacher(uid, msg, subject, 0, "app", (user.get("first_name") or ""))
        STATS["requests"] += 1
        if subject: set_ctx(uid, subject)
        return {"answer": ans, "sources": fmt_sources(hits)}
    return heavy_api(uid, run)

def api_quiz(uid, user, body):
    subject = _subj(body)
    try: count = max(1, min(10, int(body.get("count") or 5)))
    except (TypeError, ValueError): count = 5
    qtype = body.get("type") if body.get("type") in ("mcq", "tf", "mixed") else "mixed"
    mode = body.get("mode") if body.get("mode") in ("normal", "wazari", "exam") else "normal"
    ids = body.get("chapter_ids") if isinstance(body.get("chapter_ids"), list) else ([body.get("chapter_id")] if body.get("chapter_id") else [])
    titles, valid = [], []
    for i in ids[:6]:
        c = chapter_get(i) if str(i).isdigit() and int(i) else None
        if c and c["subject"] == subject and c["id"] not in valid: valid.append(c["id"]); titles.append(c["title"])
    extra = str(body.get("topic") or "").strip()[:150]
    if extra: titles.append(extra)
    topic = "، ".join(titles)
    ch_id = valid[0] if len(valid) == 1 and not extra else 0
    sid = int(body["source_id"]) if str(body.get("source_id") or "").isdigit() else None
    def run():
        if sid: qs, info = make_quiz(subject, "", count, "mixed", source_id=sid)
        elif mode == "wazari": qs, info = make_quiz(subject, topic, count, qtype, wazari=True)
        elif mode == "exam":
            try: qs, info = make_quiz(subject, topic, count, qtype, wazari=True)
            except NoSource: qs, info = make_quiz(subject, topic, count, qtype)
        else: qs, info = make_quiz(subject, topic, count, qtype, topics=titles)
        if not qs: return err("ما قدرت أكوّن أسئلة، جرّب مرة ثانية.", 502)
        STATS["quizzes"] += 1
        if subject: set_ctx(uid, subject, ch_id, topic)
        return {"questions": qs, "kind": info["kind"], "grounded": info["grounded"], "chapter_id": ch_id, "title": topic}
    return heavy_api(uid, run)

def api_review(uid, user, body):
    subject = _subj(body)
    qs = mistakes_quiz(uid, subject, 5)
    if qs: return {"questions": qs, "kind": "review", "grounded": False, "chapter_id": 0, "title": "مراجعة أخطائك"}
    if not body.get("quick"): return err("ما عندك أخطاء محفوظة بهذه المادة. أي سؤال تغلط فيه يُحفظ هنا للمراجعة 👌", 404)
    if not subject:
        ctx = get_ctx(uid); subject = (ctx or {}).get("subject") or random.choice([k for k in SUBJECTS if k != "general"])
    return api_quiz(uid, user, {"subject": subject, "count": 5, "type": "mixed"})

def api_quiz_result(uid, user, body):
    if hit_limited(_API_HITS, ("qr", uid), 20, 3600): return err("حفظت نتائج كثيرة خلال ساعة، حاول لاحقاً.", 429)
    subject = _subj(body, True) or "general"
    try:
        score, total = int(body.get("score")), int(body.get("total"))
        cid = int(body.get("chapter_id") or 0)
    except (TypeError, ValueError): return err("بيانات غير صالحة")
    if not (0 <= score <= total <= 12 and total >= 1): return err("بيانات غير صالحة")
    kind = body.get("kind") if body.get("kind") in ("chapter", "wazari", "review", "exam", "normal") else "normal"
    with db() as con:
        con.execute("INSERT INTO quiz_results(user_id,subject,score,total,created_at,chapter_id,kind) VALUES(?,?,?,?,?,?,?)", (uid, subject, score, total, now_iso(), cid or None, kind))
    wrong = body.get("wrong") if isinstance(body.get("wrong"), list) else []
    save_mistakes(uid, subject, cid, [w for w in wrong if isinstance(w, dict)])
    fixed = body.get("fixed") if isinstance(body.get("fixed"), list) else []
    delete_mistakes(uid, fixed)
    touch_activity(uid, score + 2)
    gain = 0
    try: gain = game_on_quiz(uid, score, total)
    except Exception as e: log.warning("quiz stars: %s", e)
    out = {"ok": True, "stars_gain": gain}
    if gain:
        try: out["game"] = game_state(uid)
        except Exception: pass
    return out

def _sniff_image(b):
    if b[:3] == b"\xff\xd8\xff": return "image/jpeg"
    if b[:8] == b"\x89PNG\r\n\x1a\n": return "image/png"
    if b[:4] == b"RIFF" and b[8:12] == b"WEBP": return "image/webp"
    return None

def api_ocr(uid, user, body):
    raw = str(body.get("image") or "")
    if raw.startswith("data:"): raw = raw.partition(",")[2]
    if len(raw) > 6_500_000: return err("الصورة كبيرة. جرّب صورة أصغر.", 413)
    try: img = base64.b64decode(raw, validate=True)
    except Exception: return err("صورة غير صالحة")
    mime = _sniff_image(img)
    if not mime: return err("الصيغة المدعومة: JPG / PNG / WEBP")
    question = str(body.get("question") or "").strip()[:500]
    def run():
        try: text, _ = read_image(img, mime)
        except Exception as e:
            log.warning("api ocr: %s", e); return err("ما كدرت أقرأ الصورة، جرّب صورة أوضح.", 422)
        if len(text.strip()) < 8: return err("ما قرأت نصاً واضحاً بالصورة، جرّب صورة أوضح وأقرب.", 422)
        STATS["photos"] += 1
        q = (question + "\n" if question else "") + text[:6000]
        ans, hits = ask_teacher(uid, q, _subj(body) or guess_subject(q) or "", 0, "app", user.get("first_name") or "", shown=("[صورة] " + question).strip(), from_image=True, max_tokens=1800)
        return {"answer": ans, "sources": fmt_sources(hits)}       # النص المقروء لا يُرجَع أبداً
    return heavy_api(uid, run)

def api_opinion(uid, user, body):
    kind = body.get("kind"); rid = body.get("id")
    if kind == "teacher" and str(rid).isdigit():
        with db() as con: r = con.execute("SELECT name,subject,notes FROM teachers WHERE id=? AND active=1", (int(rid),)).fetchone()
        if not r: return err("المدرّس غير موجود", 404)
        key, subj, target, extra = f"t:{rid}", subject_key(r["subject"]), f"مدرّس يوتيوب: {r['name']}", r["notes"] or ""
    elif kind == "handout" and str(rid).isdigit():
        with db() as con: r = con.execute("SELECT title,subject FROM materials WHERE id=? AND active=1", (int(rid),)).fetchone()
        if not r: return err("الملزمة غير موجودة", 404)
        key, subj, target, extra = f"m:{rid}", subject_key(r["subject"]), f"ملزمة: {r['title']}", ""
    elif kind == "youtube" and _subj(body):
        subj = _subj(body); names = "، ".join(t["name"] for t in teachers_of(subj)[:8])
        key, target, extra = "y:" + subj, f"أفضل مدرّسي يوتيوب لمادة {subj_name(subj)} للسادس الإعدادي في العراق وكيف يختار الطالب بينهم", (f"مدرّسون مسجّلون عندنا: {names}" if names else "")
    else:
        return err("طلب غير صالح")
    cached = opinion_get(key)
    if cached: return {"text": cached}
    return heavy_api(uid, lambda: {"text": web_opinion(key, subj, target, extra)})

_CH_LOCK = threading.Lock()

def get_daily_challenge():
    day = today_key()
    def load():
        with db() as con: r = con.execute("SELECT subject,data FROM daily_challenge WHERE day=?", (day,)).fetchone()
        return {"subject": r["subject"], "questions": json.loads(r["data"])} if r else None
    c = load()
    if c: return c
    if not API_HEAVY.acquire(blocking=False): raise Busy("جاري تجهيز تحدّي اليوم، جرّب بعد لحظات")
    try:
        if not _CH_LOCK.acquire(timeout=80): raise Busy("جاري تجهيز تحدّي اليوم، جرّب بعد لحظات")
        try:
            return _gen_challenge(day, load)
        finally:
            _CH_LOCK.release()
    finally:
        API_HEAVY.release()

def _gen_challenge(day, load):
    if True:
        c = load()
        if c: return c
        with db() as con: have = sorted({r["subject"] for r in con.execute("SELECT DISTINCT subject FROM sources WHERE kind='book' AND subject!='general'")})
        subject = random.Random(day).choice(have or ["math", "physics", "chemistry", "biology"])
        qs, _ = make_quiz(subject, "", 5, "mcq")
        if len(qs) < 3: raise RuntimeError("challenge generation failed")
        with db() as con: con.execute("INSERT OR REPLACE INTO daily_challenge(day,subject,data) VALUES(?,?,?)", (day, subject, json.dumps(qs, ensure_ascii=False)))
        return {"subject": subject, "questions": qs}

def api_challenge(uid, user, body):
    try: c = get_daily_challenge()
    except Busy as e: return err(str(e), 503)
    except Exception as e:
        log.warning("challenge: %s", e); return err("تعذّر تجهيز تحدّي اليوم الآن، جرّب بعد قليل.", 502)
    with db() as con: d = con.execute("SELECT score,total FROM challenge_done WHERE user_id=? AND day=?", (uid, today_key())).fetchone()
    s = c["subject"]
    return {"day": today_key(), "subject": s, "name": subj_name(s), "emoji": SUBJECTS[s]["emoji"], "questions": c["questions"], "done": dict(d) if d else None}

def api_challenge_result(uid, user, body):
    try: score, total = int(body.get("score")), int(body.get("total"))
    except (TypeError, ValueError): return err("بيانات غير صالحة")
    if not (0 <= score <= total <= 10 and total >= 1): return err("بيانات غير صالحة")
    with db() as con:
        n = con.execute("INSERT OR IGNORE INTO challenge_done(user_id,day,score,total) VALUES(?,?,?,?)", (uid, today_key(), score, total)).rowcount
        if n: con.execute("INSERT INTO quiz_results(user_id,subject,score,total,created_at,kind) VALUES(?,?,?,?,?,'challenge')", (uid, body.get("subject") if body.get("subject") in SUBJECTS else "general", score, total, now_iso()))
    if n: touch_activity(uid, score * 2 + 3)
    return {"ok": True, "counted": bool(n)}


# ── المتصدّرون: دراسة يومي/أسبوعي، نقاط، كوينز ──
LB_KINDS = {"study_day": ("دراسة اليوم", "د"), "study_week": ("دراسة الأسبوع", "د"), "points": ("النقاط", "نقطة"), "coins": ("الكوينز", "كوينز")}
def lb_rows(kind, limit=30):
    wk = (datetime.now() - timedelta(days=6)).strftime("%Y-%m-%d")
    ok = "u.hide_lb=0 AND u.banned=0"
    with db() as con:
        if kind == "study_day":
            q, a = f"SELECT a.user_id uid, a.focus_min v, u.first_name n FROM activity a JOIN users u ON u.user_id=a.user_id WHERE a.day=? AND a.focus_min>0 AND {ok} ORDER BY v DESC LIMIT ?", (today_key(), limit)
        elif kind == "study_week":
            q, a = f"SELECT a.user_id uid, SUM(a.focus_min) v, u.first_name n FROM activity a JOIN users u ON u.user_id=a.user_id WHERE a.day>=? AND {ok} GROUP BY a.user_id HAVING v>0 ORDER BY v DESC LIMIT ?", (wk, limit)
        elif kind == "points":
            q, a = f"SELECT a.user_id uid, SUM(a.points) v, u.first_name n FROM activity a JOIN users u ON u.user_id=a.user_id WHERE {ok} GROUP BY a.user_id HAVING v>0 ORDER BY v DESC LIMIT ?", (limit,)
        else:
            q, a = f"SELECT g.user_id uid, g.earned v, u.first_name n FROM game g JOIN users u ON u.user_id=g.user_id WHERE g.earned>0 AND {ok} ORDER BY v DESC LIMIT ?", (limit,)
        return [dict(r) for r in con.execute(q, a)]

def api_lb(uid, user, body):
    kind = str(body.get("kind") or "study_day")
    if kind not in LB_KINDS: return err("نوع غير معروف")
    if not _rate(("lb", uid), 40, 60): return err("بطّئ شوي", 429)
    rows = lb_rows(kind, 30)
    out = [{"rank": i, "name": (r["n"] or "طالب")[:18], "v": int(r["v"] or 0), "me": r["uid"] == int(uid)} for i, r in enumerate(rows, 1)]
    me = next((x for x in out if x["me"]), None)
    mine = None
    if not me:
        allr = lb_rows(kind, 1000)
        for i, r in enumerate(allr, 1):
            if r["uid"] == int(uid): mine = {"rank": i, "name": "أنت", "v": int(r["v"] or 0), "me": True}; break
    return {"kind": kind, "title": LB_KINDS[kind][0], "unit": LB_KINDS[kind][1], "rows": out[:20], "mine": me or mine}

def api_board(uid, user, body):
    return {"board": leaderboard(uid), "week_points": activity_totals(uid)["week_points"]}

def api_focus(uid, user, body):
    try: mins = int(body.get("minutes"))
    except (TypeError, ValueError): return err("بيانات غير صالحة")
    if not 1 <= mins <= 240: return err("المدة بين 1 و240 دقيقة")
    if hit_limited(_API_HITS, ("focus", uid), 12, 3600): return err("سجّلت جلسات كثيرة خلال ساعة.", 429)
    sub = _subj(body)
    focus_log(uid, mins, sub, str(body.get("note") or "")); touch_activity(uid, max(1, mins // 5), mins); _LIM.pop(uid, None)
    return {"ok": True, **focus_stats(uid)}

def api_focus_stats(uid, user, body):
    return focus_stats(uid)

def api_focus_goal(uid, user, body):
    g = _to_int(body.get("goal"), 10, 600)
    if not g: return err("الهدف بين 10 و600 دقيقة")
    with db() as con: con.execute("UPDATE users SET focus_goal=? WHERE user_id=?", (g, uid))
    return {"ok": True, "goal": g}

def api_player(uid, user, body):
    t, pos = player_by_ref(str(body.get("ref") or "")[:20])
    if not t: return err("هذا اللاعب غير متاح للعرض", 404)
    p = player_profile(t); p["rank"] = pos
    return p

def api_profile(uid, user, body):
    grade = str(body.get("grade") or "").strip()[:30]
    pref = body.get("preferred_subject") if body.get("preferred_subject") in SUBJECTS else ""
    with db() as con:
        con.execute("UPDATE users SET grade=?, preferred_subject=?, hide_lb=? WHERE user_id=?", (grade, pref, 1 if body.get("hide_lb") else 0, uid))
    return {"ok": True}

def api_clear(uid, user, body):
    db_clear_history(uid, 0); db_event(uid, "history_cleared", "miniapp")
    return {"ok": True}

ROUTES_GET = {"/api/init": api_init, "/api/board": api_board}
ROUTES_POST = {"/api/subject": api_subject, "/api/library": api_library, "/api/chat": api_chat, "/api/quiz": api_quiz, "/api/review": api_review,
               "/api/quiz_result": api_quiz_result, "/api/ocr": api_ocr, "/api/opinion": api_opinion, "/api/challenge": api_challenge,
               "/api/challenge_result": api_challenge_result, "/api/focus": api_focus, "/api/profile": api_profile, "/api/clear": api_clear}

# ── لوحة الإدارة داخل التطبيق (للمشرفين) ──
_UP_JOBS = {}
_UPLOAD_SEM = threading.BoundedSemaphore(2)

def _need_admin(uid):
    return None if is_admin(uid) else err("هذه الصفحة للمشرفين فقط.", 403)

def _need_owner(uid):
    return None if is_owner(uid) else err("هذا الإجراء للمالك فقط.", 403)

def _to_int(v, lo, hi):
    try: x = int(v)
    except (TypeError, ValueError): return None
    return x if lo <= x <= hi else None

def api_admin_overview(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    with db() as con:
        counts = {"users": con.execute("SELECT COUNT(*) n FROM users").fetchone()["n"],
                  "materials": con.execute("SELECT COUNT(*) n FROM materials WHERE active=1").fetchone()["n"],
                  "chapters": con.execute("SELECT COUNT(*) n FROM chapters").fetchone()["n"],
                  "teachers": con.execute("SELECT COUNT(*) n FROM teachers WHERE active=1").fetchone()["n"]}
        chs = [dict(r) for r in con.execute("SELECT id,subject,title FROM chapters ORDER BY subject,ord,id")]
        mats = [dict(r) for r in con.execute("SELECT id,title,subject,telegram_url url FROM materials WHERE active=1 ORDER BY id DESC LIMIT 60")]
    out = {"is_owner": is_owner(uid), "counts": counts, "chapters": chs, "materials": mats,
           "sources": [{"id": r["id"], "title": r["title"], "subject": r["subject"], "kind": r["kind"], "pages": r["pages"], "chunks": r["chunks"], "year": r["year"], "round": r["round"]} for r in sources_list(("book", "wazari"))],
           "teachers": [{"id": t["id"], "name": t["name"], "subject": subject_key(t["subject"]), "youtube": t["youtube_url"], "notes": t["notes"]} for t in teachers_of()],
           "max_upload_mb": MAX_UPLOAD // (1024 * 1024), "ocr_pages": OCR_MAX_PAGES}
    if is_owner(uid): out["admins"] = admin_list()
    return out

def api_admin_chapter(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    sub, title = _subj(body), str(body.get("title") or "").strip()[:120]
    if not sub or not title: return err("اختر المادة واكتب عنوان الفصل")
    order = len(chapters_of(sub)) + 1
    with db() as con: con.execute("INSERT INTO chapters(subject,title,ord) VALUES(?,?,?) ON CONFLICT(subject,title) DO NOTHING", (sub, title, order))
    return {"ok": True}

def api_admin_chapter_del(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    with db() as con: con.execute("DELETE FROM chapters WHERE id=?", (_to_int(body.get("id"), 1, 10**9) or 0,))
    return {"ok": True}

def api_admin_autochapters(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    sub = _subj(body)
    if not sub: return err("اختر المادة")
    def run():
        try: titles = autochapters(sub)
        except RuntimeError as e: return err(str(e), 422)
        return {"ok": True, "count": len(titles), "titles": titles}
    return heavy_api(uid, run, count=False)

def api_admin_teacher(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    name, sub = str(body.get("name") or "").strip()[:80], _subj(body)
    url, notes = str(body.get("youtube") or "").strip()[:300], str(body.get("notes") or "").strip()[:400]
    if not name or not sub: return err("اكتب اسم المدرّس واختر المادة")
    if url and not url.startswith("https://"): return err("رابط القناة يجب أن يبدأ بـ https://")
    with db() as con: con.execute("INSERT INTO teachers(name,subject,youtube_url,notes,created_at) VALUES(?,?,?,?,?)", (name, sub, url, notes, now_iso()))
    return {"ok": True}

def api_admin_teacher_del(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    with db() as con: con.execute("UPDATE teachers SET active=0 WHERE id=?", (_to_int(body.get("id"), 1, 10**9) or 0,))
    return {"ok": True}

def api_admin_material(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    title, sub, url = str(body.get("title") or "").strip()[:120], _subj(body), str(body.get("url") or "").strip()[:300]
    if not title or not sub: return err("اكتب عنوان الملزمة واختر المادة")
    if not url.startswith(("https://t.me/", "https://telegram.me/")): return err("الرابط يجب أن يكون رابط تيليجرام (https://t.me/…)")
    material_add(title, sub, url, by=uid)
    return {"ok": True}

def api_admin_material_del(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    with db() as con: con.execute("UPDATE materials SET active=0 WHERE id=?", (_to_int(body.get("id"), 1, 10**9) or 0,))
    return {"ok": True}

def api_admin_source_del(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    sid = _to_int(body.get("id"), 1, 10**9)
    return {"ok": True} if sid and source_delete(sid) else err("المصدر غير موجود", 404)

def api_admin_admin_add(uid, user, body):
    g = _need_owner(uid)
    if g: return g
    t = _to_int(body.get("user_id"), 1, 10**15)
    if not t: return err("اكتب رقم المستخدم (رقم تيليجرام) بالأرقام فقط")
    if t in OWNER_IDS: return err("هذا الحساب مالك أصلاً")
    admin_add(t, uid); db_event(t, "admin_added", uid)
    try: say(t, "تمت إضافتك مشرفاً في المدرّس الذكي. افتح التطبيق من حسابي ثم لوحة الإدارة.")
    except Exception: pass
    return {"ok": True}

def api_admin_admin_del(uid, user, body):
    g = _need_owner(uid)
    if g: return g
    t = _to_int(body.get("user_id"), 1, 10**15)
    return {"ok": True} if t and admin_del(t) else err("هذا الحساب ليس مشرفاً", 404)

def api_admin_job(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    j = _UP_JOBS.get(str(body.get("id") or ""))
    if not j or (j["uid"] != uid and not is_owner(uid)): return err("المهمة غير موجودة", 404)
    return {k: j[k] for k in ("status", "pct", "msg") } | ({"result": j["result"]} if j.get("result") else {})

def validate_upload_meta(m):
    if not isinstance(m, dict): return None, "بيانات غير صالحة"
    fn = str(m.get("filename") or "")[:200]
    if Path(fn).suffix.lower() not in (".pdf", ".docx", ".txt"): return None, "المدعوم: PDF / DOCX / TXT"
    kind = m.get("kind") if m.get("kind") in ("book", "wazari") else None
    sub = m.get("subject") if m.get("subject") in SUBJECTS else None
    title = str(m.get("title") or "").strip()[:120]
    if not kind or not sub or not title: return None, "اختر النوع والمادة واكتب العنوان"
    out = {"filename": fn, "kind": kind, "subject": sub, "title": title, "year": _to_int(m.get("year"), 1990, 2035),
           "round": (str(m.get("round") or "")[:30] or None), "append_to": None, "start_page": _to_int(m.get("start_page"), 1, 5000)}
    a = _to_int(m.get("append_to"), 1, 10**9)
    if a:
        with db() as con: r = con.execute("SELECT kind FROM sources WHERE id=?", (a,)).fetchone()
        if not r or r["kind"] != kind: return None, "الملف الذي اخترته للإضافة إليه غير موجود"
        out["append_to"] = a
    return out, None

def start_upload_job(uid, raw, meta):
    st, key = claim(uid, "ingest")
    if st: return err("عندك رفع جارٍ، انتظر حتى ينتهي." if st == 1 else "النظام مشغول، جرّب بعد قليل.", 429 if st == 1 else 503)
    jid = secrets.token_hex(6)
    _UP_JOBS[jid] = {"uid": uid, "status": "run", "pct": 0, "msg": "جاري قراءة الملف…", "t": time.time()}
    def job():
        j = _UP_JOBS[jid]
        try:
            prog = lambda i, n: j.update(pct=int(100 * i / max(n, 1)), msg=f"قراءة الصفحات الممسوحة {i} من {n}")
            res = ingest_bytes(raw, meta["filename"], meta["title"], meta["subject"], meta["kind"], uid, meta["year"], meta["round"], prog, meta["append_to"], meta["start_page"])
            db_event(uid, f"{meta['kind']}_added", f"{meta['title']}:{res['chunks']}")
            j.update(status="done", pct=100, msg="تمت الإضافة", result={"pages": res["pages"], "chunks": res["chunks"], "notes": res["notes"], "appended": bool(res.get("appended"))})
        except Exception as e:
            STATS["errors"] += 1; log.warning("upload job: %s", e)
            j.update(status="error", msg=str(e)[:250])
    submit(key, job)
    return {"job": jid}

ROUTES_GET.update({"/api/admin/overview": api_admin_overview, "/api/admin/job": api_admin_job})
ROUTES_POST.update({"/api/admin/chapter": api_admin_chapter, "/api/admin/chapter_del": api_admin_chapter_del, "/api/admin/autochapters": api_admin_autochapters,
                    "/api/admin/teacher": api_admin_teacher, "/api/admin/teacher_del": api_admin_teacher_del,
                    "/api/admin/material": api_admin_material, "/api/admin/material_del": api_admin_material_del,
                    "/api/admin/source_del": api_admin_source_del, "/api/admin/admin_add": api_admin_admin_add, "/api/admin/admin_del": api_admin_admin_del})
ROUTES_GET.update({"/api/player": api_player, "/api/focus_stats": api_focus_stats})
ROUTES_POST.update({"/api/focus_goal": api_focus_goal})


# ══════════════════════════════════════════════════════════════════
#        النجوم + المزرعة + المتجر + تيوب (محاضرات المدرّسين)
# ══════════════════════════════════════════════════════════════════

CROPS = [("potato", "بطاطا", 25, 15), ("carrot", "جزر", 50, 25), ("tomato", "طماطم", 75, 40), ("watermelon", "بطيخ", 100, 50), ("corn", "ذرة", 120, 60),
         ("kiwi", "كيوي", 150, 75), ("strawberry", "فراولة", 250, 125), ("cherry", "كرز", 350, 175), ("banana", "موز", 400, 200), ("lemon", "ليمون", 500, 250),
         ("grape", "عنب", 600, 300), ("orange", "برتقال", 700, 350), ("blueberry", "توت", 800, 400), ("peach", "خوخ", 900, 450), ("apple", "تفاح", 1000, 500),
         ("pineapple", "أناناس", 1200, 600), ("pumpkin", "قرع", 1500, 750)]
CROP = {k: {"key": k, "name": n, "cost": c, "need": m, "sell": int(c * 1.6)} for k, n, c, m in CROPS}
FREE_THEMES = ("dark", "light", "emerald", "sunset")
THEME_SHOP = {"sky": ("سماوي", 150), "lime": ("ليموني", 150), "rose": ("وردي", 200), "royal": ("ملكي", 400), "gold": ("ذهبي", 800)}
STICKER_SHOP = {"star": ("نجمة", 50), "moon": ("هلال", 100), "book": ("كتاب", 100), "cloud": ("غيمة", 150), "heart": ("قلب", 150),
                "cat": ("قطة", 250), "rocket": ("صاروخ", 300), "crown": ("تاج", 600)}
MAX_PLOTS = 9
START_STARS = 50

def game_init():
    with db() as con:
        con.executescript("""
            CREATE TABLE IF NOT EXISTS game(user_id INTEGER PRIMARY KEY, stars INTEGER NOT NULL DEFAULT 50, plots INTEGER NOT NULL DEFAULT 1,
                theme TEXT NOT NULL DEFAULT 'dark', sticker TEXT NOT NULL DEFAULT '', owned TEXT NOT NULL DEFAULT '[]', earned INTEGER NOT NULL DEFAULT 0);
            CREATE TABLE IF NOT EXISTS farm(user_id INTEGER NOT NULL, plot INTEGER NOT NULL, crop TEXT NOT NULL, progress INTEGER NOT NULL DEFAULT 0,
                planted_at TEXT NOT NULL, PRIMARY KEY(user_id, plot));
            CREATE TABLE IF NOT EXISTS yt_follow(user_id INTEGER NOT NULL, teacher_id INTEGER NOT NULL, PRIMARY KEY(user_id, teacher_id));
            CREATE TABLE IF NOT EXISTS yt_sub(user_id INTEGER NOT NULL, cid TEXT NOT NULL, title TEXT NOT NULL DEFAULT '', added_at TEXT NOT NULL DEFAULT '', PRIMARY KEY(user_id, cid));
            CREATE TABLE IF NOT EXISTS house(user_id INTEGER PRIMARY KEY, name TEXT NOT NULL DEFAULT 'بيتي', owned TEXT NOT NULL DEFAULT '[]', placed TEXT NOT NULL DEFAULT '{}',
                chars TEXT NOT NULL DEFAULT '[]', created_at TEXT NOT NULL DEFAULT '');
            CREATE TABLE IF NOT EXISTS world(user_id INTEGER PRIMARY KEY, objs TEXT NOT NULL DEFAULT '{}', lands TEXT NOT NULL DEFAULT '[]', cleared TEXT NOT NULL DEFAULT '[]', bx INTEGER NOT NULL DEFAULT 50, by INTEGER NOT NULL DEFAULT 50, nid INTEGER NOT NULL DEFAULT 1);
            CREATE TABLE IF NOT EXISTS town(user_id INTEGER PRIMARY KEY, items TEXT NOT NULL DEFAULT '{}', bx INTEGER NOT NULL DEFAULT 4, by INTEGER NOT NULL DEFAULT 5, ready TEXT NOT NULL DEFAULT '{}');
            CREATE TABLE IF NOT EXISTS house_invites(id INTEGER PRIMARY KEY AUTOINCREMENT, from_id INTEGER NOT NULL, to_id INTEGER NOT NULL, created_at REAL NOT NULL);
            CREATE TABLE IF NOT EXISTS friends(user_id INTEGER NOT NULL, friend_id INTEGER NOT NULL, PRIMARY KEY(user_id, friend_id));
            CREATE TABLE IF NOT EXISTS rooms(code TEXT PRIMARY KEY, name TEXT NOT NULL, owner INTEGER NOT NULL, created_at TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS room_members(code TEXT NOT NULL, user_id INTEGER NOT NULL, last_seen REAL NOT NULL DEFAULT 0, studying INTEGER NOT NULL DEFAULT 0,
                since REAL NOT NULL DEFAULT 0, PRIMARY KEY(code, user_id));
        """)
        gc = {r["name"] for r in con.execute("PRAGMA table_info(game)")}
        if "fcode" not in gc: con.execute("ALTER TABLE game ADD COLUMN fcode TEXT DEFAULT ''")
        if "qday" not in gc: con.execute("ALTER TABLE game ADD COLUMN qday TEXT DEFAULT ''")
        if "qstars" not in gc: con.execute("ALTER TABLE game ADD COLUMN qstars INTEGER DEFAULT 0")
        if "channel_id" not in {r["name"] for r in con.execute("PRAGMA table_info(teachers)")}:
            con.execute("ALTER TABLE teachers ADD COLUMN channel_id TEXT DEFAULT ''")

def game_row(uid):
    with db() as con:
        con.execute("INSERT OR IGNORE INTO game(user_id,stars) VALUES(?,?)", (int(uid), START_STARS))
        return dict(con.execute("SELECT * FROM game WHERE user_id=?", (int(uid),)).fetchone())

def game_owned(row):
    try: return [x for x in json.loads(row.get("owned") or "[]") if isinstance(x, str)]
    except Exception: return []

def game_on_focus(uid, mins):
    """كل دقيقة تركيز = نجمة (+20% للجلسات 50 دقيقة فأكثر) وتسقي كل نباتات المزرعة."""
    game_row(uid)
    gain = mins + (mins // 5 if mins >= 50 else 0)
    with db() as con:
        con.execute("UPDATE game SET stars=stars+?, earned=earned+? WHERE user_id=?", (gain, gain, int(uid)))
        con.execute("UPDATE farm SET progress=progress+? WHERE user_id=?", (int(mins), int(uid)))
    return gain

QUIZ_DAILY_CAP = 120
def game_on_quiz(uid, score, total):
    """نجمتان لكل إجابة صحيحة (+5 للعلامة الكاملة بـ5 أسئلة فأكثر)، بسقف يومي حتى لا تُستغل."""
    g = game_row(uid); day = today_key()
    used = g.get("qstars") or 0 if g.get("qday") == day else 0
    gain = min(max(0, score * 2 + (5 if total >= 5 and score == total else 0)), max(0, QUIZ_DAILY_CAP - used))
    if gain <= 0: return 0
    with db() as con:
        con.execute("UPDATE game SET stars=stars+?, earned=earned+?, qday=?, qstars=? WHERE user_id=?", (gain, gain, day, used + gain, int(uid)))
    return gain

def game_state(uid):
    g = game_row(uid)
    with db() as con:
        plots = [dict(r) for r in con.execute("SELECT plot,crop,progress FROM farm WHERE user_id=? ORDER BY plot", (int(uid),))]
    farm = []
    for r in plots:
        c = CROP.get(r["crop"])
        if c: farm.append({"plot": r["plot"], "crop": r["crop"], "progress": min(r["progress"], c["need"]), "need": c["need"], "ready": r["progress"] >= c["need"], "sell": c["sell"]})
    owned = game_owned(g)
    return {"stars": g["stars"], "earned": g["earned"], "plots": g["plots"], "max_plots": MAX_PLOTS, "unlock_cost": (300 * g["plots"]) if g["plots"] < MAX_PLOTS else 0,
            "theme": g["theme"], "sticker": g["sticker"], "owned": owned, "farm": farm,
            "crops": list(CROP.values()),
            "themes": [{"key": k, "name": n, "price": p, "owned": k in owned} for k, (n, p) in THEME_SHOP.items()],
            "stickers": [{"key": k, "name": n, "price": p, "owned": k in owned} for k, (n, p) in STICKER_SHOP.items()]}

def _spend(uid, cost):
    with db() as con:
        return con.execute("UPDATE game SET stars=stars-? WHERE user_id=? AND stars>=?", (int(cost), int(uid), int(cost))).rowcount == 1

def api_game(uid, user, body):
    return game_state(uid)

def api_farm(uid, user, body):
    act = str(body.get("action") or "")
    g = game_row(uid)
    if act == "unlock":
        if g["plots"] >= MAX_PLOTS: return err("كل المربعات مفتوحة")
        cost = 300 * g["plots"]
        if not _spend(uid, cost): return err("كوينزاتك لا تكفي لفتح مربع جديد")
        with db() as con: con.execute("UPDATE game SET plots=plots+1 WHERE user_id=?", (int(uid),))
    elif act in ("plant", "harvest"):
        try: plot = int(body.get("plot"))
        except (TypeError, ValueError): return err("مربع غير صالح")
        if not 0 <= plot < g["plots"]: return err("هذا المربع مقفل")
        with db() as con: cur = con.execute("SELECT crop,progress FROM farm WHERE user_id=? AND plot=?", (int(uid), plot)).fetchone()
        if act == "plant":
            c = CROP.get(str(body.get("crop") or ""))
            if cur: return err("المربع مزروع مسبقاً")
            if not c: return err("محصول غير معروف")
            if not _spend(uid, c["cost"]): return err("كوينزاتك لا تكفي لهذا المحصول")
            with db() as con: con.execute("INSERT OR REPLACE INTO farm(user_id,plot,crop,progress,planted_at) VALUES(?,?,?,0,?)", (int(uid), plot, c["key"], now_iso()))
        else:
            if not cur: return err("المربع فارغ")
            c = CROP.get(cur["crop"])
            if not c or cur["progress"] < c["need"]: return err("المحصول لم ينضج بعد، يحتاج مزيداً من ساعات الدراسة")
            with db() as con:
                if con.execute("DELETE FROM farm WHERE user_id=? AND plot=?", (int(uid), plot)).rowcount:
                    con.execute("UPDATE game SET stars=stars+? WHERE user_id=?", (c["sell"], int(uid)))
    else:
        return err("طلب غير معروف")
    return game_state(uid)

def api_shop(uid, user, body):
    act, kind, key = str(body.get("action") or ""), str(body.get("kind") or ""), str(body.get("key") or "")
    g = game_row(uid); owned = game_owned(g)
    cat = THEME_SHOP if kind == "theme" else STICKER_SHOP if kind == "sticker" else None
    if cat is None: return err("طلب غير صالح")
    if act == "buy":
        if key not in cat: return err("عنصر غير موجود")
        if key in owned: return err("تملكه مسبقاً")
        if not _spend(uid, cat[key][1]): return err("كوينزاتك لا تكفي")
        owned.append(key)
        with db() as con: con.execute("UPDATE game SET owned=? WHERE user_id=?", (json.dumps(owned), int(uid)))
    elif act == "equip":
        col = "theme" if kind == "theme" else "sticker"
        if kind == "theme" and key not in FREE_THEMES and key not in owned: return err("اشترِ الثيم أولاً")
        if kind == "sticker" and key and key not in owned: return err("اشترِ الملصق أولاً")
        if kind == "theme" and key not in FREE_THEMES and key not in THEME_SHOP: return err("ثيم غير موجود")
        with db() as con: con.execute(f"UPDATE game SET {col}=? WHERE user_id=?", (key, int(uid)))
    else:
        return err("طلب غير صالح")
    return game_state(uid)

# ── تيوب: محاضرات من قنوات المدرّسين عبر RSS الرسمي (بلا مفتاح API) ──
YT_TTL = 1800
_YT = {}; _YT_BAD = {}; _YT_META = {}; _YT_L = threading.Lock()
_YT_HDR = {"User-Agent": "Mozilla/5.0 (Linux; Android 12) AppleWebKit/537.36 Chrome/120 Mobile Safari/537.36", "Accept-Language": "ar,en;q=0.8", "Cookie": "CONSENT=YES+1; SOCS=CAI"}

def yt_channel_id(t):
    cid = (t.get("channel_id") or "").strip()
    if cid: return cid
    url = (t.get("youtube_url") or "").strip()
    m = re.search(r"/channel/(UC[\w-]{20,})", url)
    if not m and (url.startswith("https://") and "youtube.com" in url or "youtu.be" in url):
        if time.time() - _YT_BAD.get(url, 0) < 3600: return ""
        try:
            r = requests.get(url, headers=_YT_HDR, timeout=10)
            m = (re.search(r'"channelId":"(UC[\w-]{22})"', r.text) or re.search(r'channel_id=(UC[\w-]{22})', r.text)
                 or re.search(r'itemprop="(?:identifier|channelId)" content="(UC[\w-]{22})"', r.text) or re.search(r'/channel/(UC[\w-]{22})', r.text))
        except Exception as e:
            log.warning("yt resolve: %s", e); m = None
    if not m:
        _YT_BAD[url] = time.time(); return ""
    cid = m.group(1)
    try:
        with db() as con: con.execute("UPDATE teachers SET channel_id=? WHERE id=?", (cid, t["id"]))
    except Exception: pass
    return cid

def yt_fetch(cid):
    """آخر فيديوهات القناة (بالكاش 30 دقيقة، وعند الفشل نُبقي القديم)."""
    now = time.time()
    with _YT_L: hit = _YT.get(cid)
    if hit and now - hit[0] < YT_TTL: return hit[1]
    try:
        import xml.etree.ElementTree as ET
        r = requests.get("https://www.youtube.com/feeds/videos.xml", params={"channel_id": cid}, headers=_YT_HDR, timeout=10)
        r.raise_for_status()
        ns = {"a": "http://www.w3.org/2005/Atom", "yt": "http://www.youtube.com/xml/schemas/2015"}
        root = ET.fromstring(r.content); out = []
        _YT_META[cid] = (root.findtext("a:title", "", ns) or "").strip()[:80]
        for e in root.findall("a:entry", ns)[:15]:
            vid = (e.findtext("yt:videoId", "", ns) or "").strip()
            if not re.fullmatch(r"[\w-]{11}", vid): continue
            out.append({"id": vid, "title": (e.findtext("a:title", "", ns) or "").strip()[:160], "published": (e.findtext("a:published", "", ns) or "")[:19]})
        with _YT_L: _YT[cid] = (now, out)
        return out
    except Exception as e:
        log.warning("yt feed %s: %s", cid, str(e)[:80])
        with _YT_L: _YT[cid] = (now - YT_TTL + 120, hit[1] if hit else [])
        return hit[1] if hit else []

def yt_teachers():
    with db() as con:
        return [dict(r) for r in con.execute("SELECT id,name,subject,youtube_url,COALESCE(channel_id,'') channel_id FROM teachers WHERE active=1 AND youtube_url<>'' ORDER BY id")]

def api_yt_home(uid, user, body):
    mode = "mine" if body.get("mode") == "mine" else "all"
    q = ar_norm(str(body.get("q") or "")[:60]); subj = _subj(body)
    ts = yt_teachers()
    with db() as con: fol = {r["teacher_id"] for r in con.execute("SELECT teacher_id FROM yt_follow WHERE user_id=?", (int(uid),))}
    pool = [t for t in ts if (t["id"] in fol if mode == "mine" else True) and (not subj or subject_key(t["subject"]) == subj)]
    pool.sort(key=lambda t: (t["id"] not in fol, t["id"]))
    pool = pool[:14]
    def load(t):
        cid = yt_channel_id(t); t["channel_id"] = cid
        return t, (yt_fetch(cid) if cid else [])
    feed = []
    if pool:
        with concurrent.futures.ThreadPoolExecutor(max_workers=5) as ex:
            for t, vids in ex.map(load, pool):
                for v in vids:
                    if q and q not in ar_norm(v["title"] + " " + t["name"]): continue
                    feed.append({**v, "teacher_id": t["id"], "teacher": t["name"], "cid": t.get("channel_id") or "", "subject": subject_key(t["subject"]), "thumb": f"https://i.ytimg.com/vi/{v['id']}/mqdefault.jpg"})
    with db() as con: subs = [dict(r) for r in con.execute("SELECT cid,title FROM yt_sub WHERE user_id=? ORDER BY added_at DESC LIMIT 40", (int(uid),))]
    if mode == "mine" and not subj:
        def lsub(sc): return sc, yt_fetch(sc["cid"])
        if subs:
            with concurrent.futures.ThreadPoolExecutor(max_workers=5) as ex:
                for sc, vids in ex.map(lsub, subs[:12]):
                    for v in vids:
                        if q and q not in ar_norm(v["title"] + " " + sc["title"]): continue
                        feed.append({**v, "teacher_id": 0, "teacher": sc["title"], "cid": sc["cid"], "subject": "", "thumb": f"https://i.ytimg.com/vi/{v['id']}/mqdefault.jpg"})
    feed.sort(key=lambda v: v["published"], reverse=True)
    return {"teachers": [{"id": t["id"], "name": t["name"], "subject": subject_key(t["subject"]), "followed": t["id"] in fol, "cid": t.get("channel_id") or ""} for t in ts],
            "subs": subs, "feed": feed[:40], "followed": len(fol) + len(subs), "suggest": (yt_suggest(ts, q, subj) if mode == "all" else []), "api": bool(YT_KEY)}

# ── اقتراحات تيوب عبر YouTube Data API (مفتاح YOUTUBE_API_KEY) ──
YT_SUG_TTL = 6 * 3600
_YT_SUG = {}
YT_MOTIVE = ["تحفيز طلاب السادس الاعدادي للدراسة", "كيف تذاكر بتركيز وتنجح", "طريقة تنظيم الوقت للمذاكرة", "قصة نجاح طالب بعد السادس الاعدادي"]
YT_SUBJ_Q = {}
def yt_search(q, n=8):
    if not YT_KEY: return []
    now = time.time(); hit = _YT_SUG.get(q)
    if hit and now - hit[0] < YT_SUG_TTL: return hit[1]
    try:
        r = requests.get("https://www.googleapis.com/youtube/v3/search", params={"part": "snippet", "q": q, "type": "video", "maxResults": n, "videoEmbeddable": "true",
            "safeSearch": "strict", "relevanceLanguage": "ar", "regionCode": "IQ", "key": YT_KEY}, timeout=10)
        r.raise_for_status(); out = []
        for it in r.json().get("items", []):
            vid = (it.get("id") or {}).get("videoId", ""); sn = it.get("snippet") or {}
            if not re.fullmatch(r"[\w-]{11}", vid or ""): continue
            out.append({"id": vid, "title": html.unescape(sn.get("title", ""))[:160], "published": (sn.get("publishedAt") or "")[:19], "channel": html.unescape(sn.get("channelTitle", ""))[:60], "cid": sn.get("channelId", "")})
        _YT_SUG[q] = (now, out); return out
    except Exception as e:
        log.warning("yt search: %s", str(e)[:100])
        _YT_SUG[q] = (now - YT_SUG_TTL + 300, hit[1] if hit else []); return hit[1] if hit else []

def yt_suggest(ts, q="", subj=""):
    """اقتراحات: غالبها عن المدرّسين المضافين ثم التحفيز."""
    if not YT_KEY: return []
    qs = []
    if q: qs = [q, q + " سادس اعدادي"]
    else:
        import random
        names = [t["name"] for t in ts]; random.shuffle(names)
        qs = [f"{n} شرح سادس اعدادي" for n in names[:3]] + [f"شرح {subj_name(subj)} سادس اعدادي العراق"] * bool(subj) + random.sample(YT_MOTIVE, 2)
        if not names: qs = [f"شرح {subj_name(subj)} سادس اعدادي العراق"] * bool(subj) + YT_MOTIVE[:3] + ["شرح منهج السادس الاعدادي العراق"]
    seen, out = set(), []
    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as ex:
        for lst in ex.map(yt_search, qs):
            for v in lst:
                if v["id"] not in seen: seen.add(v["id"]); out.append({**v, "thumb": f"https://i.ytimg.com/vi/{v['id']}/mqdefault.jpg", "teacher": v["channel"], "sug": 1})
    # تداخل بين المصادر حتى يختلط المحتوى
    return out[:24]

def api_yt_related(uid, user, body):
    vid = str(body.get("id") or ""); title = str(body.get("title") or "")[:80]
    out = []
    if YT_KEY and title:
        out = [{**v, "thumb": f"https://i.ytimg.com/vi/{v['id']}/mqdefault.jpg", "teacher": v["channel"], "sug": 1} for v in yt_search(title, 12) if v["id"] != vid]
    if len(out) < 6:
        for t in yt_teachers()[:6]:
            cid = yt_channel_id(t)
            for v in (yt_fetch(cid) if cid else [])[:4]:
                if v["id"] != vid and not any(o["id"] == v["id"] for o in out):
                    out.append({**v, "teacher": t["name"], "thumb": f"https://i.ytimg.com/vi/{v['id']}/mqdefault.jpg"})
    return {"related": out[:16]}

_CID = re.compile(r"^UC[\w-]{22}$")
def api_yt_sub(uid, user, body):
    cid = str(body.get("cid") or "").strip(); title = str(body.get("title") or "").strip()[:80]
    if not _CID.match(cid): return err("قناة غير صالحة")
    with db() as con:
        if body.get("on"):
            if con.execute("SELECT COUNT(*) n FROM yt_sub WHERE user_id=?", (int(uid),)).fetchone()["n"] >= 40: return err("وصلت حد الاشتراكات (40)")
            con.execute("INSERT OR REPLACE INTO yt_sub VALUES(?,?,?,?)", (int(uid), cid, title or _YT_META.get(cid, "قناة"), now_iso()))
        else: con.execute("DELETE FROM yt_sub WHERE user_id=? AND cid=?", (int(uid), cid))
    return {"ok": True}

def api_yt_channel(uid, user, body):
    cid = str(body.get("cid") or "").strip()
    if not _CID.match(cid): return err("قناة غير صالحة")
    vids = yt_fetch(cid)
    with db() as con:
        sub = con.execute("SELECT title FROM yt_sub WHERE user_id=? AND cid=?", (int(uid), cid)).fetchone()
        t = con.execute("SELECT id,name FROM teachers WHERE channel_id=? AND active=1", (cid,)).fetchone()
    title = (sub["title"] if sub else "") or (t["name"] if t else "") or _YT_META.get(cid) or str(body.get("title") or "")[:80] or "قناة"
    return {"cid": cid, "title": title, "subscribed": bool(sub), "teacher_id": t["id"] if t else 0,
            "videos": [{**v, "cid": cid, "teacher": title, "thumb": f"https://i.ytimg.com/vi/{v['id']}/mqdefault.jpg"} for v in vids]}

_PRES = {}
def api_presence(uid, user, body):
    now = time.time(); _PRES[int(uid)] = (now, bool(body.get("studying")))
    for k in [k for k, v in _PRES.items() if now - v[0] > 90]: _PRES.pop(k, None)
    return {"studying": sum(1 for v in _PRES.values() if v[1]), "online": len(_PRES)}

def api_admin_stars_list(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    q = str(body.get("q") or "").strip()[:40]
    with db() as con:
        rows = con.execute("SELECT g.user_id id, g.stars, g.fcode, COALESCE(u.first_name,'') n FROM game g LEFT JOIN users u ON u.user_id=g.user_id ORDER BY g.stars DESC LIMIT 200").fetchall()
    out = [{"id": r["id"], "name": r["n"] or "طالب", "stars": r["stars"], "code": r["fcode"] or ""} for r in rows]
    if q: out = [o for o in out if q in str(o["id"]) or q in o["name"] or q.upper() == o["code"]]
    return {"users": out[:50]}

def api_admin_stars(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    try: tid = int(body.get("id")); delta = int(body.get("delta"))
    except (TypeError, ValueError): return err("اكتب رقم الطالب والعدد")
    if not delta or abs(delta) > 1000000: return err("العدد غير صالح")
    game_row(tid)
    with db() as con:
        cur = con.execute("SELECT stars FROM game WHERE user_id=?", (tid,)).fetchone()["stars"]
        new = max(0, cur + delta)
        con.execute("UPDATE game SET stars=? WHERE user_id=?", (new, tid))
    log.info("admin %s stars %s: %+d -> %d", uid, tid, delta, new)
    return {"id": tid, "stars": new}

# ── لوحة تحكم الأزرار والإحصائيات (للمالك/المشرفين) ──
_BTN_KINDS = ("cb", "url", "webapp", "text"); _BTN_STYLES = ("", "primary", "success", "danger")

def api_admin_btns(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    return {"btns": menu_load(True), "styles_native": _styles_supported()}

def _styles_supported():
    try: types.InlineKeyboardButton("x", callback_data="x", style="success"); return True
    except Exception: return False

def api_admin_btn_save(uid, user, body):
    g = _need_owner(uid)
    if g: return g
    label = str(body.get("label") or "").strip()[:40]; kind = str(body.get("kind") or "cb"); value = str(body.get("value") or "").strip()
    style = str(body.get("style") or ""); full = 1 if body.get("full") else 0; active = 0 if body.get("active") is False else 1
    if not label: return err("اكتب اسم الزر")
    if kind not in _BTN_KINDS or style not in _BTN_STYLES: return err("نوع غير صالح")
    if kind == "url" and not re.match(r"^(https?://|tg://)", value): return err("الرابط لازم يبدأ بـ https://")
    if kind == "text" and not value: return err("اكتب النص اللي يظهر عند الضغط")
    if kind == "cb" and not value.startswith("menu:"): return err("قيمة الزر الداخلي غير صالحة")
    if kind in ("text",): value = value[:1500]
    if kind == "url": value = value[:300]
    bid = body.get("id")
    with db() as con:
        if bid:
            if not con.execute("SELECT 1 FROM menu_btns WHERE id=?", (int(bid),)).fetchone(): return err("الزر غير موجود", 404)
            con.execute("UPDATE menu_btns SET label=?,kind=?,value=?,style=?,full=?,active=? WHERE id=?", (label, kind, value, style, full, active, int(bid)))
        else:
            if con.execute("SELECT COUNT(*) n FROM menu_btns").fetchone()["n"] >= 40: return err("وصلت الحد الأقصى (40 زر)")
            pos = (con.execute("SELECT COALESCE(MAX(pos),0)+1 p FROM menu_btns").fetchone()["p"])
            con.execute("INSERT INTO menu_btns(label,kind,value,style,pos,active,full) VALUES(?,?,?,?,?,?,?)", (label, kind, value, style, pos, active, full))
    return {"btns": menu_load(True)}

def api_admin_btn_del(uid, user, body):
    g = _need_owner(uid)
    if g: return g
    with db() as con: con.execute("DELETE FROM menu_btns WHERE id=?", (int(body.get("id") or 0),))
    return {"btns": menu_load(True)}

def api_admin_btn_move(uid, user, body):
    g = _need_owner(uid)
    if g: return g
    rows = menu_load(True); bid = int(body.get("id") or 0); d = -1 if int(body.get("dir") or 0) < 0 else 1
    ids = [r["id"] for r in rows]
    if bid not in ids: return err("الزر غير موجود", 404)
    i = ids.index(bid); j = i + d
    if 0 <= j < len(ids):
        ids[i], ids[j] = ids[j], ids[i]
        with db() as con:
            for p, x in enumerate(ids): con.execute("UPDATE menu_btns SET pos=? WHERE id=?", (p, x))
    return {"btns": menu_load(True)}

def api_admin_btn_reset(uid, user, body):
    g = _need_owner(uid)
    if g: return g
    with db() as con: con.execute("DELETE FROM menu_btns")
    return {"btns": menu_load(True)}

def _primary_channel():
    for c in CHANNEL_ENV:
        return int(c) if re.fullmatch(r"-?\d+", c) else c
    with db() as con: r = con.execute("SELECT chat_id FROM channels WHERE approved=1 ORDER BY added_at LIMIT 1").fetchone()
    return r["chat_id"] if r else None

def _stats_init():
    with db() as con:
        try: con.execute("ALTER TABLE users ADD COLUMN sub_chan INTEGER")
        except Exception: pass

def _day_ago(n): return (datetime.now() - timedelta(days=n)).strftime("%Y-%m-%d")

def api_admin_stats(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    _stats_init()
    d1, d7, d30 = _day_ago(1) , _day_ago(7), _day_ago(30)
    with db() as con:
        q1 = lambda sql, *a: con.execute(sql, a).fetchone()[0]
        out = {"total": q1("SELECT COUNT(*) FROM users"), "banned": q1("SELECT COUNT(*) FROM users WHERE banned=1"),
               "started": q1("SELECT COUNT(*) FROM users WHERE private_started=1"),
               "new1": q1("SELECT COUNT(*) FROM users WHERE joined_at>=?", d1), "new7": q1("SELECT COUNT(*) FROM users WHERE joined_at>=?", d7), "new30": q1("SELECT COUNT(*) FROM users WHERE joined_at>=?", d30),
               "act1": q1("SELECT COUNT(*) FROM users WHERE last_seen>=?", d1), "act7": q1("SELECT COUNT(*) FROM users WHERE last_seen>=?", d7), "act30": q1("SELECT COUNT(*) FROM users WHERE last_seen>=?", d30),
               "sub_yes": q1("SELECT COUNT(*) FROM users WHERE sub_chan=1"), "sub_no": q1("SELECT COUNT(*) FROM users WHERE sub_chan=0"), "sub_unk": q1("SELECT COUNT(*) FROM users WHERE sub_chan IS NULL"),
               "msgs": q1("SELECT COUNT(*) FROM messages WHERE role='user'"), "quizzes": q1("SELECT COUNT(*) FROM quiz_results"), "focus": q1("SELECT COALESCE(SUM(minutes),0) FROM focus_sessions")}
        days = {r["d"]: r["n"] for r in con.execute("SELECT substr(joined_at,1,10) d, COUNT(*) n FROM users WHERE joined_at>=? GROUP BY d", (_day_ago(13),))}
        out["series"] = [{"d": _day_ago(i)[5:], "n": days.get(_day_ago(i), 0)} for i in range(13, -1, -1)]
        out["top_btns"] = [dict(r) for r in con.execute("SELECT label,clicks FROM menu_btns WHERE clicks>0 ORDER BY clicks DESC LIMIT 8")] if con.execute("SELECT name FROM sqlite_master WHERE name='menu_btns'").fetchone() else []
    out["channel"] = bool(_primary_channel())
    return out

def api_admin_users(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    _stats_init()
    q = str(body.get("q") or "").strip()[:40]; flt = str(body.get("f") or "all"); page = max(0, int(body.get("page") or 0))
    where, args = [], []
    if q.isdigit(): where.append("u.user_id=?"); args.append(int(q))
    elif q: where.append("(u.first_name LIKE ? OR u.username LIKE ?)"); args += [f"%{q}%", f"%{q}%"]
    if flt == "banned": where.append("u.banned=1")
    elif flt == "active": where.append("u.last_seen>=?"); args.append(_day_ago(7))
    elif flt == "sub": where.append("u.sub_chan=1")
    elif flt == "nosub": where.append("(u.sub_chan=0 OR u.sub_chan IS NULL)")
    w = ("WHERE " + " AND ".join(where)) if where else ""
    with db() as con:
        total = con.execute(f"SELECT COUNT(*) FROM users u {w}", args).fetchone()[0]
        rows = [dict(r) for r in con.execute(f"SELECT u.user_id id,u.first_name n,u.username un,u.joined_at j,u.last_seen ls,u.banned b,u.sub_chan sc FROM users u {w} ORDER BY u.last_seen DESC LIMIT 25 OFFSET ?", args + [page * 25])]
        ids = [r["id"] for r in rows]
        if ids:
            ph = ",".join("?" * len(ids))
            msgs = {r[0]: r[1] for r in con.execute(f"SELECT user_id,COUNT(*) FROM messages WHERE role='user' AND user_id IN ({ph}) GROUP BY user_id", ids)}
            qz = {r[0]: r[1] for r in con.execute(f"SELECT user_id,COUNT(*) FROM quiz_results WHERE user_id IN ({ph}) GROUP BY user_id", ids)}
            fm = {r[0]: r[1] for r in con.execute(f"SELECT user_id,COALESCE(SUM(minutes),0) FROM focus_sessions WHERE user_id IN ({ph}) GROUP BY user_id", ids)}
            try: gm = {r[0]: r[1] for r in con.execute(f"SELECT user_id,stars FROM game WHERE user_id IN ({ph})", ids)}
            except Exception: gm = {}
        else: msgs = qz = fm = gm = {}
    for r in rows:
        r.update(name=r.pop("n") or "طالب", msgs=msgs.get(r["id"], 0), quizzes=qz.get(r["id"], 0), focus=fm.get(r["id"], 0), coins=gm.get(r["id"], 0))
    return {"users": rows, "total": total, "page": page}

def api_admin_ban(uid, user, body):
    g = _need_owner(uid)
    if g: return g
    try: tid = int(body.get("id"))
    except (TypeError, ValueError): return err("رقم غير صالح")
    flag = 1 if body.get("flag") else 0
    if flag and tid in OWNER_IDS: return err("لا يمكن حظر المالك")
    with db() as con:
        con.execute("INSERT INTO users(user_id,joined_at,last_seen,banned) VALUES(?,?,?,?) ON CONFLICT(user_id) DO UPDATE SET banned=excluded.banned", (tid, now_iso(), now_iso(), flag))
    _BAN.pop(tid, None); db_event(tid, "banned" if flag else "unbanned", uid)
    return {"id": tid, "banned": flag}

def api_admin_subcheck(uid, user, body):
    g = _need_admin(uid)
    if g: return g
    _stats_init(); ch = _primary_channel()
    if not ch: return err("ما أكو قناة معتمدة. اضبط CHANNEL_ID وخلّ البوت مشرفاً بالقناة.")
    one = body.get("id")
    with db() as con:
        if one: ids = [int(one)]
        else: ids = [r[0] for r in con.execute("SELECT user_id FROM users WHERE sub_chan IS NULL AND banned=0 ORDER BY last_seen DESC LIMIT 30")]
    done = bad = 0
    for tid in ids:
        try:
            st = bot.get_chat_member(ch, tid).status; val = 1 if st in ("member", "administrator", "creator") else 0
        except Exception as e:
            bad += 1
            if "chat not found" in str(e).lower() or "not enough rights" in str(e).lower() or "administrator" in str(e).lower(): return err("البوت لازم يكون مشرف بالقناة حتى يفحص الاشتراك.")
            continue
        with db() as con: con.execute("UPDATE users SET sub_chan=? WHERE user_id=?", (val, tid))
        done += 1; time.sleep(0.05)
    return {"checked": done, "failed": bad, "left": len(ids) - done - bad}

def api_yt_follow(uid, user, body):
    try: tid = int(body.get("id"))
    except (TypeError, ValueError): return err("مدرّس غير صالح")
    on = bool(body.get("on"))
    with db() as con:
        if not con.execute("SELECT 1 FROM teachers WHERE id=? AND active=1", (tid,)).fetchone(): return err("مدرّس غير موجود", 404)
        if on:
            if con.execute("SELECT COUNT(*) n FROM yt_follow WHERE user_id=?", (int(uid),)).fetchone()["n"] >= 60: return err("وصلت حد المتابعة")
            con.execute("INSERT OR IGNORE INTO yt_follow(user_id,teacher_id) VALUES(?,?)", (int(uid), tid))
        else: con.execute("DELETE FROM yt_follow WHERE user_id=? AND teacher_id=?", (int(uid), tid))
    return {"ok": True}

# ── الأصدقاء + غرف الدراسة ──
_CODE_AL = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
_RATE = {}
def _rate(uid, n=20, per=60):
    now = time.time(); lst = [t for t in _RATE.get(uid, []) if now - t < per]
    if len(lst) >= n: _RATE[uid] = lst; return False
    lst.append(now); _RATE[uid] = lst; return True
def _rcode(n):
    return "".join(random.choice(_CODE_AL) for _ in range(n))
def my_fcode(uid):
    g = game_row(uid)
    if g.get("fcode"): return g["fcode"]
    with db() as con:
        for _ in range(20):
            c = _rcode(6)
            if not con.execute("SELECT 1 FROM game WHERE fcode=?", (c,)).fetchone():
                con.execute("UPDATE game SET fcode=? WHERE user_id=?", (c, int(uid))); return c
    return ""
def _uname(con, uid):
    r = con.execute("SELECT first_name FROM users WHERE user_id=?", (int(uid),)).fetchone()
    return ((r["first_name"] if r else "") or "طالب")[:18]
def _studying(con, uid):
    r = con.execute("SELECT 1 FROM room_members WHERE user_id=? AND studying=1 AND last_seen>?", (int(uid), time.time() - 75)).fetchone()
    return bool(r)
def _are_friends(con, a, b):
    return bool(con.execute("SELECT 1 FROM friends WHERE user_id=? AND friend_id=?", (int(a), int(b))).fetchone())
def _share_room(con, a, b):
    return bool(con.execute("SELECT 1 FROM room_members x JOIN room_members y ON x.code=y.code WHERE x.user_id=? AND y.user_id=?", (int(a), int(b))).fetchone())
def _brief(con, fid):
    g = con.execute("SELECT stars,sticker,plots FROM game WHERE user_id=?", (int(fid),)).fetchone()
    td = con.execute("SELECT COALESCE(focus_min,0) f FROM activity WHERE user_id=? AND day=?", (int(fid), today_key())).fetchone()
    ready = 0
    for r in con.execute("SELECT crop,progress FROM farm WHERE user_id=?", (int(fid),)):
        c = CROP.get(r["crop"])
        if c and r["progress"] >= c["need"]: ready += 1
    return {"id": int(fid), "name": _uname(con, fid), "sticker": (g["sticker"] if g else ""), "stars": (g["stars"] if g else 0), "plots": (g["plots"] if g else 1),
            "ready": ready, "today": (td["f"] if td else 0), "studying": _studying(con, fid)}

def api_social(uid, user, body):
    code = my_fcode(uid)
    with db() as con:
        fr = [_brief(con, r["friend_id"]) for r in con.execute("SELECT friend_id FROM friends WHERE user_id=?", (int(uid),))]
    fr.sort(key=lambda f: (not f["studying"], -f["today"]))
    with db() as con:
        inv = [{"id": r["id"], "from": r["from_id"], "name": _uname(con, r["from_id"])} for r in con.execute(
            "SELECT id,from_id FROM house_invites WHERE to_id=? AND created_at>? ORDER BY id DESC LIMIT 10", (int(uid), time.time() - 3 * 86400))]
    return {"code": code, "friends": fr, "invites": inv}

def api_friend_add(uid, user, body):
    if not _rate(("fa", uid), 12): return err("حاول بعد قليل", 429)
    code = re.sub(r"[^A-Z2-9]", "", str(body.get("code") or "").upper())[:6]
    if len(code) != 6: return err("الكود 6 أحرف")
    my_fcode(uid)
    with db() as con:
        r = con.execute("SELECT user_id FROM game WHERE fcode=?", (code,)).fetchone()
        if not r: return err("ما لقيت هذا الكود", 404)
        fid = r["user_id"]
        if int(fid) == int(uid): return err("هذا كودك أنت")
        if con.execute("SELECT COUNT(*) n FROM friends WHERE user_id=?", (int(uid),)).fetchone()["n"] >= 60: return err("وصلت حد الأصدقاء")
        con.execute("INSERT OR IGNORE INTO friends VALUES(?,?)", (int(uid), int(fid)))
        con.execute("INSERT OR IGNORE INTO friends VALUES(?,?)", (int(fid), int(uid)))
    return api_social(uid, user, {})

def api_friend_del(uid, user, body):
    try: fid = int(body.get("id"))
    except (TypeError, ValueError): return err("غير صالح")
    with db() as con:
        con.execute("DELETE FROM friends WHERE (user_id=? AND friend_id=?) OR (user_id=? AND friend_id=?)", (int(uid), fid, fid, int(uid)))
    return api_social(uid, user, {})

def api_farm_view(uid, user, body):
    try: fid = int(body.get("id"))
    except (TypeError, ValueError): return err("غير صالح")
    with db() as con:
        if fid != int(uid) and not (_are_friends(con, uid, fid) or _share_room(con, uid, fid)): return err("لازم يكون صديقك أو بنفس غرفتك", 403)
        b = _brief(con, fid)
        wk = (datetime.now() - timedelta(days=6)).strftime("%Y-%m-%d") if "datetime" in globals() else ""
        tot = con.execute("SELECT COALESCE(SUM(focus_min),0) f FROM activity WHERE user_id=?", (fid,)).fetchone()["f"]
        wkm = con.execute("SELECT COALESCE(SUM(focus_min),0) f FROM activity WHERE user_id=? AND day>=?", (fid, wk)).fetchone()["f"] if wk else 0
    st = game_state(fid)
    b.update({"farm": st["farm"], "earned": st["earned"], "total": tot, "week": wkm, "crops": [{"key": c["key"], "name": c["name"]} for c in CROP.values()]})
    return b

def _room_members(con, code, uid):
    now = time.time(); out = []
    for r in con.execute("SELECT user_id,last_seen,studying,since FROM room_members WHERE code=? ORDER BY studying DESC, last_seen DESC", (code,)):
        b = _brief(con, r["user_id"]); on = now - r["last_seen"] < 75
        out.append({"id": b["id"], "name": b["name"], "sticker": b["sticker"], "me": b["id"] == int(uid), "online": on, "studying": bool(r["studying"]) and on,
                    "mins": int(max(0, now - r["since"]) // 60) if (r["studying"] and on and r["since"]) else 0, "today": b["today"], "stars": b["stars"]})
    return out

def api_room_create(uid, user, body):
    if not _rate(("rc", uid), 6): return err("حاول بعد قليل", 429)
    name = str(body.get("name") or "").strip()[:30] or "غرفة دراسة"
    with db() as con:
        if con.execute("SELECT COUNT(*) n FROM rooms WHERE owner=?", (int(uid),)).fetchone()["n"] >= 5: return err("عندك 5 غرف، احذف غرفة أولاً")
        for _ in range(20):
            code = _rcode(5)
            if not con.execute("SELECT 1 FROM rooms WHERE code=?", (code,)).fetchone(): break
        con.execute("INSERT INTO rooms VALUES(?,?,?,?)", (code, name, int(uid), now_iso()))
        con.execute("INSERT OR REPLACE INTO room_members VALUES(?,?,?,0,0)", (code, int(uid), time.time()))
    return {"code": code, "name": name}

def api_room_join(uid, user, body):
    if not _rate(("rj", uid), 15): return err("حاول بعد قليل", 429)
    code = re.sub(r"[^A-Z2-9]", "", str(body.get("code") or "").upper())[:5]
    with db() as con:
        r = con.execute("SELECT name FROM rooms WHERE code=?", (code,)).fetchone()
        if not r: return err("ما لقيت غرفة بهذا الكود", 404)
        n = con.execute("SELECT COUNT(*) n FROM room_members WHERE code=? AND user_id<>?", (code, int(uid))).fetchone()["n"]
        if n >= 12: return err("الغرفة ممتلئة (12 طالب)")
        con.execute("INSERT OR IGNORE INTO room_members VALUES(?,?,?,0,0)", (code, int(uid), time.time()))
    return {"code": code, "name": r["name"]}

def api_room_state(uid, user, body):
    code = re.sub(r"[^A-Z2-9]", "", str(body.get("code") or "").upper())[:5]
    studying = 1 if body.get("studying") else 0
    try: el = max(0, min(int(body.get("elapsed") or 0), 14400))
    except (TypeError, ValueError): el = 0
    with db() as con:
        room = con.execute("SELECT code,name,owner FROM rooms WHERE code=?", (code,)).fetchone()
        if not room: return err("الغرفة غير موجودة", 404)
        if not con.execute("SELECT 1 FROM room_members WHERE code=? AND user_id=?", (code, int(uid))).fetchone(): return err("لست عضواً بالغرفة", 403)
        con.execute("UPDATE room_members SET last_seen=?, studying=?, since=? WHERE code=? AND user_id=?", (time.time(), studying, time.time() - el if studying else 0, code, int(uid)))
        return {"code": code, "name": room["name"], "owner": room["owner"] == int(uid), "members": _room_members(con, code, uid)}

def api_room_leave(uid, user, body):
    code = re.sub(r"[^A-Z2-9]", "", str(body.get("code") or "").upper())[:5]
    with db() as con:
        con.execute("DELETE FROM room_members WHERE code=? AND user_id=?", (code, int(uid)))
        room = con.execute("SELECT owner FROM rooms WHERE code=?", (code,)).fetchone()
        if room and room["owner"] == int(uid) and body.get("delete"):
            con.execute("DELETE FROM room_members WHERE code=?", (code,)); con.execute("DELETE FROM rooms WHERE code=?", (code,))
    return api_rooms(uid, user, {})

def api_rooms(uid, user, body):
    with db() as con:
        rs = []
        for r in con.execute("SELECT r.code,r.name,r.owner FROM rooms r JOIN room_members m ON m.code=r.code WHERE m.user_id=? ORDER BY r.created_at DESC", (int(uid),)):
            ms = _room_members(con, r["code"], uid)
            rs.append({"code": r["code"], "name": r["name"], "owner": r["owner"] == int(uid), "count": len(ms), "studying": sum(1 for m in ms if m["studying"])})
    return {"rooms": rs}

# ── البيت: سوبر ماركت (أدوات منزلية/أثاث/إضاءة فضائية) + شخصيات + دعوات الأصدقاء ──
HOUSE_SECTIONS = [
    ("tools", "أدوات منزلية", [("pot", "قدر", 40), ("oil", "زيت", 30), ("kettle", "غلاية", 60), ("fan", "مروحة", 90), ("vacuum", "مكنسة", 120), ("stove", "طباخ", 150),
                              ("fridge", "ثلاجة", 300), ("washer", "غسالة", 350), ("ac", "سبلت", 400), ("tv", "تلفزيون", 450)]),
    ("furniture", "أثاث", [("plant", "نبتة", 40), ("clock", "ساعة", 50), ("curtain", "ستارة", 70), ("rug", "سجادة", 80), ("shelf", "رف كتب", 100), ("desk", "مكتب دراسة", 120),
                          ("bed", "سرير", 200), ("sofa", "كنبة", 250)]),
    ("lamps", "مصابيح الفضاء", [("lamp_star", "مصباح نجمة", 100), ("lamp_moon", "مصباح هلال", 150), ("lamp_planet", "مصباح كوكب", 200), ("lamp_comet", "مصباح مذنب", 250), ("lamp_galaxy", "مصباح مجرّة", 300)]),
]
HOUSE_ITEM = {k: {"key": k, "name": n, "price": p, "section": sec} for sec, _, its in HOUSE_SECTIONS for k, n, p in its}
HOUSE_CHARS = [("bear", "دبدوب", 0), ("duck", "بطبوط", 100), ("bunny", "أرنوب", 120), ("catc", "قطوط", 150), ("fox", "ثعلوب", 200), ("owl", "بومة", 250)]
HOUSE_CHAR = {k: {"key": k, "name": n, "price": p} for k, n, p in HOUSE_CHARS}
HOUSE_SLOTS = 12
MAX_CHARS = 5
def _clean_name(t, n):
    t = re.sub(r"[<>&\"'`\\\\]", "", str(t or "")).strip()
    return re.sub(r"\s+", " ", t)[:n]

def house_row(uid):
    with db() as con:
        r = con.execute("SELECT * FROM house WHERE user_id=?", (int(uid),)).fetchone()
        if not r:
            con.execute("INSERT INTO house(user_id,name,owned,placed,chars,created_at) VALUES(?,?,?,?,?,?)", (int(uid), "بيتي", "[]", "{}", json.dumps([{"k": "bear", "name": "دبدوب"}], ensure_ascii=False), now_iso()))
            r = con.execute("SELECT * FROM house WHERE user_id=?", (int(uid),)).fetchone()
        d = dict(r)
    def jl(x, dflt):
        try: v = json.loads(x or "")
        except Exception: return dflt
        return v if isinstance(v, type(dflt)) else dflt
    d["owned"] = [k for k in jl(d["owned"], []) if k in HOUSE_ITEM]
    d["placed"] = {str(k): v for k, v in jl(d["placed"], {}).items() if str(k).isdigit() and 0 <= int(k) < HOUSE_SLOTS and v in d["owned"]}
    d["chars"] = [c for c in jl(d["chars"], []) if isinstance(c, dict) and c.get("k") in HOUSE_CHAR][:MAX_CHARS]
    return d

def _house_save(uid, d):
    with db() as con:
        con.execute("UPDATE house SET name=?, owned=?, placed=?, chars=? WHERE user_id=?",
                    (d["name"], json.dumps(d["owned"]), json.dumps(d["placed"]), json.dumps(d["chars"], ensure_ascii=False), int(uid)))

def house_state(uid, catalog=False):
    d = house_row(uid)
    val = sum(HOUSE_ITEM[k]["price"] for k in d["placed"].values() if k in HOUSE_ITEM)
    out = {"name": d["name"], "owned": d["owned"], "placed": d["placed"], "chars": d["chars"], "value": val, "slots": HOUSE_SLOTS, "max_chars": MAX_CHARS}
    if catalog:
        own = set(d["owned"]); have = {c["k"] for c in d["chars"]}
        out["sections"] = [{"key": sec, "name": nm, "items": [{**HOUSE_ITEM[k], "owned": k in own} for k, _, _ in its]} for sec, nm, its in HOUSE_SECTIONS]
        out["char_shop"] = [{**c, "owned": c["key"] in have} for c in HOUSE_CHAR.values()]
    return out

def api_house(uid, user, body):
    return house_state(uid, True)

def api_house_act(uid, user, body):
    act = str(body.get("action") or ""); d = house_row(uid)
    if act == "rename":
        n = _clean_name(body.get("name"), 24)
        if not n: return err("اكتب اسماً لبيتك")
        d["name"] = n
    elif act == "buy":
        k = str(body.get("item") or ""); it = HOUSE_ITEM.get(k)
        if not it: return err("قطعة غير موجودة")
        if k in d["owned"]: return err("تملكها من قبل")
        if not _spend(uid, it["price"]): return err("كوينزاتك لا تكفي")
        d["owned"].append(k)
    elif act == "place":
        k = str(body.get("item") or ""); sl = str(body.get("slot"))
        if k not in d["owned"]: return err("لا تملك هذه القطعة")
        if not sl.isdigit() or not 0 <= int(sl) < HOUSE_SLOTS: return err("مكان غير صالح")
        d["placed"] = {a: b for a, b in d["placed"].items() if b != k}; d["placed"][sl] = k
    elif act == "remove":
        d["placed"].pop(str(body.get("slot")), None)
    elif act == "char_buy":
        k = str(body.get("key") or ""); c = HOUSE_CHAR.get(k)
        if not c: return err("شخصية غير موجودة")
        if any(x["k"] == k for x in d["chars"]): return err("موجودة ببيتك")
        if len(d["chars"]) >= MAX_CHARS: return err(f"البيت يتسع لـ{MAX_CHARS} شخصيات")
        if c["price"] and not _spend(uid, c["price"]): return err("كوينزاتك لا تكفي")
        d["chars"].append({"k": k, "name": c["name"]})
    elif act == "char_name":
        try: i = int(body.get("idx")); n = _clean_name(body.get("name"), 14)
        except (TypeError, ValueError): return err("غير صالح")
        if not 0 <= i < len(d["chars"]) or not n: return err("اكتب اسماً")
        d["chars"][i]["name"] = n
    elif act == "char_remove":
        try: i = int(body.get("idx"))
        except (TypeError, ValueError): return err("غير صالح")
        if 0 <= i < len(d["chars"]): d["chars"].pop(i)
    else: return err("إجراء غير معروف")
    _house_save(uid, d)
    return {"house": house_state(uid, True), "game": game_state(uid)}

def api_house_view(uid, user, body):
    try: fid = int(body.get("id"))
    except (TypeError, ValueError): return err("غير صالح")
    with db() as con:
        if fid != int(uid) and not (_are_friends(con, uid, fid) or _share_room(con, uid, fid)): return err("لازم يكون صديقك أو بنفس غرفتك", 403)
    return house_state(fid)


# ── المدينة: خريطة تشتري إليها أشياء وتحطها، والدب يمشي عليها ──
TOWN_COLS, TOWN_ROWS, TOWN_MAX = 9, 11, 80
TOWN_SECTIONS = [
    ("nature", "طبيعة", [("tree", "شجرة", 30), ("apple", "شجرة تفاح", 60), ("palm", "نخلة", 80), ("bush", "شجيرة", 15), ("flower", "ورود", 20), ("pond", "بركة", 100)]),
    ("build", "مباني", [("tent", "خيمة", 150), ("house", "بيت صغير", 200), ("barn", "حظيرة", 250), ("bakery", "مخبز", 300), ("mill", "طاحونة", 350), ("school", "مدرسة", 500), ("well", "بئر", 120)]),
    ("decor", "ديكور", [("fence", "سياج", 10), ("crate", "صندوق", 25), ("hay", "كومة قش", 35), ("lamp", "عمود إنارة", 50), ("bench", "مقعد", 50), ("fountain", "نافورة", 180)]),
]
TOWN_PROD = {"apple": (20, 12, "تفاح"), "palm": (30, 35, "تمر"), "well": (15, 15, "ماء"), "pond": (20, 14, "سمك"), "barn": (25, 30, "بيض"), "bakery": (30, 40, "خبز"), "mill": (40, 55, "طحين")}
TOWN_ITEM = {k: {"key": k, "name": n, "price": p, "section": sec} for sec, _, its in TOWN_SECTIONS for k, n, p in its}

def town_row(uid):
    with db() as con:
        r = con.execute("SELECT * FROM town WHERE user_id=?", (int(uid),)).fetchone()
        if not r:
            con.execute("INSERT INTO town(user_id) VALUES(?)", (int(uid),))
            r = con.execute("SELECT * FROM town WHERE user_id=?", (int(uid),)).fetchone()
        d = dict(r)
    try: raw = json.loads(d["items"] or "{}")
    except Exception: raw = {}
    items = {}
    for k, v in (raw.items() if isinstance(raw, dict) else []):
        try: x, y = (int(t) for t in str(k).split(","))
        except Exception: continue
        if 0 <= x < TOWN_COLS and 0 <= y < TOWN_ROWS and v in TOWN_ITEM: items[f"{x},{y}"] = v
    d["items"] = items
    try: rd = json.loads(d.get("ready") or "{}")
    except Exception: rd = {}
    d["ready"] = {k: float(v) for k, v in (rd.items() if isinstance(rd, dict) else []) if k in items and isinstance(v, (int, float))}
    for k, v in items.items():
        if v in TOWN_PROD and k not in d["ready"]: d["ready"][k] = time.time() + TOWN_PROD[v][0] * 60
    if f"{d['bx']},{d['by']}" in items or not (0 <= d["bx"] < TOWN_COLS and 0 <= d["by"] < TOWN_ROWS): d["bx"], d["by"] = 4, 5
    return d

def town_state(uid, catalog=False):
    d = town_row(uid)
    out = {"cols": TOWN_COLS, "rows": TOWN_ROWS, "items": d["items"], "bear": {"x": d["bx"], "y": d["by"]},
           "value": sum(TOWN_ITEM[v]["price"] for v in d["items"].values()), "count": len(d["items"]), "max": TOWN_MAX,
           "now": time.time(), "ready": {k: max(0, int(t - time.time())) for k, t in d["ready"].items()},
           "prod": {k: {"min": m, "coins": c, "what": w} for k, (m, c, w) in TOWN_PROD.items()}}
    if catalog: out["sections"] = [{"key": sec, "name": nm, "items": [TOWN_ITEM[k] for k, _, _ in its]} for sec, nm, its in TOWN_SECTIONS]
    return out

def _town_save(uid, d):
    with db() as con: con.execute("UPDATE town SET items=?, bx=?, by=?, ready=? WHERE user_id=?", (json.dumps(d["items"]), int(d["bx"]), int(d["by"]), json.dumps(d["ready"]), int(uid)))

def _xy(body, a="x", b="y"):
    try: x, y = int(body.get(a)), int(body.get(b))
    except (TypeError, ValueError): return None
    return (x, y) if 0 <= x < TOWN_COLS and 0 <= y < TOWN_ROWS else None

def api_town(uid, user, body):
    return town_state(uid, True)

def api_town_act(uid, user, body):
    act = str(body.get("action") or ""); d = town_row(uid); it = d["items"]
    p = _xy(body)
    if not p: return err("مكان غير صالح")
    key = f"{p[0]},{p[1]}"
    if act == "place":
        k = str(body.get("item") or ""); c = TOWN_ITEM.get(k)
        if not c: return err("قطعة غير موجودة")
        if key in it: return err("المكان مشغول")
        if p == (d["bx"], d["by"]): return err("الدب واقف هنا")
        if len(it) >= TOWN_MAX: return err("المدينة امتلأت")
        if not _spend(uid, c["price"]): return err("كوينزاتك لا تكفي")
        it[key] = k
        if k in TOWN_PROD: d["ready"][key] = time.time() + TOWN_PROD[k][0] * 60
    elif act == "collect":
        k = it.get(key)
        if k not in TOWN_PROD: return err("ما كو شي يُجمع هنا")
        if d["ready"].get(key, 0) > time.time(): return err("لسه ما جهز")
        mins, coins, what = TOWN_PROD[k]
        d["ready"][key] = time.time() + mins * 60
        with db() as con: con.execute("UPDATE game SET stars=stars+?, earned=earned+? WHERE user_id=?", (coins, coins, int(uid)))
    elif act == "sell":
        k = it.pop(key, None); d["ready"].pop(key, None)
        if not k: return err("ما كو شي هنا")
        with db() as con: con.execute("UPDATE game SET stars=stars+? WHERE user_id=?", (TOWN_ITEM[k]["price"] // 2, int(uid)))
    elif act == "move":
        q = _xy(body, "tx", "ty")
        if not q: return err("مكان غير صالح")
        tk = f"{q[0]},{q[1]}"
        if key not in it: return err("ما كو شي هنا")
        if tk in it: return err("المكان مشغول")
        if q == (d["bx"], d["by"]): return err("الدب واقف هنا")
        it[tk] = it.pop(key)
        if key in d["ready"]: d["ready"][tk] = d["ready"].pop(key)
    elif act == "bear":
        if key in it: return err("مكان مشغول")
        d["bx"], d["by"] = p
    else: return err("إجراء غير معروف")
    _town_save(uid, d)
    return {"town": town_state(uid), "game": game_state(uid)}

def api_town_view(uid, user, body):
    try: fid = int(body.get("id"))
    except (TypeError, ValueError): return err("غير صالح")
    with db() as con:
        if fid != int(uid) and not (_are_friends(con, uid, fid) or _share_room(con, uid, fid)): return err("لازم يكون صديقك أو بنفس غرفتك", 403)
    return town_state(fid)

# ── عالم المزرعة الكبير (إيزومتري): خريطة 100×100 مربع، أراضي تُفتح، حقول وأبنية ──
WORLD_N, WORLD_BLOCK, WORLD_SEED = 100, 10, 2026
WORLD_START_BLOCKS = ["4,4", "5,4", "4,5", "5,5"]
WORLD_MAX_OBJS = 400
WORLD_ITEMS = {
    # key: (الاسم, القسم, عرض, عمق, السعر, الإنتاج (دقائق, كوينز) أو None)
    "field": ("حقل", "fields", 2, 2, 40, None),
    "barn": ("حظيرة", "build", 3, 3, 400, (30, 15)), "windmill": ("طاحونة", "build", 3, 3, 450, (35, 18)),
    "house": ("بيت ريفي", "build", 3, 3, 350, (40, 20)), "shed": ("كوخ أدوات", "build", 2, 2, 150, (20, 8)),
    "market": ("بسطة السوق", "market", 2, 2, 250, (20, 10)), "cart": ("عربة خضار", "market", 1, 1, 90, (12, 5)),
    "apple": ("شجرة تفاح", "nature", 1, 1, 60, (15, 4)), "berry": ("شجرة توت", "nature", 1, 1, 80, (20, 6)),
    "oak": ("شجرة سنديان", "nature", 1, 1, 30, None), "pine": ("شجرة صنوبر", "nature", 1, 1, 30, None),
    "bushf": ("شجيرة ورد", "nature", 1, 1, 25, None), "bush": ("شجيرة", "nature", 1, 1, 20, None), "flowers": ("زهور", "nature", 1, 1, 15, None),
    "cow": ("بقرة", "animals", 1, 1, 120, (20, 7)), "sheep": ("خروف", "animals", 1, 1, 90, (15, 5)), "pig": ("خنزير", "animals", 1, 1, 80, (15, 4)),
    "hen": ("دجاجة", "animals", 1, 1, 30, (10, 2)), "hen2": ("دجاجة بنية", "animals", 1, 1, 30, (10, 2)),
    "goose": ("إوزة", "animals", 1, 1, 40, (15, 3)), "rabbit": ("أرنب", "animals", 1, 1, 35, (10, 2)), "chick": ("كتكوت", "animals", 1, 1, 20, (8, 1)),
    "barrel": ("برميل", "decor", 1, 1, 15, None), "crate": ("صندوق", "decor", 1, 1, 15, None), "hay": ("قشّ", "decor", 1, 1, 20, None),
    "wall": ("سور حجري", "decor", 1, 1, 20, None), "lantern": ("فانوس", "decor", 1, 1, 25, None),
    "fence": ("سياج", "decor", 1, 1, 5, None), "road": ("طريق", "decor", 1, 1, 2, None),
}
WORLD_CATS = [("fields", "حقول"), ("build", "مباني"), ("market", "أسواق"), ("nature", "طبيعة"), ("animals", "حيوانات"), ("decor", "ديكور")]
WORLD_WALK = {"field", "road"}          # قطع يمشي عليها الدب
WORLD_CROPS = {   # key: (الاسم, تكلفة البذور, الدقائق, المكافأة)
    "wheat": ("قمح", 6, 2, 9), "potato": ("بطاطا", 10, 4, 15), "carrot": ("جزر", 16, 8, 24), "corn": ("ذرة", 28, 15, 42),
    "tomato": ("طماطم", 45, 25, 68), "sunflower": ("دوار الشمس", 70, 40, 105), "strawberry": ("فراولة", 110, 60, 165), "watermelon": ("بطيخ", 170, 90, 255),
}
_WT = {}

def _wh(x, y, s):
    n = (x * 374761393 + y * 668265263 + s * 1274126177) & 0xFFFFFFFF
    n = ((n ^ (n >> 13)) * 1274126177) & 0xFFFFFFFF
    n ^= n >> 16
    return n / 4294967295.0

def _wnoise(x, y, s):
    x0, y0 = math.floor(x), math.floor(y); fx, fy = x - x0, y - y0
    fx, fy = fx * fx * (3 - 2 * fx), fy * fy * (3 - 2 * fy)
    a, b, c, d = _wh(x0, y0, s), _wh(x0 + 1, y0, s), _wh(x0, y0 + 1, s), _wh(x0 + 1, y0 + 1, s)
    return (a * (1 - fx) + b * fx) * (1 - fy) + (c * (1 - fx) + d * fx) * fy

def _wfbm(x, y, s):
    return (_wnoise(x, y, s) * 0.6 + _wnoise(x * 2.1, y * 2.1, s + 7) * 0.28 + _wnoise(x * 4.3, y * 4.3, s + 13) * 0.12)

def world_terrain():
    """سلسلة من 10000 خانة: 0 عشب، 1 ماء، 2 شجرة، 3 صخرة. نفس الخريطة لكل اللاعبين."""
    if "t" in _WT: return _WT["t"]
    N = WORLD_N; rows = []
    for y in range(N):
        r = []
        for x in range(N):
            dc = math.hypot(x - 50, y - 50)
            wv = _wfbm(x / 11.0, y / 11.0, WORLD_SEED)
            tv = _wfbm(x / 7.0, y / 7.0, WORLD_SEED + 101)
            c = "0"
            if wv > 0.63 and dc > 10: c = "1"
            elif tv > 0.60 and _wh(x, y, 5) < 0.72 and dc > 8.5: c = "2"
            elif _wh(x, y, 9) < 0.014 and dc > 3: c = "3"
            elif dc <= 8.5 and _wh(x, y, 11) < 0.035 and dc > 5.5: c = "2"
            r.append(c)
        rows.append("".join(r))
    _WT["t"] = "".join(rows)
    return _WT["t"]

def world_row(uid):
    with db() as con:
        r = con.execute("SELECT * FROM world WHERE user_id=?", (int(uid),)).fetchone()
        if not r:
            objs = {"1": {"id": 1, "k": "field", "x": 45, "y": 46, "c": "", "t": 0}, "2": {"id": 2, "k": "field", "x": 48, "y": 46, "c": "", "t": 0},
                    "3": {"id": 3, "k": "apple", "x": 53, "y": 47, "c": "", "t": time.time()}}
            con.execute("INSERT INTO world(user_id,objs,lands,cleared,bx,by,nid) VALUES(?,?,?,?,?,?,?)",
                        (int(uid), json.dumps(objs), json.dumps(WORLD_START_BLOCKS), "[]", 50, 50, 4))
            r = con.execute("SELECT * FROM world WHERE user_id=?", (int(uid),)).fetchone()
        return _world_parse(dict(r))

def _world_parse(d):
    def jl(x, dflt):
        try: v = json.loads(x or "")
        except Exception: return dflt
        return v if isinstance(v, type(dflt)) else dflt
    d["objs"] = {str(k): v for k, v in jl(d["objs"], {}).items() if isinstance(v, dict) and v.get("k") in WORLD_ITEMS}
    d["lands"] = [x for x in jl(d["lands"], []) if isinstance(x, str)] or list(WORLD_START_BLOCKS)
    d["cleared"] = [x for x in jl(d["cleared"], []) if isinstance(x, str)]
    return d

def _world_save(con, uid, d):
    con.execute("UPDATE world SET objs=?, lands=?, cleared=?, bx=?, by=?, nid=? WHERE user_id=?",
                (json.dumps(d["objs"], separators=(",", ":")), json.dumps(d["lands"]), json.dumps(d["cleared"]), int(d["bx"]), int(d["by"]), int(d["nid"]), int(uid)))

def world_public(d, coins):
    return {"n": WORLD_N, "block": WORLD_BLOCK, "lands": d["lands"], "cleared": d["cleared"], "bear": {"x": d["bx"], "y": d["by"]},
            "objs": list(d["objs"].values()), "coins": coins, "now": time.time(), "max": WORLD_MAX_OBJS,
            "unlock_cost": _world_unlock_cost(d)}

def _world_unlock_cost(d): return min(1500, 100 + 60 * max(0, len(d["lands"]) - len(WORLD_START_BLOCKS)))

def _world_cells(d, skip=None):
    """الخانات المشغولة: {(x,y): key} لكل القطع (ما عدا skip)."""
    occ = {}
    for o in d["objs"].values():
        if skip is not None and o["id"] == skip: continue
        it = WORLD_ITEMS[o["k"]]
        for i in range(it[2]):
            for j in range(it[3]): occ[(o["x"] + i, o["y"] + j)] = o["k"]
    return occ

def _world_free(d, k, x, y, skip=None):
    it = WORLD_ITEMS[k]; t = world_terrain(); occ = _world_cells(d, skip); lands = set(d["lands"]); cl = set(d["cleared"])
    for i in range(it[2]):
        for j in range(it[3]):
            cx, cy = x + i, y + j
            if not (0 <= cx < WORLD_N and 0 <= cy < WORLD_N): return "خارج الخريطة"
            if f"{cx // WORLD_BLOCK},{cy // WORLD_BLOCK}" not in lands: return "هذي الأرض مقفلة"
            ch = t[cy * WORLD_N + cx]
            if ch == "1": return "ما تكدر تبني على الماء"
            if ch in "23" and f"{cx},{cy}" not in cl: return "امسح الشجرة أو الصخرة أول"
            if (cx, cy) in occ: return "المكان مشغول"
    return None

def _world_walkable(d, x, y):
    if not (0 <= x < WORLD_N and 0 <= y < WORLD_N): return False
    if f"{x // WORLD_BLOCK},{y // WORLD_BLOCK}" not in set(d["lands"]): return False
    ch = world_terrain()[y * WORLD_N + x]
    if ch == "1" or (ch in "23" and f"{x},{y}" not in set(d["cleared"])): return False
    k = _world_cells(d).get((x, y))
    return k is None or k in WORLD_WALK

def api_world(uid, user, body):
    game_row(uid)
    d = world_row(uid)
    with db() as con: coins = con.execute("SELECT stars FROM game WHERE user_id=?", (int(uid),)).fetchone()["stars"]
    out = world_public(d, coins)
    out["terrain"] = world_terrain()
    out["catalog"] = {"cats": [{"key": k, "name": n} for k, n in WORLD_CATS],
                      "items": [{"key": k, "name": v[0], "cat": v[1], "fw": v[2], "fh": v[3], "price": v[4], "prod": ({"min": v[5][0], "coins": v[5][1]} if v[5] else None)} for k, v in WORLD_ITEMS.items()],
                      "crops": [{"key": k, "name": v[0], "cost": v[1], "min": v[2], "reward": v[3]} for k, v in WORLD_CROPS.items()]}
    return out

def api_world_act(uid, user, body):
    if not _rate(("wa", uid), 240, 60): return err("بطّئ شوي", 429)
    act = str(body.get("action") or "")
    game_row(uid)
    def num(k):
        try: return int(body.get(k))
        except (TypeError, ValueError): return None
    d = world_row(uid)
    with db() as con:
        coins = con.execute("SELECT stars FROM game WHERE user_id=?", (int(uid),)).fetchone()["stars"]
        now = time.time(); msg = ""

        def spend(c):
            nonlocal coins
            if coins < c: return False
            coins -= c; return True
        def find(oid):
            o = d["objs"].get(str(oid)); return o
        if act == "place":
            k = str(body.get("k") or ""); x, y = num("x"), num("y")
            if k not in WORLD_ITEMS or x is None or y is None: return err("طلب غير صالح")
            if len(d["objs"]) >= WORLD_MAX_OBJS: return err("العالم امتلأ بالقطع")
            e = _world_free(d, k, x, y)
            if e: return err(e)
            if not spend(WORLD_ITEMS[k][4]): return err("كوينزاتك لا تكفي")
            oid = d["nid"]; d["nid"] += 1
            d["objs"][str(oid)] = {"id": oid, "k": k, "x": x, "y": y, "c": "", "t": now if WORLD_ITEMS[k][5] else 0}
        elif act == "move":
            o = find(num("id")); x, y = num("x"), num("y")
            if not o or x is None or y is None: return err("طلب غير صالح")
            e = _world_free(d, o["k"], x, y, skip=o["id"])
            if e: return err(e)
            o["x"], o["y"] = x, y
        elif act == "sell":
            o = find(num("id"))
            if not o: return err("القطعة غير موجودة")
            refund = WORLD_ITEMS[o["k"]][4] // 2
            if o["c"] and o["k"] == "field": refund += WORLD_CROPS.get(o["c"], ("", 0))[1] // 2
            coins += refund; d["objs"].pop(str(o["id"]))
        elif act == "clear":
            x, y = num("x"), num("y")
            if x is None or y is None or not (0 <= x < WORLD_N and 0 <= y < WORLD_N): return err("طلب غير صالح")
            ch = world_terrain()[y * WORLD_N + x]
            if ch not in "23" or f"{x},{y}" in d["cleared"]: return err("ما كو شي يُمسح هنا")
            if f"{x // WORLD_BLOCK},{y // WORLD_BLOCK}" not in d["lands"]: return err("هذي الأرض مقفلة")
            if not spend(12 if ch == "2" else 25): return err("كوينزاتك لا تكفي")
            d["cleared"].append(f"{x},{y}")
        elif act == "unlock":
            bx, by = num("bx"), num("by")
            if bx is None or by is None or not (0 <= bx < WORLD_N // WORLD_BLOCK and 0 <= by < WORLD_N // WORLD_BLOCK): return err("طلب غير صالح")
            key = f"{bx},{by}"
            if key in d["lands"]: return err("الأرض مفتوحة")
            if not any(f"{bx + a},{by + b}" in d["lands"] for a, b in ((1, 0), (-1, 0), (0, 1), (0, -1))): return err("افتح الأراضي المجاورة أول")
            if not spend(_world_unlock_cost(d)): return err("كوينزاتك لا تكفي")
            d["lands"].append(key)
        elif act == "plant":
            o = find(num("id")); c = WORLD_CROPS.get(str(body.get("crop") or ""))
            if not o or o["k"] != "field" or not c: return err("طلب غير صالح")
            if o["c"]: return err("الحقل مزروع")
            if not spend(c[1]): return err("كوينزاتك لا تكفي لهذي البذور")
            o["c"] = str(body.get("crop")); o["t"] = now
        elif act == "harvest":
            o = find(num("id"))
            if not o or o["k"] != "field" or not o["c"]: return err("ما كو شي يُحصد")
            c = WORLD_CROPS.get(o["c"])
            if not c: o["c"] = ""
            else:
                if now < o["t"] + c[2] * 60 - 2: return err("لسه ما نضج")
                coins += c[3]; msg = f"+{c[3]}"; o["c"] = ""; o["t"] = 0
        elif act == "collect":
            o = find(num("id")); pr = WORLD_ITEMS[o["k"]][5] if o else None
            if not o or not pr: return err("ما كو شي يُجمع هنا")
            if now < o["t"] + pr[0] * 60 - 2: return err("لسه ما جهز")
            coins += pr[1]; msg = f"+{pr[1]}"; o["t"] = now
        elif act == "bear":
            x, y = num("x"), num("y")
            if x is None or y is None or not _world_walkable(d, x, y): return err("ما أكدر أوصل هناك")
            d["bx"], d["by"] = x, y
        else: return err("إجراء غير معروف")
        con.execute("UPDATE game SET stars=? WHERE user_id=?", (coins, int(uid)))
        _world_save(con, uid, d)
    r = world_public(d, coins); r["msg"] = msg
    return r

def api_invite(uid, user, body):
    try: to = int(body.get("to"))
    except (TypeError, ValueError): return err("غير صالح")
    if not _rate(("inv", uid), 10, 3600): return err("دعوات كثيرة، حاول لاحقاً", 429)
    with db() as con:
        if not _are_friends(con, uid, to): return err("لازم يكون صديقك", 403)
        if con.execute("SELECT 1 FROM house_invites WHERE from_id=? AND to_id=? AND created_at>?", (int(uid), to, time.time() - 600)).fetchone(): return err("دعوتَه قبل قليل")
        con.execute("INSERT INTO house_invites(from_id,to_id,created_at) VALUES(?,?,?)", (int(uid), to, time.time()))
        me = _uname(con, uid); hn = house_row(uid)["name"]
    try: bot.send_message(to, f"🏠 {esc(me)} يدعوك لزيارة بيته «{esc(hn)}» بتطبيق الأستاذ. افتح التطبيق ← أصدقائي.", parse_mode="HTML")
    except Exception: pass
    return {"ok": True}

def api_invite_clear(uid, user, body):
    with db() as con: con.execute("DELETE FROM house_invites WHERE to_id=? AND id=?", (int(uid), int(body.get("id") or 0)))
    return {"ok": True}

_api_init_base, _api_focus_base = api_init, api_focus
def api_init(uid, user, body):
    d = _api_init_base(uid, user, body)
    if isinstance(d, dict) and "error" not in d: d["game"] = game_state(uid)
    return d

def api_focus(uid, user, body):
    r = _api_focus_base(uid, user, body)
    if isinstance(r, dict) and r.get("ok"):
        try:
            gain = game_on_focus(uid, int(body.get("minutes")))
            r["stars_gain"] = gain; r["game"] = game_state(uid)
        except Exception: log.exception("game focus")
    return r

ROUTES_GET.update({"/api/init": api_init, "/api/game": api_game, "/api/social": api_social, "/api/rooms": api_rooms, "/api/house": api_house, "/api/town": api_town, "/api/lb": api_lb, "/api/world": api_world})
ROUTES_POST.update({"/api/focus": api_focus, "/api/farm": api_farm, "/api/shop": api_shop, "/api/yt_home": api_yt_home, "/api/yt_follow": api_yt_follow, "/api/yt_related": api_yt_related, "/api/yt_sub": api_yt_sub, "/api/yt_channel": api_yt_channel, "/api/presence": api_presence,
                    "/api/admin/stars_list": api_admin_stars_list, "/api/admin/stars": api_admin_stars,
                    "/api/admin/btns": api_admin_btns, "/api/admin/btn_save": api_admin_btn_save, "/api/admin/btn_del": api_admin_btn_del, "/api/admin/btn_move": api_admin_btn_move, "/api/admin/btn_reset": api_admin_btn_reset,
                    "/api/admin/stats": api_admin_stats, "/api/admin/users": api_admin_users, "/api/admin/ban": api_admin_ban, "/api/admin/subcheck": api_admin_subcheck,
                    "/api/social": api_social, "/api/house": api_house, "/api/house_act": api_house_act, "/api/town": api_town, "/api/lb": api_lb, "/api/world": api_world, "/api/world_act": api_world_act, "/api/town_act": api_town_act, "/api/town_view": api_town_view, "/api/house_view": api_house_view, "/api/invite": api_invite, "/api/invite_clear": api_invite_clear, "/api/friend_add": api_friend_add, "/api/friend_del": api_friend_del, "/api/farm_view": api_farm_view,
                    "/api/room_create": api_room_create, "/api/room_join": api_room_join, "/api/room_state": api_room_state, "/api/room_leave": api_room_leave, "/api/rooms": api_rooms})



# ═══════════════════════════════════════════════════════════════════
#               خادم الويب: خيوط محدودة + مهلات + رفض عند الامتلاء
# ═══════════════════════════════════════════════════════════════════

INDEX_PATH = os.path.join(APP_DIR, "index.html")
_INDEX = {"mtime": 0, "data": b""}
MAX_BODY = 6 * 1024 * 1024
SECURITY_HEADERS = {
    "Content-Security-Policy": ("default-src 'self'; script-src 'self' 'unsafe-inline' https://telegram.org https://cdn.jsdelivr.net; "
                                "style-src 'self' 'unsafe-inline' https://cdn.jsdelivr.net https://fonts.googleapis.com; "
                                "font-src 'self' data: https://cdn.jsdelivr.net https://fonts.gstatic.com; img-src 'self' data: blob: https://i.ytimg.com https://yt3.ggpht.com; connect-src 'self'; frame-src https://www.youtube-nocookie.com https://www.youtube.com; "
                                "frame-ancestors https://web.telegram.org https://*.telegram.org https://telegram.org; base-uri 'none'; form-action 'none'"),
    "X-Content-Type-Options": "nosniff", "Referrer-Policy": "no-referrer", "Cache-Control": "no-store"}

def load_index():
    try:
        mt = os.path.getmtime(INDEX_PATH)
        if mt != _INDEX["mtime"]:
            with open(INDEX_PATH, "rb") as f: raw = f.read()
            ver = MAIN_HASH + hashlib.sha1(raw).hexdigest()[:8]
            _INDEX.update(mtime=mt, ver=ver, data=raw.replace(b"__APP_VERSION__", ver.encode()))
        return _INDEX["data"]
    except OSError:
        return None

def app_version():
    load_index()
    return _INDEX.get("ver") or MAIN_HASH

class PoolHTTPServer(HTTPServer):
    """خادم محمي من الإغراق:
    1) بوابة: الاتصال لا يصل لعمّال الويب إلا بعد أن يُرسل بيانات فعلاً (الاتصالات الخاملة لا تأخذ عاملاً).
    2) مجمّع خيوط ثابت؛ إن امتلأ يرد 503 فوراً.  3) سقف اتصالات لكل IP.  4) مهلة كلية لاستلام الطلب (انظر Handler)."""
    daemon_threads = True
    request_queue_size = 128
    GATE_WAIT = 10          # أقصى انتظار لأول بيانات
    MAX_PENDING = 512

    def __init__(self, addr, handler, workers=HTTP_WORKERS, backlog=48):
        super().__init__(addr, handler)
        self.pool = ThreadPoolExecutor(max_workers=workers, thread_name_prefix="http")
        self._sem = threading.BoundedSemaphore(workers + backlog)
        self._glock = threading.Lock()
        self._conns = defaultdict(int)
        self._incoming = queue.SimpleQueue()
        self._wake_r, self._wake_w = socket.socketpair()
        self._wake_r.setblocking(False)
        self._workers_n = workers
        threading.Thread(target=self._gate_loop, daemon=True, name="http-gate").start()

    def _reject(self, request, ip=None):
        try: request.sendall(b"HTTP/1.0 503 Service Unavailable\r\nConnection: close\r\nRetry-After: 5\r\nContent-Length: 0\r\n\r\n")
        except Exception: pass
        self.shutdown_request(request)
        if ip is not None:
            with self._glock: self._conns[ip] = max(0, self._conns[ip] - 1)

    def process_request(self, request, client_address):
        ip = client_address[0]
        cap = 10**6 if ip in ("127.0.0.1", "::1") else 10      # خلف النفق كل الاتصالات محلية فلا سقف لها
        with self._glock:
            full = self._conns[ip] >= cap
            if not full: self._conns[ip] += 1
        if full: return self._reject(request)
        self._incoming.put((request, client_address, time.time() + self.GATE_WAIT))
        try: self._wake_w.send(b"x")
        except Exception: pass

    def _dispatch(self, request, client_address):
        ip = client_address[0]
        if not self._sem.acquire(blocking=False): return self._reject(request, ip)
        def run():
            try:
                request.setblocking(True)
                self.finish_request(request, client_address)
            except Exception: pass
            finally:
                self.shutdown_request(request); self._sem.release()
                with self._glock: self._conns[ip] = max(0, self._conns[ip] - 1)
        try: self.pool.submit(run)
        except Exception:
            self._sem.release(); self._reject(request, ip)

    def _drop(self, sock, addr):
        self.shutdown_request(sock)
        with self._glock: self._conns[addr[0]] = max(0, self._conns[addr[0]] - 1)

    def _gate_loop(self):
        """لا يُسلَّم اتصال لعامل إلا بعد وصول ترويسات HTTP كاملة؛ الخامل والناقص يُقطعان بالمهلة دون أخذ عامل."""
        sel = selectors.DefaultSelector()
        sel.register(self._wake_r, selectors.EVENT_READ, None)
        waiting, partial = {}, {}      # waiting: لا بيانات بعد (مسجّل بالمحدِّد) | partial: بدأت البيانات ولم تكتمل الترويسات
        while True:
            try:
                try:
                    while True:
                        req, addr, dl = self._incoming.get_nowait()
                        if len(waiting) + len(partial) >= self.MAX_PENDING: self._reject(req, addr[0]); continue
                        try:
                            req.setblocking(False); sel.register(req, selectors.EVENT_READ, None); waiting[req] = (addr, dl)
                        except Exception: self._reject(req, addr[0])
                except queue.Empty:
                    pass
                for key, _ in sel.select(timeout=0.25):
                    if key.fileobj is self._wake_r:
                        try: self._wake_r.recv(4096)
                        except Exception: pass
                        continue
                    sock = key.fileobj; sel.unregister(sock)
                    partial[sock] = waiting.pop(sock)
                now = time.time()
                for sock in list(partial):
                    addr, dl = partial[sock]
                    try: data = sock.recv(8192, socket.MSG_PEEK)
                    except (BlockingIOError, InterruptedError): data = None
                    except Exception: data = b""
                    if data == b"":
                        del partial[sock]; self._drop(sock, addr)
                    elif data and (b"\r\n\r\n" in data or b"\n\n" in data or len(data) >= 8192):
                        del partial[sock]; self._dispatch(sock, addr)
                    elif dl < now:
                        del partial[sock]; self._drop(sock, addr)
                for sock in [s_ for s_, (_a, dl) in waiting.items() if dl < now]:
                    addr, _dl = waiting.pop(sock)
                    try: sel.unregister(sock)
                    except Exception: pass
                    self._drop(sock, addr)
            except Exception:
                log.exception("http gate"); time.sleep(0.2)

    def handle_error(self, request, client_address):
        log.debug("http error from %s", client_address)

class Handler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.0"      # اتصال لكل طلب: لا تبقى خيوط معلّقة
    timeout = 6                        # مهلة كل قراءة
    server_version = "Mudarris"
    HEADER_DEADLINE, BODY_DEADLINE = 10, 30    # مهلة كلية لاستلام الترويسات / الجسم (تقطع slowloris)

    def setup(self):
        super().setup()
        self._timer = None
        self._arm(self.HEADER_DEADLINE)

    def _kill(self):
        try: self.connection.shutdown(socket.SHUT_RDWR)
        except Exception: pass

    def _arm(self, secs):
        self._disarm()
        self._timer = threading.Timer(secs, self._kill); self._timer.daemon = True; self._timer.start()

    def _disarm(self):
        t = getattr(self, "_timer", None)
        if t: t.cancel()

    def finish(self):
        self._disarm()
        try: super().finish()
        except Exception: pass

    def log_message(self, *a): pass

    def _ip(self):
        ip = self.client_address[0]
        if ip in ("127.0.0.1", "::1"):                        # خلف النفق/البروكسي المحلي فقط نصدّق الترويسات
            fwd = self.headers.get("CF-Connecting-IP") or (self.headers.get("X-Forwarded-For") or "").split(",")[0].strip()
            if fwd: return fwd[:64]
        return ip

    def _send(self, code, body, ctype, extra=None):
        try:
            self.send_response(code)
            self.send_header("Content-Type", ctype); self.send_header("Content-Length", str(len(body)))
            self.send_header("Connection", "close")
            for k, v in {**SECURITY_HEADERS, **(extra or {})}.items(): self.send_header(k, v)
            self.end_headers()
            if self.command != "HEAD": self.wfile.write(body)
        except Exception: pass

    def _json(self, code, obj):
        self._send(code, json.dumps(obj, ensure_ascii=False).encode(), "application/json; charset=utf-8")

    def _reject_early(self):
        self.close_connection = True

    def _auth(self):
        user = validate_init_data(self.headers.get("X-Telegram-Init-Data", "")) or session_user(self.headers.get("X-Session", ""))
        if not user: self._json(401, {"error": "الجلسة غير صالحة أو انتهت. أغلق التطبيق وافتحه من تيليجرام من جديد."}); return None
        return user

    def _dispatch(self, fn, user, body):
        uid = int(user["id"])
        if hit_limited(_API_HITS, ("u", uid), 120):
            return self._json(429, {"error": "طلبات كثيرة، تمهّل قليلاً."})
        try: res = fn(uid, user, body)
        except Exception:
            log.exception("api %s", self.path); res = err("خطأ داخلي، جرّب بعد قليل.", 500)
        code = res.pop("_code", 200) if isinstance(res, dict) else 200
        self._json(code, res)

    def _upload(self, user):
        """رفع كتاب خام (بدون JSON) من التطبيق: الوصف في الترويسة X-Meta، والقراءة في مهمة خلفية."""
        uid = int(user["id"]); self.close_connection = True
        if not is_admin(uid): return self._json(403, {"error": "هذه الصفحة للمشرفين فقط."})
        try: n = int(self.headers.get("Content-Length") or 0)
        except ValueError: n = -1
        if n <= 0 or n > MAX_UPLOAD:
            return self._json(413, {"error": f"حجم الملف يجب أن يكون أقل من {MAX_UPLOAD // (1024 * 1024)} ميغابايت. قسّمه وأضف كل جزء على حدة."})
        try: meta = json.loads(unquote(self.headers.get("X-Meta") or "{}"))
        except Exception: return self._json(400, {"error": "بيانات الملف غير صالحة"})
        meta, bad = validate_upload_meta(meta)
        if bad: return self._json(400, {"error": bad})
        if not _UPLOAD_SEM.acquire(blocking=False): return self._json(503, {"error": "رفع آخر جارٍ، حاول بعد قليل."})
        try:
            self._arm(UPLOAD_DEADLINE)
            try:
                self.connection.settimeout(30)
                raw = self.rfile.read(n)
            except Exception:
                raw = b""
            finally:
                self._disarm()
        finally:
            _UPLOAD_SEM.release()
        if len(raw) != n: return self._json(400, {"error": "انقطع الرفع، أعد المحاولة."})
        res = start_upload_job(uid, raw, meta)
        self._json(res.pop("_code", 200), res)

    def do_HEAD(self): self.do_GET()

    def do_GET(self):
        self._disarm()
        path = urlparse(self.path).path
        if hit_limited(_IP_HITS, self._ip(), 240): return self._json(429, {"error": "too many requests"})
        if path in ("/", "/index.html"):
            data = load_index()
            if data is None: return self._send(404, "index.html غير موجود بجانب main.py".encode(), "text/plain; charset=utf-8")
            return self._send(200, data, "text/html; charset=utf-8")
        if path == "/health":
            return self._json(200, {"ok": True, "jobs": len(INFLIGHT)})
        if path == "/api/version": return self._json(200, {"v": app_version()})
        if path == "/api/auth/poll":
            tok = str(dict(parse_qsl(urlparse(self.path).query)).get("token", ""))[:40]
            return self._json(200, login_poll(tok, self.headers.get("User-Agent", "")))
        if path == "/manifest.webmanifest":
            return self._send(200, json.dumps(manifest_dict(), ensure_ascii=False).encode(), "application/manifest+json; charset=utf-8", {"Cache-Control": "public, max-age=3600"})
        if path == "/sw.js":
            return self._send(200, sw_js().encode(), "application/javascript; charset=utf-8", {"Cache-Control": "no-cache", "Service-Worker-Allowed": "/"})
        if path in ICONS:
            return self._send(200, icon_bytes(path), "image/png", {"Cache-Control": "public, max-age=86400"})
        if path.startswith("/static/"):
            fp = os.path.realpath(os.path.join(STATIC_DIR, path[len("/static/"):]))
            if not fp.startswith(os.path.realpath(STATIC_DIR) + os.sep) or not os.path.isfile(fp) or os.path.getsize(fp) > 5 * 1024 * 1024:
                return self._send(404, b"not found", "text/plain")
            stt = os.stat(fp); etag = '"%x-%x"' % (int(stt.st_mtime), stt.st_size)
            if self.headers.get("If-None-Match") == etag: return self._send(304, b"", "text/plain", {"ETag": etag, "Cache-Control": "no-cache"})
            with open(fp, "rb") as f: data = f.read()
            return self._send(200, data, mimetypes.guess_type(fp)[0] or "application/octet-stream", {"ETag": etag, "Cache-Control": "no-cache"})
        fn = ROUTES_GET.get(path)
        if not fn: return self._send(404, b"not found", "text/plain")
        user = self._auth()
        if user: self._dispatch(fn, user, dict(parse_qsl(urlparse(self.path).query)))

    def do_POST(self):
        path = urlparse(self.path).path
        if hit_limited(_IP_HITS, self._ip(), 240):
            self.close_connection = True; return self._json(429, {"error": "too many requests"})
        if path == "/api/auth/start":
            self.close_connection = True
            if hit_limited(_IP_HITS, ("auth", self._ip()), 12, 600): return self._json(429, {"error": "محاولات كثيرة، انتظر قليلاً ثم جرّب."})
            if not BOT_USERNAME: return self._json(503, {"error": "البوت غير جاهز بعد، جرّب بعد لحظات."})
            tok, code = login_create(self.headers.get("User-Agent", ""))
            return self._json(200, {"token": tok, "code": code, "link": f"https://t.me/{BOT_USERNAME}?start=login_{tok}", "expires": LOGIN_TTL})
        if path == "/api/auth/logout":
            self.close_connection = True; session_revoke(self.headers.get("X-Session", "")); return self._json(200, {"ok": True})
        if path == "/api/admin/upload":
            user = self._auth()
            if not user:
                self.close_connection = True; return
            return self._upload(user)
        fn = ROUTES_POST.get(path)
        if not fn:
            self.close_connection = True; return self._send(404, b"not found", "text/plain")
        user = self._auth()
        if not user:
            self.close_connection = True; return
        try: n = int(self.headers.get("Content-Length") or 0)
        except ValueError: n = -1
        if n < 0 or n > MAX_BODY:
            self.close_connection = True; return self._json(413, {"error": "الطلب كبير جداً"})
        self._arm(self.BODY_DEADLINE)
        try: body = json.loads(self.rfile.read(n) or b"{}") if n else {}
        except Exception:
            return self._json(400, {"error": "JSON غير صالح"})
        finally:
            self._disarm()
        if not isinstance(body, dict): return self._json(400, {"error": "طلب غير صالح"})
        self._dispatch(fn, user, body)

def web_supervisor():
    """يعيد تشغيل خادم الويب تلقائياً إن سقط."""
    while True:
        try:
            srv = PoolHTTPServer(("0.0.0.0", WEB_PORT), Handler)
            log.info("🌐 خادم التطبيق يعمل على المنفذ %d (خيوط: %d)", WEB_PORT, HTTP_WORKERS)
            srv.serve_forever(poll_interval=0.5)
        except Exception:
            log.exception("web server crashed — restarting in 3s")
        time.sleep(3)


# ═══════════════════════════════════════════════════════════════════
#                  نفق Cloudflare (إن لم تضبط WEBAPP_URL)
# ═══════════════════════════════════════════════════════════════════

def ensure_cloudflared():
    import shutil
    exe = shutil.which("cloudflared")
    if exe: return exe
    win = platform.system() == "Windows"
    local = os.path.join(APP_DIR, "cloudflared.exe" if win else "cloudflared")
    if os.path.isfile(local): return local
    mach = platform.machine().lower(); sysname = platform.system()
    if sysname == "Linux": asset = "cloudflared-linux-" + ("arm64" if mach in ("aarch64", "arm64") else "arm" if mach.startswith("arm") else "amd64")
    elif win: asset = "cloudflared-windows-amd64.exe"
    else: raise RuntimeError("ثبّت cloudflared يدوياً (brew install cloudflared) أو اضبط WEBAPP_URL")
    log.info("⬇️ تنزيل cloudflared…")
    with requests.get(f"https://github.com/cloudflare/cloudflared/releases/latest/download/{asset}", stream=True, timeout=60) as r:
        r.raise_for_status()
        with open(local, "wb") as f:
            for chunk in r.iter_content(1 << 16): f.write(chunk)
    os.chmod(local, 0o755)
    return local

def apply_webapp_url(url):
    changed = url != WEB["url"]; WEB["url"] = url
    log.info("🌐 رابط التطبيق: %s", url)
    try: bot.set_chat_menu_button(menu_button=types.MenuButtonWebApp(text="🎨 التطبيق", web_app=types.WebAppInfo(url=url)))
    except Exception as e: log.warning("menu button: %s", e)
    if changed:
        for oid in OWNER_IDS:
            try: say(oid, f"🌐 رابط التطبيق الجديد:\n<code>{esc(url)}</code>\n(للرابط الثابت اضبط WEBAPP_URL)")
            except Exception: pass

def ensure_ngrok():
    import shutil, tarfile
    exe = shutil.which("ngrok")
    if exe: return exe
    local = os.path.join(APP_DIR, "ngrok")
    if os.path.isfile(local): return local
    mach = platform.machine().lower()
    arch = "arm64" if mach in ("aarch64", "arm64") else "arm" if mach.startswith("arm") else "amd64"
    if platform.system() != "Linux": raise RuntimeError("ثبّت ngrok يدوياً")
    log.info("⬇️ تنزيل ngrok…")
    url = f"https://bin.equinox.io/c/bNyj1mQVY4c/ngrok-v3-stable-linux-{arch}.tgz"
    tgz = local + ".tgz"
    with requests.get(url, stream=True, timeout=90) as r:
        r.raise_for_status()
        with open(tgz, "wb") as f:
            for chunk in r.iter_content(1 << 16): f.write(chunk)
    with tarfile.open(tgz) as t: t.extract("ngrok", APP_DIR)
    os.remove(tgz); os.chmod(local, 0o755)
    return local

NGROK_PROC = None
def _kill_ngrok(strays=True):
    """يوقف ngrok الحالي، وأي ngrok يتيم بقي من تشغيل سابق (execv لا يقتل العمليات الفرعية)."""
    global NGROK_PROC
    try:
        if NGROK_PROC and NGROK_PROC.poll() is None: NGROK_PROC.terminate()
    except Exception: pass
    if not strays: return
    me = os.getpid()
    try:
        for d in os.listdir("/proc"):
            if not d.isdigit() or int(d) == me: continue
            try: cmd = open(f"/proc/{d}/cmdline", "rb").read().decode("utf-8", "ignore").replace("\0", " ")
            except OSError: continue
            if "ngrok" in cmd and " http " in cmd and "python" not in cmd:
                try: os.kill(int(d), 15)
                except OSError: pass
    except OSError: pass

def ngrok_loop():
    """نفق ngrok بدومين ثابت (مجاني): NGROK_AUTHTOKEN + NGROK_DOMAIN. الرابط لا يتغير بين التشغيلات."""
    delay = 10
    if NGROK_DOMAIN: apply_webapp_url("https://" + NGROK_DOMAIN)
    _kill_ngrok(); time.sleep(3)
    while True:
        try:
            exe = ensure_ngrok()
            cmd = [exe, "http", str(WEB_PORT), "--authtoken", NGROK_TOKEN, "--log", "stdout", "--log-format", "json"]
            if NGROK_DOMAIN: cmd += ["--url", "https://" + NGROK_DOMAIN]
            proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
            global NGROK_PROC; NGROK_PROC = proc
            for line in proc.stdout:
                try: j = json.loads(line)
                except Exception: continue
                if j.get("lvl") == "eror" or j.get("err"): log.warning("ngrok: %s", str(j.get("err") or j.get("msg"))[:200])
                u = j.get("url") or ""
                if j.get("msg") == "started tunnel" and u.startswith("https://") and not NGROK_DOMAIN:
                    apply_webapp_url(u); delay = 10
                if j.get("msg") == "started tunnel": delay = 10
            proc.wait()
        except Exception as e:
            log.warning("ngrok: %s", e)
        time.sleep(delay); delay = min(delay * 2, 300)

def tunnel_loop():
    delay = 10
    while True:
        try:
            exe = ensure_cloudflared()
            cmd = [exe, "tunnel", "--no-autoupdate", "run", "--token", CF_TUNNEL_TOKEN] if CF_TUNNEL_TOKEN else [exe, "tunnel", "--url", f"http://127.0.0.1:{WEB_PORT}", "--no-autoupdate"]
            proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, text=True)
            for line in proc.stderr:
                m = re.search(r"https://[a-z0-9-]+\.trycloudflare\.com", line)
                if m: apply_webapp_url(m.group(0)); delay = 10
            proc.wait()
        except Exception as e:
            log.warning("tunnel: %s", e)
        time.sleep(delay); delay = min(delay * 2, 300)


# ═══════════════════════════════════════════════════════════════════
#                        صيانة دورية + التشغيل
# ═══════════════════════════════════════════════════════════════════

# ── تحديث تلقائي من GitHub: يسحب main.py و index.html ويعيد التشغيل إذا تغيّرا ──
GH_MAIN = _env("GITHUB_MAIN_FILE") or os.path.basename(os.path.abspath(__file__))     # اسم ملف البوت بالمستودع (افتراضياً نفس اسم الملف اللي يشتغل)
GH_FILES = ((GH_MAIN, os.path.abspath(__file__)), ("index.html", os.path.join(APP_DIR, "index.html")), ("world.js", os.path.join(STATIC_DIR, "world.js")), ("atlas.png", os.path.join(STATIC_DIR, "atlas.png")), ("atlas.json", os.path.join(STATIC_DIR, "atlas.json")))
GH_OPTIONAL = {"world.js", "atlas.png", "atlas.json"}
def _gh_hdr():
    h = {"User-Agent": "teacher-bot-updater"}
    if GH_TOKEN: h["Authorization"] = f"token {GH_TOKEN}"
    return h

def _gh_names():
    r = requests.get(f"https://api.github.com/repos/{GH_REPO}/contents/", params={"ref": GH_BRANCH}, headers=_gh_hdr(), timeout=25)
    if r.status_code != 200: raise RuntimeError(f"قائمة المستودع: HTTP {r.status_code}")
    return [x["name"] for x in r.json() if x.get("type") == "file"]

def _gh_find(name):
    """إذا اسم الملف ما طابق ما بالمستودع (مثلاً main-23.py) نختار أقرب ملف من نفس النوع."""
    names = _gh_names(); ext = os.path.splitext(name)[1]
    if name in names: return name
    same = [n for n in names if n.endswith(ext) and n not in ("requirements.txt",)]
    stem = os.path.splitext(name)[0].lower()
    pick = [n for n in same if n.lower().startswith(stem.split("-")[0])] or same
    if len(pick) == 1: return pick[0]
    pick.sort(key=lambda n: [int(x) if x.isdigit() else 0 for x in re.split(r"(\d+)", n)], reverse=True)   # أحدث رقم (main-24 قبل main-23)
    if pick: return pick[0]
    raise RuntimeError(f"ما لقيت ملف {ext} بالمستودع")

def _gh_get(name):
    url = f"https://raw.githubusercontent.com/{GH_REPO}/{GH_BRANCH}/{name}"
    r = requests.get(url, headers=_gh_hdr(), timeout=25)
    if r.status_code == 404:
        alt = _gh_find(name)
        if alt != name: r = requests.get(f"https://raw.githubusercontent.com/{GH_REPO}/{GH_BRANCH}/{alt}", headers=_gh_hdr(), timeout=25); name = alt
    if r.status_code != 200: raise RuntimeError(f"{name}: HTTP {r.status_code}")
    return r.content

def selfupdate(force=False):
    """يرجع قائمة الملفات التي تغيّرت (وطبّقها) أو [] ."""
    if not GH_REPO: return None
    new = {}
    for name, path in GH_FILES:
        try: data = _gh_get(name)
        except Exception:
            if name in GH_OPTIONAL: continue
            raise
        try: cur = open(path, "rb").read()
        except OSError: cur = b""
        if hashlib.sha256(data).digest() == hashlib.sha256(cur).digest(): continue
        if name.endswith(".py"):
            try: compile(data.decode("utf-8"), name, "exec")
            except Exception as e: raise RuntimeError(f"{name} فيه خطأ برمجي، ما انطبّق: {e}")
        elif name.endswith(".html"):
            if b"</html>" not in data: raise RuntimeError(f"{name} غير مكتمل")
        elif name.endswith(".png"):
            if not data.startswith(b"\x89PNG"): raise RuntimeError(f"{name} ليس صورة صالحة")
        elif name.endswith(".json"):
            try: json.loads(data.decode("utf-8"))
            except Exception: raise RuntimeError(f"{name} JSON غير صالح")
        elif name.endswith(".js"):
            if len(data) < 2000 or b"WorldGame" not in data: raise RuntimeError(f"{name} غير مكتمل")
        new[name] = (path, data, cur)
    for name, (path, data, cur) in new.items():
        if cur:
            try: open(path + ".bak", "wb").write(cur)
            except OSError: pass
        os.makedirs(os.path.dirname(path), exist_ok=True)
        tmp = path + ".new"
        with open(tmp, "wb") as f: f.write(data)
        os.replace(tmp, path)
    return list(new)

def _restart():
    log.warning("🔄 تحديث من GitHub: إعادة تشغيل البوت")
    _kill_ngrok(); time.sleep(2)
    try: os.execv(sys.executable, [sys.executable] + sys.argv)
    except Exception as e: log.error("restart failed: %s", e); os._exit(3)

def update_loop():
    time.sleep(20)
    while True:
        try:
            ch = selfupdate()
            if ch: log.info("GitHub: تحدّث %s", ", ".join(ch)); _restart()
        except Exception as e: log.warning("GitHub update: %s", str(e)[:160])
        time.sleep(GH_EVERY * 60)

def maintenance_loop():
    n = 0
    while True:
        time.sleep(600); n += 1
        try:
            now = time.time()
            for store in (_IP_HITS, _API_HITS, _CHAT_HITS):
                for k in [k for k, dq in list(store.items()) if not dq or now - dq[-1] > 3600]: store.pop(k, None)
            for d in (USER_LAST, _UPS):
                for k in [k for k, t in list(d.items()) if now - t > 3600]: d.pop(k, None)
            for k in [k for k, t in list(_BAN.items()) if now - t[1] > 300]: _BAN.pop(k, None)
            for k in [k for k, j in list(_UP_JOBS.items()) if now - j['t'] > 3600]: _UP_JOBS.pop(k, None)
            with db() as con:
                con.execute("DELETE FROM events WHERE id NOT IN (SELECT id FROM events ORDER BY id DESC LIMIT 20000)")
                con.execute("DELETE FROM daily_usage WHERE day < ?", ((datetime.now() - timedelta(days=40)).strftime("%Y-%m-%d"),))
                con.execute("DELETE FROM daily_challenge WHERE day < ?", ((datetime.now() - timedelta(days=14)).strftime("%Y-%m-%d"),))
                con.execute("DELETE FROM opinions WHERE created_at < ?", ((datetime.utcnow() - timedelta(days=30)).isoformat(timespec="seconds"),))
                if n % 6 == 0: con.execute("PRAGMA wal_checkpoint(TRUNCATE)")
        except Exception:
            log.exception("maintenance")

def main():
    global BOT_USERNAME
    if not BOT_TOKEN: raise SystemExit("❌ BOT_TOKEN غير موجود في .env")
    if not OWNER_IDS: log.warning("⚠️ OWNER_ID غير مضبوط — لن تعمل أوامر المالك")
    if not (GROQ_KEYS or GEMINI_KEYS or MISTRAL_KEYS or OPENROUTER_KEYS): log.warning("لا يوجد أي مفتاح ذكاء اصطناعي — البوت لن يجيب")
    if not GEMINI_KEYS: log.warning("ℹ️ بدون Gemini: لا بحث حيّ لآراء الملازم/المدرّسين ولا OCR للصور بالرؤية")
    db_init(); game_init()
    try: BOT_USERNAME = bot.get_me().username or ""
    except Exception as e: log.warning("get_me: %s", e)
    threading.Thread(target=web_supervisor, daemon=True, name="web").start()
    threading.Thread(target=maintenance_loop, daemon=True, name="maint").start()
    if GH_REPO: threading.Thread(target=update_loop, daemon=True, name="ghupdate").start(); log.info("🔁 تحديث تلقائي من GitHub: %s@%s كل %d د", GH_REPO, GH_BRANCH, GH_EVERY)
    if NGROK_TOKEN: threading.Thread(target=ngrok_loop, daemon=True, name="ngrok").start()
    elif not WEB["url"] or CF_TUNNEL_TOKEN: threading.Thread(target=tunnel_loop, daemon=True, name="tunnel").start()
    else: log.info("🌐 رابط التطبيق من الإعدادات: %s", WEB["url"])
    try:
        bot.set_my_commands([types.BotCommand(c, d) for c, d in (("start", "البداية"), ("help", "المساعدة"), ("wazari", "أسئلة وزارية"), ("stats", "إحصائياتي"),
                                                                    ("webapp", "فتح التطبيق"), ("materials", "آخر الملازم"), ("clear", "مسح ذاكرتي"), ("id", "معرّفي"))])
    except Exception as e: log.warning("set commands: %s", e)
    log.info("المدرّس الذكي v3 يعمل — مفاتيح: Groq %d | Gemini %d | Mistral %d | OpenRouter %d | حدّ التزامن %d", len(GROQ_KEYS), len(GEMINI_KEYS), len(MISTRAL_KEYS), len(OPENROUTER_KEYS), LLM_CONCURRENCY)
    while True:
        try:
            bot.remove_webhook()
            bot.infinity_polling(timeout=30, long_polling_timeout=30, allowed_updates=["message", "callback_query", "channel_post"])
        except KeyboardInterrupt:
            raise SystemExit(0)
        except Exception:
            log.exception("polling crashed — restarting in 5s")
            time.sleep(5)

if __name__ == "__main__":
    main()
